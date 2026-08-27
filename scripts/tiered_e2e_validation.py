"""Tiered end-to-end validation for IT/OT reporting parity.

Tier 1 — smoke (low N): wiring, metric availability, reconciliations
Tier 2 — medium (fixed seed): representative numerical checks
Tier 3 — production (500k): run separately; see docs/LIMITATIONS_REGISTER.md

Usage:
  PYTHONPATH=src:. python scripts/tiered_e2e_validation.py --tier 1
  PYTHONPATH=src:. python scripts/tiered_e2e_validation.py --tier 2
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import tempfile
import time
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
COMBINED = ROOT / "model" / "Guided_IT_OT_CRQ_Model_v1_0.xlsx"
OUT_DIR = ROOT / "outputs" / "tiered_e2e"


def _configure(wb, sector: str, asset: str, years: int, seed: int = 20260821, insurance: bool = False):
    from tests.helpers import configure_run

    configure_run(wb, sector, asset)
    if sector == "Financial Services":
        wb["IT 03 - Organisation Inputs"]["C30"] = years
        wb["IT 03 - Organisation Inputs"]["C31"] = seed
    else:
        adj = wb["OT 06 - Assessment Adjustments"]
        for r in range(1, 80):
            key = str(adj.cell(r, 1).value or "").strip().upper()
            if key == "SIMULATIONS":
                adj.cell(r, 4).value = years
            if key == "RANDOM_SEED":
                adj.cell(r, 4).value = seed
    if insurance and "06 Appetite & Insurance" in wb.sheetnames:
        ws = wb["06 Appetite & Insurance"]
        ws["B13"] = 5_000_000  # retention
        ws["B14"] = 25_000_000  # primary limit
        ws["B15"] = 50_000_000  # aggregate programme limit
        # Multi-layer tower
        ws.cell(18, 1).value = "Primary"
        ws.cell(18, 2).value = 5_000_000
        ws.cell(18, 3).value = 20_000_000
        ws.cell(18, 4).value = 1.0
        ws.cell(19, 1).value = "Excess"
        ws.cell(19, 2).value = 25_000_000
        ws.cell(19, 3).value = 25_000_000
        ws.cell(19, 4).value = 1.0


def _reconcile(result: dict) -> list[str]:
    errs = []
    er = result.get("engine_result") or result
    view = str(er.get("reporting_view") or er.get("detailed_view") or "Prudent")
    aal = er.get("prudent_aal") if view.startswith("P") else er.get("best_aal")
    aal = aal if aal is not None else er.get("AAL") or er.get("prudent_aal") or er.get("best_aal")
    actor = er.get("actor_aal") or {}
    scen = er.get("scenario_aal") or {}
    if actor and aal is not None:
        if abs(sum(actor.values()) - aal) > max(1.0, 1e-4 * abs(aal)):
            errs.append(f"actor AAL {sum(actor.values())} != total {aal}")
    if scen and aal is not None:
        if abs(sum(scen.values()) - aal) > max(1.0, 1e-4 * abs(aal)):
            errs.append(f"scenario AAL {sum(scen.values())} != total {aal}")
    comps = er.get("loss_components") or []
    if comps and aal is not None:
        ca = sum(c.get("aal") or 0 for c in comps)
        if abs(ca - aal) > max(1.0, 1e-3 * abs(aal)):
            errs.append(f"component AAL {ca} != total {aal}")
        t95 = er.get("primary_annual_tvar95") or er.get("prudent_tvar95") or er.get("TVaR95")
        t99 = er.get("primary_annual_tvar99") or er.get("prudent_tvar99") or er.get("TVaR99")
        if t95 is not None:
            c95 = sum(c.get("contrib_tvar95") or 0 for c in comps)
            if abs(c95 - t95) > max(1.0, 1e-3 * abs(t95)):
                errs.append(f"component TVaR95 {c95} != {t95}")
        if t99 is not None:
            c99 = sum(c.get("contrib_tvar99") or 0 for c in comps)
            if abs(c99 - t99) > max(1.0, 1e-3 * abs(t99)):
                errs.append(f"component TVaR99 {c99} != {t99}")
    for key in ("best_aal", "best_var95", "best_tvar95", "best_var99", "best_tvar99"):
        if er.get(key) is None:
            errs.append(f"missing {key}")
    pkgs = er.get("control_packages") or []
    for p in pkgs:
        if not p.get("simulated"):
            errs.append(f"package {p.get('name')} not simulated")
        for k in ("aal", "var95", "tvar95", "var99", "tvar99"):
            if p.get(k) is None:
                errs.append(f"package {p.get('name')} missing {k}")
    ins = er.get("insurance_analysis")
    if ins and not ins.get("reconcile_ok", True):
        errs.append("insurance trial reconcile failed")
    return errs


CASES_SMOKE = [
    ("Financial Services", "Organisation", "IT"),
    ("Power Generation", "CCGT", "OT"),
    ("Energy Assets", "Upstream Onshore", "OT"),
    ("Manufacturing", "Process Manufacturing", "OT"),
]


def run_case(sector, asset, domain, years, run_whatifs, sensitivity, label, run_packages=None, insurance=False):
    from it_ot_crq.router import run_combined

    if run_packages is None:
        run_packages = run_whatifs
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prev_e2e = os.environ.get("CRQ_E2E_ALLOW_OT_N")
    # Proportionate smoke/medium OT runs: allow non-500k N when explicitly validating.
    if years != 500_000:
        os.environ["CRQ_E2E_ALLOW_OT_N"] = "1"
    # Allow packages without individual what-ifs for production-scale package checks
    prev_pkg = os.environ.get("CRQ_RUN_PACKAGES")
    os.environ["CRQ_RUN_PACKAGES"] = "1" if run_packages else "0"
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        model = td / "model.xlsx"
        shutil.copy2(COMBINED, model)
        wb = openpyxl.load_workbook(model)
        _configure(wb, sector, asset, years, insurance=insurance)
        wb.save(model)
        wb.close()
        out = OUT_DIR / f"{label}_{sector.replace(' ', '_')}.xlsx"
        env_sens = os.environ.get("CRQ_RUN_SENSITIVITY")
        if sensitivity:
            os.environ["CRQ_RUN_SENSITIVITY"] = "1"
        else:
            os.environ.pop("CRQ_RUN_SENSITIVITY", None)
        t0 = time.time()
        result = run_combined(
            model,
            output=out,
            sector_pack_dir=ROOT / "sector_packs",
            work_dir=td / "work",
            run_whatifs=run_whatifs,
        )
        elapsed = time.time() - t0
        if env_sens is None:
            os.environ.pop("CRQ_RUN_SENSITIVITY", None)
        else:
            os.environ["CRQ_RUN_SENSITIVITY"] = env_sens
        if prev_e2e is None:
            os.environ.pop("CRQ_E2E_ALLOW_OT_N", None)
        else:
            os.environ["CRQ_E2E_ALLOW_OT_N"] = prev_e2e
        if prev_pkg is None:
            os.environ.pop("CRQ_RUN_PACKAGES", None)
        else:
            os.environ["CRQ_RUN_PACKAGES"] = prev_pkg
        errs = _reconcile(result)
        er = result["engine_result"]
        pkgs = er.get("control_packages") or []
        target = next((p for p in pkgs if p.get("name") == "Target"), None)
        return {
            "label": label,
            "sector": sector,
            "domain": domain,
            "basis": er.get("reporting_view") or er.get("detailed_view") or "Prudent",
            "control_state": "baseline+packages" if run_packages else ("packages" if run_whatifs else "baseline"),
            "insurance_state": "active" if er.get("insurance_analysis") else "off",
            "seed": er.get("random_seed"),
            "output": str(out),
            "elapsed_s": round(elapsed, 2),
            "years": er.get("simulation_years"),
            "best_aal": er.get("best_aal"),
            "prudent_aal": er.get("prudent_aal"),
            "best_aal_se": er.get("best_aal_se"),
            "prudent_aal_se": er.get("prudent_aal_se"),
            "best_var95": er.get("best_var95"),
            "prudent_var95": er.get("prudent_var95"),
            "best_tvar95": er.get("best_tvar95"),
            "prudent_tvar95": er.get("prudent_tvar95"),
            "best_var99": er.get("best_var99"),
            "prudent_var99": er.get("prudent_var99"),
            "best_tvar99": er.get("best_tvar99"),
            "prudent_tvar99": er.get("prudent_tvar99"),
            "target_package_aal": None if target is None else target.get("aal"),
            "target_package_var95": None if target is None else target.get("var95"),
            "target_package_tvar95": None if target is None else target.get("tvar95"),
            "target_package_var99": None if target is None else target.get("var99"),
            "target_package_tvar99": None if target is None else target.get("tvar99"),
            "packages": len(pkgs),
            "insurance": bool(er.get("insurance_analysis")),
            "sensitivities": len(er.get("sensitivities") or []),
            "components": len(er.get("loss_components") or []),
            "errors": errs,
            "ok": not errs,
        }


def _convergence(tier2_cases, tier3_cases):
    """Compare medium vs production runs; wider relative tolerance for tails.

    Tolerances are documented sampling bands, not model-equality gates:
    - AAL: 5% OR |Δ| ≤ 2× estimated Tier-2 SE (SE_t3 × √(N3/N2)) when SE present
    - VaR 95: 8% generally; 35% when either VaR is sparse (≤ 1.5× AAL or zero) —
      sparse annual aggregates make the 95% order statistic highly unstable at N=10k
    - TVaR 95/99 and VaR 99: wider bands (12–15% / 10%) for extreme-tail sampling noise
    """
    tolerances = {
        "aal": 0.05,
        "var95": 0.08,
        "tvar95": 0.12,
        "var99": 0.10,
        "tvar99": 0.15,
    }
    by_sector = {c["sector"]: c for c in tier2_cases}
    rows = []
    for t3 in tier3_cases:
        t2 = by_sector.get(t3["sector"])
        if not t2:
            continue
        entry = {"sector": t3["sector"], "domain": t3["domain"], "metrics": {}}
        aal_t2 = t2.get("prudent_aal")
        aal_t3 = t3.get("prudent_aal")
        for metric, tol in tolerances.items():
            k = f"prudent_{metric}" if metric != "aal" else "prudent_aal"
            a, b = t2.get(k), t3.get(k)
            if a is None or b is None:
                continue
            abs_d = float(b) - float(a)
            pct = abs_d / float(a) if a else None
            n2, n3 = t2.get("years") or 10000, t3.get("years") or 500000
            aal_se_t2 = t2.get("prudent_aal_se") if metric == "aal" else None
            aal_se_t3 = t3.get("prudent_aal_se") if metric == "aal" else None
            # Infer Tier-2 SE from Tier-3 when Tier-2 report predates SE fields
            if metric == "aal" and aal_se_t2 is None and aal_se_t3 is not None and n2 and n3:
                aal_se_t2 = float(aal_se_t3) * ((n3 / n2) ** 0.5)
            effective_tol = tol
            sparse_var = False
            if metric == "var95" and aal_t2 and aal_t3:
                sparse_var = (
                    float(a) == 0
                    or float(b) == 0
                    or float(a) <= 1.5 * float(aal_t2)
                    or float(b) <= 1.5 * float(aal_t3)
                )
                if sparse_var:
                    effective_tol = 0.35
            within = abs(pct) <= effective_tol if pct is not None else None
            if metric == "aal" and within is False and aal_se_t2 is not None:
                # Accept when absolute drift is within 2× Tier-2 Monte Carlo SE
                if abs(abs_d) <= 2.0 * float(aal_se_t2):
                    within = True
            entry["metrics"][metric] = {
                "tier2": a,
                "tier3": b,
                "abs_diff": abs_d,
                "pct_diff": pct,
                "tolerance_pct": effective_tol,
                "within_tolerance": within,
                "n_tier2": n2,
                "n_tier3": n3,
                "aal_se_tier2": aal_se_t2,
                "aal_se_tier3": aal_se_t3,
                "sparse_var95": sparse_var if metric == "var95" else None,
                "note": (
                    (
                        "Sparse VaR 95 (near AAL or zero): wider ±35% band — "
                        "95% order statistic is unstable for rare-event aggregates at N=10k."
                        if sparse_var
                        else "AAL sampling uncertainty shrinks roughly with 1/√N; "
                        "VaR/TVaR use wider bands because extreme order statistics converge more slowly."
                    )
                    if metric != "aal"
                    else (
                        "AAL Monte Carlo SE = sample_std/√N from annual aggregates. "
                        "Acceptance: |Δ%|≤5% or |Δ|≤2×SE(Tier-2). 10k→500k reduces SE by ~√50 ≈ 7×."
                    )
                ),
            }
        rows.append(entry)
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tier", type=int, choices=(1, 2, 3), default=1)
    args = ap.parse_args()
    rows = []
    if args.tier == 1:
        for sector, asset, domain in CASES_SMOKE:
            rows.append(
                run_case(sector, asset, domain, years=10000, run_whatifs=False, sensitivity=False, label="tier1_smoke")
            )
    elif args.tier == 2:
        rows.append(run_case("Financial Services", "Organisation", "IT", 10000, True, True, "tier2_medium"))
        rows.append(run_case("Power Generation", "CCGT", "OT", 10000, True, True, "tier2_medium"))
        rows.append(run_case("Manufacturing", "Process Manufacturing", "OT", 10000, True, False, "tier2_medium"))
    else:
        # Production scale: baseline + packages + multi-layer insurance (no full individual what-ifs)
        rows.append(
            run_case(
                "Financial Services",
                "Organisation",
                "IT",
                500000,
                False,
                False,
                "tier3_prod",
                run_packages=True,
                insurance=True,
            )
        )
        rows.append(
            run_case(
                "Power Generation",
                "CCGT",
                "OT",
                500000,
                False,
                False,
                "tier3_prod",
                run_packages=True,
                insurance=True,
            )
        )

    report = {"tier": args.tier, "cases": rows, "all_ok": all(r["ok"] for r in rows)}
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / f"tier{args.tier}_report.json"
    if args.tier == 3:
        t2_path = OUT_DIR / "tier2_report.json"
        if t2_path.is_file():
            t2 = json.loads(t2_path.read_text())
            report["convergence_vs_tier2"] = _convergence(t2.get("cases") or [], rows)
            # Human-readable markdown
            md = ["# Tier-3 vs Tier-2 Monte Carlo convergence\n"]
            for c in report["convergence_vs_tier2"]:
                md.append(f"## {c['sector']} ({c['domain']})\n")
                for m, v in (c.get("metrics") or {}).items():
                    md.append(
                        f"- **{m}**: tier2={v['tier2']:,.2f} tier3={v['tier3']:,.2f} "
                        f"Δ%={100 * (v['pct_diff'] or 0):+.2f}% "
                        f"(tol ±{100 * v['tolerance_pct']:.0f}%) "
                        f"{'PASS' if v['within_tolerance'] else 'REVIEW'}\n"
                    )
                    md.append(f"  - {v['note']}\n")
            (OUT_DIR / "tier3_convergence.md").write_text("".join(md))
    path.write_text(json.dumps(report, indent=2))
    print(json.dumps(report, indent=2))
    if not report["all_ok"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
