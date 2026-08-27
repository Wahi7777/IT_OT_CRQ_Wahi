from it_ot_crq.router import _apply_approved_adjustments
import openpyxl
from tests.paths import COMBINED


def test_rationale_row_exists_when_parameter_changes():
    wb = openpyxl.load_workbook(COMBINED)
    for r in range(16, 30):
        if str(wb["OT 03 - Facility Inputs"].cell(r, 1).value) == "INTERNET_OT":
            wb["OT 03 - Facility Inputs"].cell(r, 3).value = "No"
    applied = _apply_approved_adjustments(wb, "OT", [{
        "finding_id": "OI-R1",
        "quantitative_target": "facility_input",
        "model_field": "INTERNET_OT",
        "conditioned_value": "Yes",
        "approved": "yes",
        "approver": "QA",
        "attribution_confidence": "high",
        "adjustment_rationale": "Confirmed asset",
        "double_counting_assessment": "feasibility only",
        "evidence_polarity": "positive",
    }], "run-r1")
    assert len(applied) == 1
    ws = wb["00 COMMON - OI Audit Trail"]
    found = False
    for r in range(16, ws.max_row + 1):
        if str(ws.cell(r, 1).value) == "OI-R1":
            found = True
            assert any(str(ws.cell(r, c).value) == "approved" for c in range(1, 20))
    assert found
