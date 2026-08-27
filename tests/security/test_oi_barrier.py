"""Outside-in exposure vs exploitability without double counting."""
import shutil

import openpyxl
import pytest

from it_ot_crq.router import _apply_approved_adjustments
from tests.helpers import configure_run
from tests.paths import COMBINED


def _base_row(**kwargs):
    row = {
        "finding_id": "OI-VULN-1",
        "approved": "yes",
        "approver": "QA",
        "attribution_confidence": "high",
        "organisation_attribution": "Example Energy Co",
        "facility_attribution": "Example Power Station",
        "adjustment_rationale": "Confirmed applicable CVE on existing remote access; KEV=Yes; high EPSS.",
        "double_counting_assessment": "Exploitability only; remote access already present so frequency/feasibility unchanged.",
        "cve": "CVE-2024-0001",
        "kev": "Yes",
        "epss": "0.92",
        "mapped_ttp": "T0866",
        "model_stage": "S1",
        "evidence_polarity": "positive",
    }
    row.update(kwargs)
    return row


def test_low_attribution_does_not_change_facility():
    wb = openpyxl.load_workbook(COMBINED)
    ws = wb["OT 03 - Facility Inputs"]
    for r in range(16, 30):
        if str(ws.cell(r, 1).value) == "INTERNET_OT":
            ws.cell(r, 3).value = "No"
    applied = _apply_approved_adjustments(wb, "OT", [_base_row(
        finding_id="OI-LOW",
        quantitative_target="facility_input",
        model_field="INTERNET_OT",
        conditioned_value="Yes",
        attribution_confidence="low",
        double_counting_assessment="feasibility only",
    )], "run-low")
    assert applied == []
    for r in range(16, 30):
        if str(ws.cell(r, 1).value) == "INTERNET_OT":
            assert ws.cell(r, 3).value == "No"


def test_same_finding_cannot_double_count_without_split_assessment():
    wb = openpyxl.load_workbook(COMBINED)
    rows = [
        _base_row(
            quantitative_target="facility_input",
            model_field="REMOTE_ACCESS",
            conditioned_value="Controlled",
            double_counting_assessment="oops",
        ),
        _base_row(
            quantitative_target="ttp_prevent_barrier",
            model_field="T0866",
            conditioned_value="0.7",
            double_counting_assessment="oops",
        ),
    ]
    with pytest.raises(ValueError, match="exposure vs exploitability"):
        _apply_approved_adjustments(wb, "OT", rows, "run-dc")


def test_exploitability_overlay_weakens_barrier_not_frequency(tmp_path, monkeypatch):
    dest = tmp_path / "ot.xlsx"
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    # Remote access already present
    for r in range(16, 30):
        if str(wb["OT 03 - Facility Inputs"].cell(r, 1).value) == "REMOTE_ACCESS":
            wb["OT 03 - Facility Inputs"].cell(r, 3).value = "Controlled"
    wb.save(dest)
    wb.close()
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    from ot_crq.engine import refresh

    base = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    p0 = base["combos"][("Nation State", "Operational Disruption")]["p_success"]
    lam0 = base["lambda_be"]
    fac0 = base["facility_mult"]

    wb = openpyxl.load_workbook(dest)
    _apply_approved_adjustments(wb, "OT", [_base_row(
        quantitative_target="ttp_prevent_barrier",
        model_field="T0866",
        conditioned_value="0.5",
    )], "run-bar")
    wb.save(dest)
    wb.close()

    cond = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    p1 = cond["combos"][("Nation State", "Operational Disruption")]["p_success"]
    assert cond["lambda_be"] == pytest.approx(lam0)
    assert cond["facility_mult"] == pytest.approx(fac0)
    assert p1 >= p0 - 1e-15
    # Barrier weakening must not decrease success probability
    rationale = openpyxl.load_workbook(dest)["00 COMMON - OI Audit Trail"]
    texts = [str(c.value) for row in rationale.iter_rows(min_row=16, max_col=16) for c in row]
    assert any("T0866" in t for t in texts if t != "None")


def test_new_exposure_may_change_feasibility_only():
    wb = openpyxl.load_workbook(COMBINED)
    for r in range(16, 30):
        if str(wb["OT 03 - Facility Inputs"].cell(r, 1).value) == "INTERNET_OT":
            wb["OT 03 - Facility Inputs"].cell(r, 3).value = "No"
    applied = _apply_approved_adjustments(wb, "OT", [_base_row(
        finding_id="OI-EXP-1",
        quantitative_target="facility_input",
        model_field="INTERNET_OT",
        conditioned_value="Yes",
        mapped_ttp="",
        adjustment_rationale="Newly observed internet-facing OT asset.",
        double_counting_assessment="Feasibility/opportunity only; no barrier overlay in this finding.",
    )], "run-exp")
    assert len(applied) == 1
    for r in range(16, 30):
        if str(wb["OT 03 - Facility Inputs"].cell(r, 1).value) == "INTERNET_OT":
            assert wb["OT 03 - Facility Inputs"].cell(r, 3).value == "Yes"
