import csv

import pytest

from crq.io_safety import UnsafePathError, sanitize_untrusted_text, validate_user_path
from it_ot_crq.router import _apply_approved_adjustments, _snapshot_outside_in
import openpyxl

from tests.paths import COMBINED


def test_formula_injection_sanitized():
    for payload in ("=HYPERLINK(\"http://evil\")", "=CMD(1)", "+SUM(1)", "@something", "-1+1"):
        out = sanitize_untrusted_text(payload)
        assert out.startswith("'")
        assert not out.startswith("=") or out.startswith("'=")


def test_path_traversal_rejected(tmp_path):
    with pytest.raises((UnsafePathError, FileNotFoundError, ValueError)):
        validate_user_path(tmp_path / ".." / "etc" / "passwd")


def test_positive_evidence_can_override_no(tmp_path):
    wb = openpyxl.load_workbook(COMBINED)
    ws = wb["OT 03 - Facility Inputs"]
    for r in range(16, 30):
        if str(ws.cell(r, 1).value) == "INTERNET_OT":
            ws.cell(r, 3).value = "No"
    rows = [{
        "finding_id": "OI-1",
        "quantitative_target": "facility_input",
        "model_field": "INTERNET_OT",
        "conditioned_value": "Yes",
        "approved": "yes",
        "adjustment_rationale": "Confirmed internet-facing OT HMI",
        "approver": "QA",
        "evidence_polarity": "positive",
        "attribution_confidence": "high",
    }]
    _apply_approved_adjustments(wb, "OT", rows, "run1")
    for r in range(16, 30):
        if str(ws.cell(r, 1).value) == "INTERNET_OT":
            assert ws.cell(r, 3).value == "Yes"


def test_negative_absence_cannot_override_yes(tmp_path):
    wb = openpyxl.load_workbook(COMBINED)
    ws = wb["OT 03 - Facility Inputs"]
    for r in range(16, 30):
        if str(ws.cell(r, 1).value) == "INTERNET_OT":
            ws.cell(r, 3).value = "Yes"
    rows = [{
        "finding_id": "OI-2",
        "quantitative_target": "facility_input",
        "model_field": "INTERNET_OT",
        "conditioned_value": "No",
        "approved": "yes",
        "adjustment_rationale": "scanner saw nothing",
        "approver": "QA",
        "evidence_polarity": "absence",
        "attribution_confidence": "high",
    }]
    with pytest.raises(ValueError, match="cannot automatically"):
        _apply_approved_adjustments(wb, "OT", rows, "run1")


def test_unapproved_scan_does_not_apply(tmp_path):
    csv_path = tmp_path / "scan.csv"
    with csv_path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["finding_id", "approved", "quantitative_target", "model_field", "conditioned_value"])
        w.writeheader()
        w.writerow({
            "finding_id": "OI-3",
            "approved": "no",
            "quantitative_target": "facility_input",
            "model_field": "INTERNET_OT",
            "conditioned_value": "Yes",
        })
    wb = openpyxl.load_workbook(COMBINED)
    status, approved = _snapshot_outside_in(wb, csv_path)
    assert approved == []
    assert "0 approved" in status


def test_aal_multiplier_target_rejected():
    wb = openpyxl.load_workbook(COMBINED)
    rows = [{
        "finding_id": "OI-X",
        "quantitative_target": "aal_multiplier",
        "model_field": "AAL",
        "conditioned_value": "1.5",
        "approved": "yes",
        "adjustment_rationale": "score",
        "approver": "nobody",
        "attribution_confidence": "high",
    }]
    with pytest.raises(ValueError, match="not supported"):
        _apply_approved_adjustments(wb, "OT", rows, "run1")


def test_injected_csv_snapshot_is_text(tmp_path):
    csv_path = tmp_path / "evil.csv"
    with csv_path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=["finding_id", "note"])
        w.writeheader()
        w.writerow({"finding_id": "OI-E", "note": "=HYPERLINK(\"http://x\")"})
    wb = openpyxl.load_workbook(COMBINED)
    _snapshot_outside_in(wb, csv_path)
    found = False
    for row in wb["00 COMMON - Outside-In"].iter_rows(min_row=22, max_row=30, max_col=4, values_only=True):
        for v in row:
            if isinstance(v, str) and "HYPERLINK" in v:
                found = True
                assert v.startswith("'")
    assert found
