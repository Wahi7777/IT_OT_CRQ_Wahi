"""Freeze production-N Control What-If results. Invoked by validate-release --production."""
from __future__ import annotations

import json
import os
import shutil
import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("OT_CRQ_USE_XLSX_BACKEND", "1")

from crq.io_safety import sha256_file  # noqa: E402
from it_ot_crq.router import run_combined  # noqa: E402
from tests.helpers import configure_run  # noqa: E402

COMBINED = ROOT / "model" / "Guided_IT_OT_CRQ_Model_v1_0.xlsx"
if not COMBINED.is_file():
    COMBINED = ROOT / "model" / "Guided_IT_OT_CRQ_Combined_Model_v0_3.xlsx"

CASES = [
    ("Financial Services", "Organisation", "whatif_fs_500k.json", True),
    ("Power Generation", "CCGT", "whatif_pg_500k.json", False),
    ("Energy Assets", "Upstream Onshore", "whatif_ea_500k.json", False),
    ("Manufacturing", "Process Manufacturing", "whatif_mf_500k.json", False),
]


def _top(whatifs, n=12):
    rows = []
    for x in whatifs[:n]:
        pr = x.get("prudent") or {}
        rows.append({
            "control_id": x.get("cid"),
            "control_name": x.get("name"),
            "channel": x.get("channel"),
            "current_maturity": x.get("current"),
            "whatif_maturity": x.get("whatif") or x.get("next"),
            "whatif_aal": pr.get("aal"),
            "aal_reduction": pr.get("reduction"),
            "whatif_tvar95": pr.get("tvar95"),
            "whatif_tvar99": pr.get("tvar99"),
            "whatif_frequency": pr.get("event_freq"),
            "mapped": x.get("mapped", True),
        })
    return rows


def freeze_one(sector, asset, filename, is_it, out_dir: Path, work: Path):
    import openpyxl

    src = work / f"in_{filename.replace('.json', '')}.xlsx"
    dest = work / f"out_{filename.replace('.json', '')}.xlsx"
    shutil.copy2(COMBINED, src)
    wb = openpyxl.load_workbook(src)
    before = {}
    ctl_name = "IT 05 - Control Assessment" if is_it else "OT 04 - Control Assessment"
    if ctl_name not in wb.sheetnames:
        ctl_name = "IT 05 Control Assessment" if is_it else "04 Control Assessment"
    ctl = wb[ctl_name]
    for r in range(16, 80):
        cid = str(ctl.cell(r, 1).value or "").strip()
        if cid:
            before[cid] = ctl.cell(r, 5).value if is_it else ctl.cell(r, 4).value
    configure_run(wb, sector, asset)
    if is_it:
        name = "IT 03 - Organisation Inputs" if "IT 03 - Organisation Inputs" in wb.sheetnames else "IT 03 Organisation Inputs"
        wb[name]["C30"] = 500_000
    wb.save(src)
    wb.close()
    input_hash = sha256_file(src)
    t0 = time.perf_counter()
    result = run_combined(
        src,
        output=dest,
        sector_pack_dir=ROOT / "sector_packs",
        work_dir=work / "it_stage",
        run_whatifs=True,
    )
    elapsed = time.perf_counter() - t0
    er = result["engine_result"]
    whatifs = er.get("whatifs") or []
    for x in whatifs:
        pr = x.get("prudent") or {}
        be = x.get("be") or {}
        if pr.get("reduction", 0) < -1e-6 or be.get("reduction", 0) < -1e-6:
            raise SystemExit(f"Monotonicity fail {sector} {x.get('cid')}: {pr.get('reduction')} / {be.get('reduction')}")
    after_wb = openpyxl.load_workbook(dest)
    ctl = after_wb["IT 05 - Control Assessment"] if is_it else after_wb["OT 04 - Control Assessment"]
    after = {}
    for r in range(16, 80):
        cid = str(ctl.cell(r, 1).value or "").strip()
        if cid:
            after[cid] = ctl.cell(r, 5).value if is_it else ctl.cell(r, 4).value
    after_wb.close()
    if before != after:
        raise SystemExit(f"What-If mutated control maturities for {sector}")
    ranking = [float((x.get("prudent") or {}).get("reduction") or 0) for x in whatifs]
    if ranking != sorted(ranking, reverse=True):
        raise SystemExit(f"What-If ranking not by AAL reduction for {sector}")
    payload = {
        "sector": sector,
        "asset": asset,
        "engine": er.get("engine_version"),
        "pack": result.get("pack_id"),
        "seed": er.get("random_seed"),
        "N": er.get("simulation_years"),
        "input_sha256": input_hash,
        "elapsed_seconds": elapsed,
        "validation": result.get("validation"),
        "baseline_aal": er.get("prudent_aal"),
        "baseline_tvar95": er.get("prudent_tvar95"),
        "baseline_tvar99": er.get("prudent_tvar99"),
        "baseline_successful_event_frequency": er.get("prudent_event_frequency"),
        "whatifs_ran": er.get("whatifs_ran"),
        "n_controls": len(whatifs),
        "monotonicity_prudent": min((x.get("prudent") or {}).get("reduction", 0) for x in whatifs) if whatifs else 0,
        "top_controls": _top(whatifs),
    }
    path = out_dir / filename
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n")
    print(f"WROTE {path} validation={payload['validation']} elapsed={elapsed:.1f}s n={payload['n_controls']}")
    return payload


def main():
    out_dir = ROOT / "tests" / "fixtures"
    work = ROOT / "_work" / "whatif_freeze"
    work.mkdir(parents=True, exist_ok=True)
    only = sys.argv[1:]
    for sector, asset, filename, is_it in CASES:
        if only and filename not in only and sector not in only:
            continue
        freeze_one(sector, asset, filename, is_it, out_dir, work)


if __name__ == "__main__":
    main()
