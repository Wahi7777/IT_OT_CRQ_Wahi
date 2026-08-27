"""Engine numeric snapshot equals Output Bridge / Engine Results (no Excel recalc required)."""
from __future__ import annotations

import openpyxl
import pytest

from tests.paths import ROOT

FREEZE = ROOT / "_work" / "golden_freeze"
CASES = [
    ("out_regression_it_fs_500k.xlsx", "regression_it_fs_500k.json"),
    ("out_regression_ot_power_generation.xlsx", "regression_ot_power_generation.json"),
    ("out_regression_ot_energy_assets.xlsx", "regression_ot_energy_assets.json"),
    ("out_regression_ot_manufacturing.xlsx", "regression_ot_manufacturing.json"),
]


def _num(v):
    if v is None:
        return None
    return float(v)


@pytest.mark.production
@pytest.mark.parametrize("xlsx_name, json_name", CASES)
def test_engine_written_bridge_matches_golden(xlsx_name, json_name):
    import json

    xlsx = FREEZE / xlsx_name
    fixture = ROOT / "tests" / "fixtures" / json_name
    if not xlsx.is_file() or not fixture.is_file():
        pytest.skip("Production freeze workbook not present")
    data = json.loads(fixture.read_text())
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    try:
        bridge_name = "00 COMMON - Output Bridge" if "00 COMMON - Output Bridge" in wb.sheetnames else "00 Output Bridge"
        engine_name = "00 COMMON - Engine Results" if "00 COMMON - Engine Results" in wb.sheetnames else "00 Engine Results"
        bridge = wb[bridge_name]
        engine = wb[engine_name]
        dash = wb["00 Dashboard"]
        assert not str(bridge["B15"].value or "").startswith("=")
        assert abs(_num(bridge["C15"].value) - data["prudent_aal"]) <= max(1.0, 1e-9 * abs(data["prudent_aal"]))
        assert abs(_num(engine["C3"].value) - data["prudent_aal"]) <= max(1.0, 1e-9 * abs(data["prudent_aal"]))
        assert abs(_num(bridge["I15"].value) - data["best_aal"]) <= max(1.0, 1e-9 * abs(data["best_aal"]))
        # Dashboard KPI formulas remain presentation selectors onto the Bridge.
        assert "Output Bridge" in str(dash["A8"].value)
        assert dash["B71"].value is not None
        view_aal = _num(dash["B71"].value)
        assert abs(view_aal - data["prudent_aal"]) <= max(1.0, 1e-9 * abs(data["prudent_aal"]))
    finally:
        wb.close()
