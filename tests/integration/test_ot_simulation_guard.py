"""OT engine fail-closed simulation count (no 500k draw)."""
import openpyxl
import pytest

from tests.helpers import configure_run
from tests.paths import COMBINED


def test_ot_rejects_non_governed_simulation_count(tmp_path, monkeypatch):
    dest = tmp_path / "ot.xlsx"
    import shutil
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    freq = wb["OT 06 - Assessment Adjustments"]
    for r in range(16, 22):
        if str(freq.cell(r, 1).value) == "SIMULATIONS":
            freq.cell(r, 4).value = 1000
    wb["OT 03 - Facility Inputs"]["C21"] = "CCGT"
    wb.save(dest)
    wb.close()
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    from ot_crq.engine import refresh
    with pytest.raises(ValueError, match="500,000"):
        refresh(str(dest), str(dest), run_whatifs=False, simulate=True)
