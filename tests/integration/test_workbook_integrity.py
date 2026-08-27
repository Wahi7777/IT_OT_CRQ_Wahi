"""Workbook structural integrity (no Excel app required)."""
import zipfile

import openpyxl
from openpyxl.cell.cell import Cell

from tests.paths import COMBINED, EXPECTED_TAB_ORDER

ERROR_TOKENS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")


def test_combined_opens_and_roundtrips(tmp_path):
    from it_ot_crq.reporting import REPORTING_SHEETS

    wb = openpyxl.load_workbook(COMBINED)
    out = tmp_path / "roundtrip.xlsx"
    wb.save(out)
    wb.close()
    wb2 = openpyxl.load_workbook(out)
    required = [n for n in EXPECTED_TAB_ORDER if n not in REPORTING_SHEETS]
    assert set(required) <= set(wb2.sheetnames)
    ordered = [n for n in wb2.sheetnames if n in required]
    assert ordered == [n for n in required if n in wb2.sheetnames]
    wb2.close()


def test_no_cached_error_tokens():
    wb = openpyxl.load_workbook(COMBINED, data_only=False)
    hits = []
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(max_row=min(ws.max_row or 1, 80), max_col=min(ws.max_column or 1, 20)):
            for cell in row:
                if isinstance(cell, Cell) and isinstance(cell.value, str):
                    if any(tok in cell.value for tok in ERROR_TOKENS) and not str(cell.value).startswith("="):
                        hits.append((name, cell.coordinate, cell.value))
    wb.close()
    assert hits == []


def test_ooxml_has_workbook_xml():
    with zipfile.ZipFile(COMBINED) as zf:
        names = zf.namelist()
    assert "xl/workbook.xml" in names
    assert "[Content_Types].xml" in names
    assert not any(n.endswith(".bin") and "vba" in n.lower() for n in names)


def test_output_bridge_has_no_independent_var_formula():
    wb = openpyxl.load_workbook(COMBINED)
    ws = wb["00 COMMON - Output Bridge"]
    aal = str(ws["B15"].value)
    wb.close()
    assert "Executive Summary" in aal
    assert "PERCENTILE" not in aal.upper()


def test_risk_transfer_is_labelled_diagnostic():
    wb = openpyxl.load_workbook(COMBINED)
    text = str(wb["00 Risk Transfer"]["A2"].value).lower()
    wb.close()
    assert "not a policy-pricing engine" in text or "indicative" in text


def test_it_06_bia_is_not_title_filled():
    wb = openpyxl.load_workbook(COMBINED)
    ws = wb["IT 06 - Impact & BIA Overrides"]
    title = str(ws["A1"].value or "")
    assert "Impact" in title
    repeats = sum(
        1
        for row in ws.iter_rows(min_row=1, max_row=20, max_col=8)
        for cell in row
        if cell.value == ws["A1"].value
    )
    assert repeats <= 2
    assert ws["A5"].value not in (None, ws["A1"].value)
    wb.close()
