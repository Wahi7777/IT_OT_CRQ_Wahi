"""Dashboard LEC chart: annual aggregate loss on X, exceedance probability on Y."""
import openpyxl

from it_ot_crq.dashboards import populate_common_dashboards
from tests.paths import COMBINED


def test_aggregate_lec_chart_axes_and_series(tmp_path):
    wb = openpyxl.load_workbook(COMBINED)
    src = wb["IT CALC - Aggregate LECs"]
    rows = [
        (0.5, 0.0, 0.0, 0.0, 0.0),
        (0.2, 500.0, 450.0, 1000.0, 900.0),
        (0.1, 2500.0, 2250.0, 5000.0, 4500.0),
        (0.05, 6000.0, 5500.0, 12000.0, 11000.0),
    ]
    for i in range(11):
        r = 16 + i
        if i < len(rows):
            p, ba, bo, pa, po = rows[i]
            src.cell(r, 2).value = p
            src.cell(r, 3).value = ba
            src.cell(r, 4).value = bo
            src.cell(r, 5).value = pa
            src.cell(r, 6).value = po
        else:
            for c in range(2, 7):
                src.cell(r, c).value = None
    result = {
        "actor_aal": {"Nation-state": 10.0, "Cybercriminal": 20.0},
        "scenario_aal": {"Critical business-service disruption": 30.0},
        "prudent_var95": 12000.0,
        "prudent_var99": 25000.0,
        "best_var95": 6000.0,
        "best_var99": 15000.0,
    }
    meta = {"domain": "IT", "reporting_view": "Prudent"}
    populate_common_dashboards(wb, result, meta)

    dash = wb["00 Dashboard"]
    assert dash["A14"].value == 0.5
    assert dash["A14"].number_format == "0.0%"
    assert dash["B16"].value == 5000.0
    assert dash["A12"].value == "ANNUAL AGGREGATE LOSS EXCEEDANCE CURVE"
    assert "TVaR" in str(dash["A28"].value or "")
    assert dash["B26"].value == 12000.0

    chart = dash._charts[0]
    assert type(chart).__name__ == "ScatterChart"
    assert chart.title.tx.rich.p[0].r[0].t == "Annual Aggregate Loss Exceedance Curve"
    assert chart.x_axis.title.tx.rich.p[0].r[0].t == "Annual aggregate loss ($)"
    assert chart.y_axis.title.tx.rich.p[0].r[0].t == "Annual exceedance probability"
    # AEP, OEP, VaR 95 marker, VaR 99 marker
    assert len(chart.series) == 4
    assert "VaR 95" in str(chart.series[2].title)
    assert "VaR 99" in str(chart.series[3].title)
    # Zero-loss point filtered from plot series
    cache_x = [float(p.v) for p in chart.series[0].xVal.numRef.numCache.pt]
    cache_y = [float(p.v) for p in chart.series[0].yVal.numRef.numCache.pt]
    assert cache_x == [1000.0, 5000.0, 12000.0]
    assert cache_y == [0.2, 0.1, 0.05]

    out = tmp_path / "lec_chart.xlsx"
    wb.save(out)
    wb.close()
    again = openpyxl.load_workbook(out)
    ch = again["00 Dashboard"]._charts[0]
    assert type(ch).__name__ == "ScatterChart"
    assert len(ch.series) == 4
    again.close()
