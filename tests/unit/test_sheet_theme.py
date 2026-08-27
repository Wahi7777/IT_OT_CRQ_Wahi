from crq.sheet_theme import (
    FONT_NAME,
    HEADER_FILL,
    ROLE_CALC,
    ROLE_PACK,
    apply_column_headers,
    role_for_sheet,
)


def test_role_table_is_track_agnostic():
    assert role_for_sheet("IT CALC - Attack Path") == ROLE_CALC
    assert role_for_sheet("OT CALC - Attack Path") == ROLE_CALC
    assert HEADER_FILL[ROLE_CALC] == "0F6B78"
    assert HEADER_FILL[ROLE_PACK] == "D99400"
    assert FONT_NAME == "Arial"


def test_apply_column_headers_uses_calc_fill(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IT CALC - Attack Path"
    ws["A15"] = "Stage"
    apply_column_headers(ws, 15, 1, ROLE_CALC)
    assert ws["A15"].font.name == "Arial"
    assert ws["A15"].font.size == 10
    assert ws["A15"].fill.fgColor.rgb in {"0F6B78", "000F6B78", "FF0F6B78"}
