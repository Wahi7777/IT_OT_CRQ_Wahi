"""OT pack sheet 09 scenario parameters: row-oriented load and fail-closed checks."""
import openpyxl
import pytest

from ot_crq.packs import load_ot_pack, rewrite_ot_pack_scenario_parameters
from tests.paths import ROOT

PACK = ROOT / "sector_packs" / "OT" / "OT_CRQ_Sector_Pack_Power_Generation_v1_6.xlsx"


def _ttp_ids(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    ids = []
    for r in range(16, 80):
        tid = str(wb["12 TTP Catalogue"].cell(r, 1).value or "").strip()
        if tid:
            ids.append(tid)
    wb.close()
    return ids


def test_pg_pack_09_is_row_oriented_and_loads():
    wb = openpyxl.load_workbook(PACK, data_only=False)
    ws = wb["09 Scenario Parameters"]
    headers = [ws.cell(15, c).value for c in range(1, 11)]
    assert headers == [
        "Scenario ID",
        "Scenario",
        "Parameter ID",
        "P50",
        "P99",
        "Unit",
        "Applies?",
        "Evidence grade",
        "Source ID",
        "Description / rationale",
    ]
    for c in range(11, 40):
        assert ws.cell(15, c).value in (None, "")
    params = [ws.cell(r, 3).value for r in range(16, 26)]
    assert params.count("DOWNTIME_DAYS") == 5
    assert params.count("CAPACITY_AFFECTED") == 5
    blob = " ".join(str(ws.cell(r, c).value or "") for r in range(15, 30) for c in range(1, 12))
    for banned in (
        "RESTORE_DAYS",
        "AFFECTED_RECORD",
        "PAYMENT_FLOW",
        "FRAUD_RECOVERY",
        "MONITORING_TAKEUP",
        "CHURN_REVENUE",
        "DIVERTED_SHARE",
        "EXTORTION_DIRECT",
        "AFFECTED_ENDPOINT",
        "AFFECTED_SERVER",
        "AFFECTED_SERVICE_SHARE",
    ):
        assert banned not in blob
    wb.close()

    loaded = load_ot_pack(PACK, "Power Generation", "CCGT", _ttp_ids(PACK), "PG-v1.6")
    # Accepted golden Impact/BIA calibration (restored into pack 09).
    assert loaded["scenario_impact"]["Safety System Compromise"]["base_days"] == 14
    assert loaded["scenario_impact"]["Safety System Compromise"]["stress_days"] == 120
    assert loaded["scenario_impact"]["Safety System Compromise"]["base_capacity"] == 0.8
    assert loaded["scenario_impact"]["Safety System Compromise"]["stress_capacity"] == 1.0
    assert loaded["scenario_impact"]["Operational Disruption"]["base_capacity"] == 0.75
    assert loaded["scenario_impact"]["Operational Disruption"]["stress_capacity"] == 1.0


def test_pack_09_rejects_removed_it_parameter(tmp_path):
    dest = tmp_path / "pg.xlsx"
    dest.write_bytes(PACK.read_bytes())
    wb = openpyxl.load_workbook(dest)
    ws = wb["09 Scenario Parameters"]
    ws.cell(16, 3, "AFFECTED_RECORD_SHARE")
    wb.save(dest)
    wb.close()
    with pytest.raises(ValueError, match="unsupported Parameter ID"):
        load_ot_pack(dest, "Power Generation", "CCGT", _ttp_ids(PACK), "PG-v1.6")


def test_pack_09_rejects_blank_p99(tmp_path):
    dest = tmp_path / "pg.xlsx"
    dest.write_bytes(PACK.read_bytes())
    wb = openpyxl.load_workbook(dest)
    ws = wb["09 Scenario Parameters"]
    ws.cell(16, 5).value = None
    wb.save(dest)
    wb.close()
    with pytest.raises(ValueError, match="P50 and P99 must be populated"):
        load_ot_pack(dest, "Power Generation", "CCGT", _ttp_ids(PACK), "PG-v1.6")


def test_rewrite_helper_is_idempotent(tmp_path):
    dest = tmp_path / "pg.xlsx"
    dest.write_bytes(PACK.read_bytes())
    rewrite_ot_pack_scenario_parameters(dest)
    first = load_ot_pack(dest, "Power Generation", "CCGT", _ttp_ids(PACK), "PG-v1.6")["scenario_impact"]
    rewrite_ot_pack_scenario_parameters(dest)
    second = load_ot_pack(dest, "Power Generation", "CCGT", _ttp_ids(PACK), "PG-v1.6")["scenario_impact"]
    assert first == second
