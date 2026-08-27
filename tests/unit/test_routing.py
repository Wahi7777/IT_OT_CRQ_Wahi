"""Sector routing fail-closed tests. Engines are mocked so IT never runs OT and vice versa."""
from pathlib import Path

import openpyxl
import pytest

from it_ot_crq.router import SECTOR_DOMAIN, run_combined
from tests.helpers import configure_run
from tests.paths import COMBINED, EXPECTED_TAB_ORDER


def _copy_with_sector(tmp_path: Path, sector: str, asset: str = "CCGT") -> Path:
    dest = tmp_path / "assessment.xlsx"
    wb = openpyxl.load_workbook(COMBINED)
    configure_run(wb, sector, asset)
    wb.save(dest)
    wb.close()
    return dest


def test_tab_order():
    from it_ot_crq.reporting import REPORTING_SHEETS

    wb = openpyxl.load_workbook(COMBINED, read_only=True)
    # Seven presentation sheets are created on first populate/run; template may omit them.
    required = [n for n in EXPECTED_TAB_ORDER if n not in REPORTING_SHEETS]
    assert set(required) <= set(wb.sheetnames)
    ordered = [n for n in wb.sheetnames if n in required]
    assert ordered == [n for n in required if n in wb.sheetnames]
    wb.close()


@pytest.mark.parametrize(
    "sector,asset,expect_domain",
    [
        ("Financial Services", "Organisation", "IT"),
        ("Power Generation", "CCGT", "OT"),
        ("Energy Assets", "Upstream Onshore", "OT"),
        ("Manufacturing", "Process Manufacturing", "OT"),
    ],
)
def test_router_dispatches_correct_engine(tmp_path, monkeypatch, sector, asset, expect_domain):
    called = {"it": 0, "ot": 0}

    def fake_it(*_a, **_k):
        called["it"] += 1
        return {
            "pack_id": "FS-v1.1.1",
            "pack_status": "working",
            "engine_label": "it_crq FS-v1.1.1",
            "validation": "PASS",
            "best_aal": 1,
            "prudent_aal": 2,
            "prudent_tvar99": 3,
            "actor_aal": {},
            "scenario_aal": {},
        }

    def fake_ot(*_a, **_k):
        called["ot"] += 1
        return {
            "sector_pack_id": "PG-v1.6",
            "sector_pack_status": "working",
            "engine_label": "ot_crq PG-v1.6",
            "pack_id": "PG-v1.6",
            "pack_status": "working",
            "validation": "PASS",
            "overall": "PASS",
            "best_aal": 1,
            "prudent_aal": 2,
            "prudent_tvar99": 3,
            "actor_aal": {},
            "scenario_aal": {},
            "asset_type": asset,
        }

    monkeypatch.setattr("it_ot_crq.router._run_it", fake_it)
    monkeypatch.setattr("it_ot_crq.router._run_ot", fake_ot)
    model = _copy_with_sector(tmp_path, sector, asset)
    out = tmp_path / "out.xlsx"
    result = run_combined(model, output=out, sector_pack_dir=Path("sector_packs"))
    assert result["domain"] == expect_domain
    if expect_domain == "IT":
        assert called["it"] == 1 and called["ot"] == 0
    else:
        assert called["ot"] == 1 and called["it"] == 0
    assert SECTOR_DOMAIN[sector] == expect_domain


def test_unknown_sector_fails_closed(tmp_path):
    dest = tmp_path / "assessment.xlsx"
    wb = openpyxl.load_workbook(COMBINED)
    wb["00 COMMON - Run Setup"]["C6"] = "OT"
    wb["00 COMMON - Run Setup"]["C7"] = "Retail Banking"
    wb["00 COMMON - Run Setup"]["C8"] = "CCGT"
    wb.save(dest)
    wb.close()
    with pytest.raises(ValueError, match="Unsupported sector"):
        run_combined(dest, output=tmp_path / "out.xlsx")


def test_refuses_canonical_overwrite(tmp_path):
    with pytest.raises(ValueError, match="canonical"):
        run_combined(COMBINED, output=COMBINED)


def test_run_setup_is_domain_first():
    wb = openpyxl.load_workbook(COMBINED)
    setup = wb["00 COMMON - Run Setup"]
    assert str(setup["A6"].value).strip() == "Model domain"
    assert str(setup["C6"].value) in {"IT", "OT"}
    assert str(setup["C7"].value) in SECTOR_DOMAIN
    wb.close()
