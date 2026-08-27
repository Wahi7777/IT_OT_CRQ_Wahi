"""Control What-If channel behaviour, CRN, monotonicity helpers, basis default."""
import shutil

import numpy as np
import openpyxl
import pytest

from it_crq.engine import _simulate
from tests.helpers import configure_run
from tests.paths import COMBINED


def test_default_basis_is_prudent():
    wb = openpyxl.load_workbook(COMBINED)
    assert str(wb["00 COMMON - Run Setup"]["C9"].value).strip() == "Prudent"
    assert str(wb["00 COMMON - Run Setup"]["A9"].value).strip() == "Basis"
    assert str(wb["00 Dashboard"]["G4"].value).strip() == "Basis"
    wb.close()


def test_dashboard_metric_labels_are_basis_independent():
    wb = openpyxl.load_workbook(COMBINED)
    labels = [str(wb["00 Dashboard"][a].value or "") for a in ("A7", "D7", "G7", "J7")]
    joined = " ".join(labels)
    assert "Prudent AAL" not in joined
    assert "Prudent TVaR" not in joined
    assert "AAL" in labels[0] or "ANNUAL AVERAGE LOSS" in labels[0]
    wb.close()


def test_it_simulate_crn_same_seed():
    cells = [{
        "event_rate": 0.2, "actor": "Nation-state", "scenario": "Ransomware",
        "blocks": [{"mu": 10.0, "sigma": 0.4}],
    }]
    a = _simulate(cells, 3000, 20260821, 1.0, ["Nation-state"], ["Ransomware"], 0.3)
    b = _simulate(cells, 3000, 20260821, 1.0, ["Nation-state"], ["Ransomware"], 0.3)
    assert np.array_equal(a["annual"], b["annual"])


def test_ot_prevent_does_not_change_lambda(tmp_path, monkeypatch):
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    dest = tmp_path / "pg.xlsx"
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    ctl = wb["OT 04 - Control Assessment"]
    prevent_row = None
    for r in range(16, 70):
        if str(ctl.cell(r, 3).value or "") == "Prevent / Resist":
            prevent_row = r
            break
    assert prevent_row
    from ot_crq.engine import refresh
    wb.save(dest)
    wb.close()
    base = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    p0 = base["combos"][("Nation State", "Operational Disruption")]["p_success"]
    lam0 = base["lambda_be"]
    wb = openpyxl.load_workbook(dest)
    ctl = wb["OT 04 - Control Assessment"]
    cur = str(ctl.cell(prevent_row, 4).value)
    levels = ["Absent", "Initial", "Developing", "Managed", "Optimised"]
    nxt = levels[min(levels.index(cur) + 1, len(levels) - 1)] if cur in levels else "Managed"
    ctl.cell(prevent_row, 4).value = nxt
    wb.save(dest)
    wb.close()
    cond = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    assert cond["lambda_be"] == pytest.approx(lam0)
    p1 = cond["combos"][("Nation State", "Operational Disruption")]["p_success"]
    assert p1 <= p0 + 1e-15


def test_ot_recover_does_not_change_attempt_rate(tmp_path, monkeypatch):
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    dest = tmp_path / "pg.xlsx"
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    ctl = wb["OT 04 - Control Assessment"]
    row = None
    for r in range(16, 70):
        if str(ctl.cell(r, 3).value or "") == "Recover / Restore":
            row = r
            break
    assert row
    from ot_crq.engine import refresh
    wb.save(dest)
    wb.close()
    base = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    lam0 = base["lambda_be"]
    rec0 = base["combos"][("Nation State", "Operational Disruption")]["recover_factor"]
    wb = openpyxl.load_workbook(dest)
    ctl = wb["OT 04 - Control Assessment"]
    cur = str(ctl.cell(row, 4).value)
    levels = ["Absent", "Initial", "Developing", "Managed", "Optimised"]
    nxt = levels[min(levels.index(cur) + 1, len(levels) - 1)] if cur in levels else "Managed"
    ctl.cell(row, 4).value = nxt
    wb.save(dest)
    wb.close()
    cond = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    assert cond["lambda_be"] == pytest.approx(lam0)
    rec1 = cond["combos"][("Nation State", "Operational Disruption")]["recover_factor"]
    assert rec1 <= rec0 + 1e-12


def test_ot_unmapped_excluded_control_no_p_success_change(tmp_path, monkeypatch):
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    dest = tmp_path / "pg.xlsx"
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    ctl = wb["OT 04 - Control Assessment"]
    row = None
    for r in range(16, 70):
        ch = str(ctl.cell(r, 3).value or "")
        if "Excluded" in ch or "no modelled" in ch.lower():
            row = r
            break
    from ot_crq.engine import refresh
    wb.save(dest)
    wb.close()
    if row is None:
        pytest.skip("No excluded control in assessment sheet")
    base = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    p0 = base["combos"][("Nation State", "Operational Disruption")]["p_success"]
    wb = openpyxl.load_workbook(dest)
    ctl = wb["OT 04 - Control Assessment"]
    ctl.cell(row, 4).value = "Optimised"
    wb.save(dest)
    wb.close()
    cond = refresh(str(dest), str(dest), run_whatifs=False, simulate=False)
    p1 = cond["combos"][("Nation State", "Operational Disruption")]["p_success"]
    assert abs(p1 - p0) < 1e-12
