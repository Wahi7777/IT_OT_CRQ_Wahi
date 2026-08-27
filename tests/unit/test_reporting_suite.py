"""Reporting suite sheet creation and population smoke test."""
from __future__ import annotations

import openpyxl

from it_ot_crq.reporting import REPORTING_SHEETS
from it_ot_crq.reporting.populate import populate_reporting_suite
from tests.paths import COMBINED


def test_reporting_suite_sheets_populate(tmp_path):
    wb = openpyxl.load_workbook(COMBINED)
    result = {
        "best_aal": 1.0,
        "prudent_aal": 2.0,
        "best_var95": 3.0,
        "prudent_var95": 4.0,
        "best_tvar95": 5.0,
        "prudent_tvar95": 6.0,
        "best_var99": 7.0,
        "prudent_var99": 8.0,
        "best_tvar99": 9.0,
        "prudent_tvar99": 10.0,
        "best_pany": 0.1,
        "prudent_pany": 0.2,
        "appetite_status": "Tolerance not set",
        "executive_narrative": "The facility has an estimated 10.0% probability of at least one successful material cyber event in the next year.",
        "formation": {
            "campaigns_per_year": 1.5,
            "applicable_campaigns_per_year": 1.5,
            "conditional_success_probability": 0.2,
            "successful_events_per_year": 0.3,
            "p_any_successful_event": 0.2,
            "aal": 2.0,
        },
        "scenario_analysis": [
            {
                "name": "Operational Disruption",
                "campaign_frequency": 0.5,
                "successful_event_frequency": 0.1,
                "annual_event_probability": 0.1,
                "aal": 1.0,
                "var95": 2.0,
                "tvar95": 3.0,
                "var99": 4.0,
                "tvar99": 5.0,
                "pct_total_aal": 0.5,
                "contrib_tvar95": 1.5,
                "contrib_tvar99": 2.5,
                "primary_operational_driver": "Downtime days",
                "primary_financial_driver": "Business interruption",
            }
        ],
        "actor_analysis": [],
        "actor_scenario_rows": [],
        "whatifs": [],
        "control_packages": [
            {
                "name": "Target",
                "n_controls": 3,
                "baseline_aal": 2.0,
                "aal": 1.0,
                "var95": 1.0,
                "tvar95": 1.0,
                "var99": 1.0,
                "tvar99": 1.0,
                "aal_reduction": 1.0,
                "tvar99_reduction": 1.0,
                "simulated": True,
            }
        ],
        "insurance_analysis": None,
        "sensitivities": [],
        "annual_revenue_at_risk": 100.0,
    }
    meta = {"domain": "OT", "reporting_view": "Prudent"}
    populate_reporting_suite(wb, result, meta)
    for name in REPORTING_SHEETS:
        assert name in wb.sheetnames
        assert wb[name]["A1"].value
    assert "VaR 95" in str(wb["01 Executive Risk Story"]["A7"].value) or wb["01 Executive Risk Story"]["A7"].value
    assert wb["01 Executive Risk Story"]["B12"].value == "Tolerance not set"
    assert wb["06 Appetite & Insurance"]["A13"].value == "INSURANCE_RETENTION"
    out = tmp_path / "reporting.xlsx"
    wb.save(out)
    again = openpyxl.load_workbook(out)
    assert again["03 Scenario Analysis"]["A6"].value == "Operational Disruption"
    again.close()


def test_formation_stage_block_does_not_overwrite_route_rows():
    """IT emits many actor×scenario×route rows; stage diagnostic must sit below them."""
    wb = openpyxl.load_workbook(COMBINED)
    rows = []
    for i in range(40):
        rows.append({
            "actor": "Cybercriminal",
            "scenario": "Cyber-enabled theft or fraud",
            "route": f"Route-{i}",
            "campaign_frequency": 0.01,
            "applicable_frequency": 0.01,
            "success_probability": 0.1,
            "successful_event_frequency": 0.001,
            "annual_event_probability": 0.001,
            "aal": 100.0,
        })
    result = {
        "best_aal": 4000.0,
        "prudent_aal": 4000.0,
        "best_var95": 1.0,
        "prudent_var95": 1.0,
        "best_tvar95": 1.0,
        "prudent_tvar95": 1.0,
        "best_var99": 1.0,
        "prudent_var99": 1.0,
        "best_tvar99": 1.0,
        "prudent_tvar99": 1.0,
        "best_pany": 0.1,
        "prudent_pany": 0.1,
        "formation": {"aal": 4000.0, "campaigns_per_year": 1.0, "applicable_campaigns_per_year": 1.0,
                      "conditional_success_probability": 0.1, "successful_events_per_year": 0.1,
                      "p_any_successful_event": 0.1},
        "actor_scenario_rows": rows,
        "scenario_analysis": [],
        "actor_analysis": [],
        "whatifs": [],
        "control_packages": [],
        "sensitivities": [],
    }
    populate_reporting_suite(wb, result, {"domain": "IT", "reporting_view": "Prudent"})
    ws = wb["02 Risk Formation"]
    assert ws["C15"].value == "Route"
    # Last data row is 16+39 = 55; stage header must be below that
    assert ws["A55"].value == "Cybercriminal"
    assert ws["C55"].value == "Route-39"
    assert ws["A34"].value != "STAGE DIAGNOSTIC (business-readable)"
    assert ws["A58"].value == "STAGE DIAGNOSTIC (business-readable)"
    assert ws["A59"].value == "Stage code"
