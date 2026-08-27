from crq.pack_sheet_guide import IT_PACK_BLURBS, OT_PACK_BLURBS, apply_pack_sheet_blurb


def test_every_standard_pack_sheet_has_a_blurb():
    expected = [
        "00 Pack Guide",
        "01 Pack Metadata",
        "02 Actor Weights",
        "03 Actor Scenario",
        "04 Route Definitions",
        "05 Actor Route",
        "06 Scenario Route",
        "07 Stage Requirements",
        "08 Stage Priors",
        "09 Scenario Parameters",
        "10 Driver Rates",
        "11 Impact Driver Matrix",
        "12 TTP Catalogue",
        "13 Route TTP Map",
        "14 Control Reference",
        "15 Control TTP Map",
        "16 Sources",
        "17 Evidence Register",
        "18 Validation",
    ]
    assert list(IT_PACK_BLURBS) == expected
    assert list(OT_PACK_BLURBS) == expected
    for table in (IT_PACK_BLURBS, OT_PACK_BLURBS):
        for purpose, how in table.values():
            assert len(purpose) > 20
            assert len(how) > 20


def test_apply_writes_how_used_on_non_guide(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "12 TTP Catalogue"
    ws["A1"] = "12 — Catalogue"
    apply_pack_sheet_blurb(ws, "OT")
    assert "ATT&CK" in str(ws["A2"].value)
    assert str(ws["A3"].value).startswith("How used in the model:")
    assert "fail closed" in str(ws["A3"].value).lower() or "governed" in str(ws["A3"].value).lower()
    wb.close()
