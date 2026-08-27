"""OT sensitivity live reruns: actor-mix, route, stage-barrier."""

from __future__ import annotations

import math
import os
import shutil

import openpyxl
import pytest

from tests.helpers import configure_run
from tests.paths import COMBINED


@pytest.fixture
def ot_pg_workbook(tmp_path, monkeypatch):
    monkeypatch.setenv("OT_CRQ_USE_XLSX_BACKEND", "1")
    monkeypatch.setenv("CRQ_E2E_ALLOW_OT_N", "1")
    dest = tmp_path / "pg.xlsx"
    shutil.copy2(COMBINED, dest)
    wb = openpyxl.load_workbook(dest)
    configure_run(wb, "Power Generation", "CCGT")
    adj = wb["OT 06 - Assessment Adjustments"]
    for r in range(1, 80):
        if str(adj.cell(r, 1).value or "").strip().upper() == "SIMULATIONS":
            adj.cell(r, 4).value = 10000
        if str(adj.cell(r, 1).value or "").strip().upper() == "RANDOM_SEED":
            adj.cell(r, 4).value = 20260821
    wb.save(dest)
    wb.close()
    return dest


def test_ot_sensitivity_live_actor_route_stage(ot_pg_workbook, monkeypatch):
    monkeypatch.setenv("CRQ_RUN_SENSITIVITY", "1")
    from ot_crq.engine import refresh

    result = refresh(str(ot_pg_workbook), str(ot_pg_workbook), run_whatifs=False, run_packages=False, run_sensitivity=True)
    kinds = {s.get("kind") for s in result["sensitivities"]}
    assert "actor_mix" in kinds
    assert "stage_barrier" in kinds
    actor = next(s for s in result["sensitivities"] if s.get("kind") == "actor_mix")
    assert actor["aal"] is not None
    assert abs(sum(actor["shocked_actor_weights"].values()) - 1.0) < 1e-12
    assert all(v >= 0 for v in actor["shocked_actor_weights"].values())
    for key in ("aal", "var95", "tvar95", "var99", "tvar99"):
        assert actor.get(key) is not None
        assert actor.get(f"delta_{key}" if key != "aal" else "delta_aal") is not None or key == "aal"
    stage = [s for s in result["sensitivities"] if s.get("kind") == "stage_barrier"]
    assert len(stage) >= 2
    for s in stage:
        assert s["aal"] is not None
        for v in (s.get("shocked_stage_priors") or {}).values():
            assert 0.0 <= v <= 1.0
    # Route may be live or limitation if no enabled route
    route = [s for s in result["sensitivities"] if s.get("kind") == "route"]
    assert route
    if route[0].get("aal") is not None:
        assert route[0].get("baseline_route_value") not in (None, "None", "No", "")


def test_ot_sensitivity_cleared_when_disabled(ot_pg_workbook, monkeypatch):
    monkeypatch.delenv("CRQ_RUN_SENSITIVITY", raising=False)
    from ot_crq.engine import refresh

    result = refresh(str(ot_pg_workbook), str(ot_pg_workbook), run_whatifs=False, run_packages=False, run_sensitivity=False)
    assert result.get("sensitivities") == []
