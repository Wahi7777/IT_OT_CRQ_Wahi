"""Authoritative sector-pack registry. Engines must not fall back to embedded packs."""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

REGISTRY_REL = Path("config/sector_pack_registry.json")
OT_PACK_SCHEMA = "CRQ-OT-PACK-1.0"
IT_PACK_SCHEMA = "CRQ-PACK-1.1"

DOMAINS = {"IT", "OT"}
SECTORS_BY_DOMAIN = {
    "IT": ("Financial Services",),
    "OT": ("Power Generation", "Energy Assets", "Manufacturing"),
}


@dataclass(frozen=True)
class PackRecord:
    domain: str
    sector: str
    pack_id: str
    pack_version: str
    relative_file_path: str
    schema_version: str
    pack_status: str
    minimum_engine_version: str
    unit_of_analysis: str
    description: str

    def path(self, project_root: Path) -> Path:
        return (project_root / self.relative_file_path).resolve()


def project_root_from(start: Path | None = None) -> Path:
    here = (start or Path(__file__)).resolve()
    for candidate in [here, *here.parents]:
        if (candidate / REGISTRY_REL).is_file():
            return candidate
        if (candidate / "config" / "sector_pack_registry.json").is_file():
            return candidate
    return Path.cwd().resolve()


def load_registry(project_root: Path | None = None) -> list[PackRecord]:
    root = project_root or project_root_from()
    path = root / REGISTRY_REL
    if not path.is_file():
        raise FileNotFoundError(f"Sector-pack registry missing: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    packs = []
    for row in data.get("packs") or []:
        packs.append(PackRecord(**{k: row[k] for k in PackRecord.__dataclass_fields__}))
    if len(packs) < 4:
        raise ValueError("Registry must list all four governed sector packs.")
    return packs


def resolve_pack(domain: str, sector: str, project_root: Path | None = None) -> PackRecord:
    domain = str(domain or "").strip()
    sector = str(sector or "").strip()
    if domain not in DOMAINS:
        raise ValueError(f"Unsupported domain {domain!r}. Expected IT or OT.")
    allowed = SECTORS_BY_DOMAIN[domain]
    if sector not in allowed:
        raise ValueError(
            f"Sector {sector!r} is not valid for domain {domain}. Expected one of: {', '.join(allowed)}"
        )
    root = project_root or project_root_from()
    matches = [p for p in load_registry(root) if p.domain == domain and p.sector == sector]
    if len(matches) != 1:
        raise ValueError(f"Registry must contain exactly one pack for {domain}/{sector}.")
    rec = matches[0]
    pack_path = rec.path(root)
    if not pack_path.is_file():
        raise FileNotFoundError(
            f"Selected sector pack is missing: {pack_path}. "
            "The engine will not use embedded workbook calibration."
        )
    return rec


def validate_pack_file(rec: PackRecord, project_root: Path, engine_version: str, domain: str) -> dict:
    import openpyxl

    path = rec.path(project_root)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if "01 Pack Metadata" not in wb.sheetnames:
            raise ValueError(f"Pack {path.name} is missing sheet '01 Pack Metadata'.")
        meta = {}
        ws = wb["01 Pack Metadata"]
        for row in ws.iter_rows(min_row=1, max_row=80, max_col=2, values_only=True):
            key = str(row[0] or "").strip()
            if key:
                meta[key] = row[1]
        sector = str(meta.get("SECTOR") or "").strip()
        pack_id = str(meta.get("PACK_ID") or "").strip()
        schema = str(meta.get("SCHEMA_VERSION") or "").strip()
        pack_domain = str(meta.get("DOMAIN") or "").strip()
        if sector != rec.sector:
            raise ValueError(f"Pack sector {sector!r} does not match registry {rec.sector!r}.")
        if pack_id != rec.pack_id:
            raise ValueError(f"Pack ID {pack_id!r} does not match registry {rec.pack_id!r}.")
        if pack_domain and pack_domain != domain:
            raise ValueError(f"Pack domain {pack_domain!r} does not match selected domain {domain!r}.")
        expected_schema = IT_PACK_SCHEMA if domain == "IT" else OT_PACK_SCHEMA
        if schema != rec.schema_version or schema != expected_schema:
            raise ValueError(
                f"Unsupported pack schema {schema!r} for {rec.pack_id}; expected {expected_schema}."
            )
        min_engine = str(meta.get("MIN_ENGINE_VERSION") or rec.minimum_engine_version).strip()
        if _version_tuple(min_engine) > _version_tuple(engine_version):
            raise ValueError(
                f"Pack {rec.pack_id} requires engine {min_engine}; running {engine_version}."
            )
        status = str(meta.get("STATUS") or rec.pack_status).strip()
        if not status:
            raise ValueError(f"Pack {rec.pack_id} has no calibration status.")
        return {"meta": meta, "path": path, "record": rec}
    finally:
        wb.close()


def _version_tuple(value) -> tuple[int, ...]:
    parts = []
    for token in str(value or "0").split("."):
        digits = "".join(ch for ch in token if ch.isdigit())
        parts.append(int(digits or 0))
    return tuple(parts or (0,))
