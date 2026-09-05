"""Read-only materialization of governed assumptions into immutable ModelBundle objects."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl

from crq.application.models import ModelBundle, json_value
from crq.io_safety import sha256_file
from crq.pack_registry import resolve_pack, validate_pack_file
from crq.versions import METHODOLOGY_VERSION, MODEL_BUNDLE_SCHEMA_VERSION
from it_ot_crq.router import IT_ENGINE_VERSION, OT_ENGINE_VERSION


def _sheet_snapshot(workbook: Path, selected_sheets: set[str] | None = None) -> dict[str, list[list[Any]]]:
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    try:
        output = {}
        for ws in wb.worksheets:
            if selected_sheets is not None and ws.title not in selected_sheets:
                continue
            rows = []
            for row in ws.iter_rows(values_only=True):
                values = [json_value(v) for v in row]
                while values and values[-1] is None:
                    values.pop()
                if values and any(v not in (None, "") for v in values):
                    rows.append(values)
            output[ws.title] = rows
        return output
    finally:
        wb.close()


def load_model_bundle(domain: str, sector: str, *, project_root: str | Path | None = None) -> ModelBundle:
    """Load and hash every governed pack/core source required by the unchanged engine."""
    root = Path(project_root).resolve() if project_root else Path(__file__).resolve().parents[3]
    rec = resolve_pack(domain, sector, root)
    engine_version = IT_ENGINE_VERSION if domain == "IT" else OT_ENGINE_VERSION
    checked = validate_pack_file(rec, root, engine_version, domain)
    pack_path = checked["path"]
    source_map_path = root / "contracts" / "mappings" / "model-bundle-source-map.json"
    source_map = json.loads(source_map_path.read_text(encoding="utf-8"))
    template = root / "model" / "Guided_IT_OT_CRQ_Model_v1_0.xlsx"
    template_wb = openpyxl.load_workbook(template, read_only=True)
    try:
        core_sheets = {name for name in template_wb.sheetnames if name.startswith(f"{domain} CORE")}
    finally:
        template_wb.close()
    if domain == "OT":
        core_sheets.update({"OT 06 - Assessment Adjustments"})
    payload = {
        "schema_version": MODEL_BUNDLE_SCHEMA_VERSION,
        "bundle_id": f"{domain}-{rec.pack_id}-engine-{engine_version}",
        "model_bundle_version": "1.0.0",
        "engine_version": engine_version,
        "methodology_version": METHODOLOGY_VERSION,
        "sector_pack": {
            "pack_id": rec.pack_id,
            "pack_version": rec.pack_version,
            "pack_schema_version": rec.schema_version,
            "pack_hash": sha256_file(pack_path),
            "source": rec.relative_file_path,
            "status": rec.pack_status,
        },
        "applicability": {"domain": domain, "sectors": [sector], "asset_types": _asset_types(root, domain, sector)},
        "assumptions": {
            "actors": {"source_map": _entries(source_map, domain, "actors")},
            "scenarios": {"source_map": _entries(source_map, domain, "scenarios")},
            "routes_or_ttps": {"source_map": _entries(source_map, domain, "routes_or_ttps")},
            "control_mappings": {"source_map": _entries(source_map, domain, "control_mappings")},
            "severity_priors": {"source_map": _entries(source_map, domain, "severity_priors")},
            "frequency_priors": {"source_map": _entries(source_map, domain, "frequency_priors")},
            "dependency": {"source_map": _entries(source_map, domain, "dependency")},
            "caps_and_floors": {"source_map": _entries(source_map, domain, "caps_and_floors")},
            "governed_mappings": {
                "pack_workbook_snapshot": _sheet_snapshot(pack_path),
                "core_workbook_snapshot": _sheet_snapshot(template, core_sheets),
                "source_map": source_map,
            },
        },
        "calibration": {
            "status": rec.pack_status,
            "limitations": ["Working/reference calibration status is preserved from the source pack; no recalibration performed."],
            "validation_status": "registry-and-pack-metadata-validated",
            "approved_by": None,
            "approved_at": None,
        },
        "effective_date": "2026-08-27",
        "supersedes_bundle_id": None,
        "source_hashes": {
            "sector_pack": sha256_file(pack_path),
            "combined_workbook": sha256_file(template),
            "pack_registry": sha256_file(root / "config" / "sector_pack_registry.json"),
            "source_map": sha256_file(source_map_path),
        },
    }
    return ModelBundle.from_dict(payload)


def _entries(source_map: dict, domain: str, prefix: str) -> list[dict[str, Any]]:
    return [row for row in source_map.get(domain, []) if str(row.get("target", "")).startswith(f"assumptions.{prefix}")]


def _asset_types(root: Path, domain: str, sector: str) -> list[str]:
    data = json.loads((root / "config" / "sector_pack_registry.json").read_text(encoding="utf-8"))
    for row in data["packs"]:
        if row["domain"] == domain and row["sector"] == sector:
            values = row.get("asset_types")
            if values:
                return list(values)
    defaults = {"Financial Services": ["Organisation"], "Power Generation": ["CCGT"], "Energy Assets": ["Upstream Onshore"], "Manufacturing": ["Process Manufacturing"]}
    return defaults[sector]
