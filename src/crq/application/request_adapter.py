"""Convert a caller-safe public assessment payload into CRQAssessment."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

import openpyxl

from crq.application.errors import ErrorCode, public_error
from crq.application.models import CRQAssessment, json_value
from crq.versions import INPUT_SCHEMA_VERSION


_TOP_LEVEL = {"schema_version", "assessment", "scope", "runtime", "risk_appetite", "insurance", "outside_in", "domain_inputs", "evidence"}
_IDENTITY = {"assessment_id", "assessment_version", "organisation", "domain", "sector", "asset_type", "reporting_basis", "currency"}
_RUNTIME = {"simulation_count", "random_seed", "run_whatifs", "run_packages", "run_sensitivity"}
_BANNED_KEYS = {"model_bundle", "model_bundle_id", "requested_model_bundle_id", "governed_overrides", "assumptions", "actor_priors", "scenario_priors", "frequency_priors", "severity_priors", "dependency", "correlation", "caps_and_floors", "control_mappings", "route_ttp_mappings", "pack_path", "workbook_path", "compatibility", "assessment_hash"}


def public_assessment_payload(assessment: CRQAssessment | Mapping[str, Any]) -> dict[str, Any]:
    """Project a canonical assessment onto fields accepted at the public boundary."""
    payload = assessment.to_dict() if isinstance(assessment, CRQAssessment) else json_value(assessment)
    result = dict(payload)
    for key in ("assessment_hash", "compatibility", "requested_model_bundle_id", "governed_overrides"):
        result.pop(key, None)
    result["runtime"] = dict(result["runtime"])
    result["runtime"].pop("legacy_simulation_mode", None)
    inputs = dict(result["domain_inputs"])
    inputs["controls"] = [
        row for row in inputs.get("controls") or []
        if row.get("control_id") != "CONTROL ASSESSMENT FIELD GUIDE"
    ]
    result["domain_inputs"] = inputs
    return result


def assessment_from_public_payload(value: Mapping[str, Any], *, project_root: str | Path | None = None) -> CRQAssessment:
    payload = json_value(value)
    if not isinstance(payload, dict):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "assessment must be a JSON object.")
    if payload.get("schema_version") != INPUT_SCHEMA_VERSION:
        raise public_error(ErrorCode.SCHEMA_VERSION_UNSUPPORTED, "The assessment schema version is unsupported.", supported=INPUT_SCHEMA_VERSION)
    _reject_unknown(payload, _TOP_LEVEL, "assessment")
    _reject_banned(payload)
    _validate_public_shapes(payload)
    identity = _object(payload, "assessment")
    _reject_unknown(identity, _IDENTITY, "assessment.assessment")
    domain = identity.get("domain")
    sector = identity.get("sector")
    if domain not in {"IT", "OT"}:
        raise public_error(ErrorCode.UNSUPPORTED_DOMAIN, "The assessment domain is unsupported.", domain=domain)
    sectors = {"IT": {"Financial Services"}, "OT": {"Power Generation", "Energy Assets", "Manufacturing"}}
    if sector not in sectors[domain]:
        raise public_error(ErrorCode.UNSUPPORTED_SECTOR, "The assessment sector is unsupported for its domain.", domain=domain, sector=sector)
    outside_in = _object(payload, "outside_in")
    if outside_in.get("apply") is True:
        raise public_error(ErrorCode.VALIDATION_FAILED, "Outside-in application requires a future server-side approved evidence resolver.")
    if outside_in.get("approved_adjustments") not in (None, []):
        raise public_error(ErrorCode.INVALID_REQUEST, "Caller-supplied governed outside-in adjustments are forbidden.")
    runtime = _object(payload, "runtime")
    _reject_unknown(runtime, _RUNTIME, "assessment.runtime")
    if runtime.get("run_whatifs") or runtime.get("run_packages") or runtime.get("run_sensitivity"):
        raise public_error(ErrorCode.VALIDATION_FAILED, "Phase 3A accepts baseline execution only; optional analysis flags must be false.")
    root = Path(project_root).resolve() if project_root else Path(__file__).resolve().parents[3]
    cells = _build_cells(payload, root)
    canonical = dict(payload)
    canonical.setdefault("evidence", [])
    canonical["requested_model_bundle_id"] = None
    canonical["governed_overrides"] = []
    canonical["compatibility"] = {
        "workbook_cells": cells,
        "inactive_legacy": _inactive(payload),
        "source_format": "combined-xlsx",
        "ownership_rule": "Server-generated allowlisted compatibility cells only; caller-supplied workbook coordinates are forbidden.",
    }
    try:
        return CRQAssessment.from_dict(canonical)
    except (KeyError, TypeError, ValueError) as exc:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "The assessment failed canonical validation.", reason=str(exc)) from None


def _reject_unknown(value: Mapping[str, Any], allowed: set[str], location: str) -> None:
    unknown = sorted(set(value) - allowed)
    if unknown:
        raise public_error(ErrorCode.INVALID_REQUEST, "Unknown fields are not permitted.", location=location, fields=unknown)


def _reject_banned(value: Any, path: str = "assessment") -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if key in _BANNED_KEYS:
                raise public_error(ErrorCode.INVALID_REQUEST, "Caller-supplied governed assumptions or implementation paths are forbidden.", location=f"{path}.{key}")
            _reject_banned(item, f"{path}.{key}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            _reject_banned(item, f"{path}[{index}]")


def _validate_public_shapes(payload: Mapping[str, Any]) -> None:
    identity = _object(payload, "assessment")
    required_identity = {"assessment_id", "assessment_version", "domain", "sector", "asset_type", "reporting_basis"}
    _require(identity, required_identity, "assessment.assessment")
    for key in ("assessment_id", "asset_type"):
        if not isinstance(identity.get(key), str) or not identity[key].strip():
            raise public_error(ErrorCode.INVALID_ASSESSMENT, f"assessment.assessment.{key} must be a non-empty string.")
    if type(identity.get("assessment_version")) is not int or identity["assessment_version"] < 1:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "assessment.assessment.assessment_version must be a positive integer.")
    if identity.get("reporting_basis") not in {"Best Estimate", "Prudent", "Both"}:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "assessment.assessment.reporting_basis is invalid.")
    _reject_unknown(_object(payload, "scope"), {"country_region", "facility", "critical_process"}, "assessment.scope")
    runtime = _object(payload, "runtime")
    _require(runtime, {"simulation_count", "random_seed", "run_whatifs", "run_packages", "run_sensitivity"}, "assessment.runtime")
    if type(runtime.get("simulation_count")) is not int or type(runtime.get("random_seed")) is not int:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Assessment runtime count and seed must be integers.")
    for key in ("run_whatifs", "run_packages", "run_sensitivity"):
        if type(runtime.get(key)) is not bool:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, f"assessment.runtime.{key} must be boolean.")
    appetite = _object(payload, "risk_appetite")
    _reject_unknown(
        appetite,
        {"annual_loss_tolerance", "max_acceptable_event_probability", "max_acceptable_tvar95", "max_acceptable_tvar99", "max_acceptable_downtime_days"},
        "assessment.risk_appetite",
    )
    for key, value in appetite.items():
        _nonnegative_or_none(value, f"assessment.risk_appetite.{key}")
    probability = appetite.get("max_acceptable_event_probability")
    if probability is not None and probability > 1:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Maximum acceptable event probability must not exceed 1.")
    insurance = _object(payload, "insurance")
    _reject_unknown(insurance, {"retention", "primary_limit", "aggregate_programme_limit", "layers"}, "assessment.insurance")
    _require(insurance, {"retention", "layers"}, "assessment.insurance")
    for key in ("retention", "primary_limit", "aggregate_programme_limit"):
        _nonnegative_or_none(insurance.get(key), f"assessment.insurance.{key}")
    if not isinstance(insurance.get("layers"), list) or len(insurance["layers"]) > 4:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Insurance layers must be an array of at most four entries.")
    for index, layer in enumerate(insurance.get("layers") or []):
        row = _mapping(layer, f"assessment.insurance.layers[{index}]")
        _reject_unknown(row, {"name", "attachment", "limit", "coinsurance"}, f"assessment.insurance.layers[{index}]")
        _require(row, {"name", "attachment", "limit", "coinsurance"}, f"assessment.insurance.layers[{index}]")
        for key in ("attachment", "limit", "coinsurance"):
            _nonnegative_or_none(row.get(key), f"assessment.insurance.layers[{index}].{key}", allow_none=False)
        if row["coinsurance"] > 1:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Insurance coinsurance must not exceed 1.")
    outside = _object(payload, "outside_in")
    _reject_unknown(outside, {"apply", "evidence_set_id", "source_file_hash", "approved_adjustments"}, "assessment.outside_in")
    evidence_fields = {"evidence_id", "description", "source", "observed_at", "hash"}
    for index, evidence in enumerate(payload.get("evidence") or []):
        _reject_unknown(_mapping(evidence, f"assessment.evidence[{index}]"), evidence_fields, f"assessment.evidence[{index}]")


def _mapping(value: Any, location: str) -> dict[str, Any]:
    if not isinstance(value, Mapping):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, f"{location} must be a JSON object.")
    return dict(value)


def _require(value: Mapping[str, Any], fields: set[str], location: str) -> None:
    missing = sorted(fields - set(value))
    if missing:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Required fields are missing.", location=location, fields=missing)


def _nonnegative_or_none(value: Any, location: str, *, allow_none: bool = True) -> None:
    if value is None and allow_none:
        return
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value < 0:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, f"{location} must be a non-negative number.")


def _object(payload: Mapping[str, Any], key: str) -> dict[str, Any]:
    value = payload.get(key)
    if not isinstance(value, dict):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, f"{key} must be a JSON object.")
    return value


def _record(sheet: str, cell: str, value: Any, classification: str, path: str) -> dict[str, Any]:
    return {"sheet": sheet, "cell": cell, "value": json_value(value), "classification": classification, "canonical_path": path}


def _build_cells(payload: dict[str, Any], root: Path) -> list[dict[str, Any]]:
    identity = payload["assessment"]
    scope = payload.get("scope") or {}
    runtime = payload["runtime"]
    inputs = payload["domain_inputs"]
    records = [
        _record("00 COMMON - Run Setup", "C6", identity["domain"], "USER_INPUT", "assessment.domain"),
        _record("00 COMMON - Run Setup", "C7", identity["sector"], "USER_INPUT", "assessment.sector"),
        _record("00 COMMON - Run Setup", "C8", identity["asset_type"], "USER_INPUT", "assessment.asset_type"),
        _record("00 COMMON - Run Setup", "C9", identity["reporting_basis"], "USER_INPUT", "assessment.reporting_basis"),
        _record("00 COMMON - Run Setup", "C10", "No", "USER_INPUT", "outside_in.apply"),
    ]
    appetite = payload.get("risk_appetite") or {}
    for row, key in enumerate(("annual_loss_tolerance", "max_acceptable_event_probability", "max_acceptable_tvar95", "max_acceptable_tvar99", "max_acceptable_downtime_days"), 6):
        records.append(_record("06 Appetite & Insurance", f"B{row}", appetite.get(key), "USER_INPUT", f"risk_appetite.{key}"))
    insurance = payload.get("insurance") or {}
    for cell, key in (("B13", "retention"), ("B14", "primary_limit"), ("B15", "aggregate_programme_limit")):
        records.append(_record("06 Appetite & Insurance", cell, insurance.get(key), "USER_INPUT", f"insurance.{key}"))
    for row, layer in enumerate(insurance.get("layers") or [], 18):
        if row > 21:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "At most four insurance layers are supported.")
        for col, key in zip("ABCD", ("name", "attachment", "limit", "coinsurance"), strict=True):
            records.append(_record("06 Appetite & Insurance", f"{col}{row}", layer.get(key), "USER_INPUT", f"insurance.layers.{key}"))
    template = root / "model" / "Guided_IT_OT_CRQ_Model_v1_0.xlsx"
    # Normal mode indexes cells in memory. Read-only worksheets make the
    # allowlist lookups below repeatedly stream the workbook and add seconds
    # to IT request validation.
    wb = openpyxl.load_workbook(template, read_only=False, data_only=False)
    try:
        if identity["domain"] == "IT":
            records.extend(_it_cells(identity, scope, runtime, inputs, wb))
        else:
            records.extend(_ot_cells(identity, scope, runtime, inputs, wb))
    finally:
        wb.close()
    return records


def _row_index(ws, id_column: int, start: int, end: int) -> dict[str, int]:
    return {str(ws.cell(r, id_column).value).strip(): r for r in range(start, end + 1) if ws.cell(r, id_column).value not in (None, "")}


def _it_cells(identity, scope, runtime, inputs, wb) -> list[dict[str, Any]]:
    _reject_unknown(inputs, {"domain", "financial_exposure", "exposure_model_enabled", "routes", "controls", "impact_overrides", "frequency_adjustments"}, "assessment.domain_inputs")
    _require(inputs, {"domain", "financial_exposure", "exposure_model_enabled", "routes", "controls", "impact_overrides", "frequency_adjustments"}, "assessment.domain_inputs")
    if inputs.get("domain") != "IT":
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "domain_inputs.domain must match assessment.domain.")
    fin = _mapping(inputs["financial_exposure"], "assessment.domain_inputs.financial_exposure")
    financial_fields = {"annual_revenue_at_risk", "annual_payment_value", "bi_loss_factor", "employees", "customers", "sensitive_records", "critical_endpoints", "critical_servers"}
    _reject_unknown(fin, financial_fields, "assessment.domain_inputs.financial_exposure")
    _require(fin, financial_fields - {"employees", "customers"}, "assessment.domain_inputs.financial_exposure")
    for key, value in fin.items():
        _nonnegative_or_none(value, f"assessment.domain_inputs.financial_exposure.{key}")
    if type(inputs.get("exposure_model_enabled")) is not bool:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "exposure_model_enabled must be boolean.")
    mapping = {
        "C16": identity["assessment_id"], "C17": identity.get("organisation"), "C18": scope.get("country_region"), "C20": identity.get("currency"),
        "C21": fin["annual_revenue_at_risk"], "C22": fin["annual_payment_value"], "C23": fin["bi_loss_factor"],
        "C24": fin.get("employees"), "C25": fin.get("customers"), "C26": fin["sensitive_records"], "C27": fin["critical_endpoints"], "C28": fin["critical_servers"],
        "C29": "Yes" if inputs["exposure_model_enabled"] else "No", "C30": runtime["simulation_count"], "C31": runtime["random_seed"],
    }
    records = [_record("IT 03 - Organisation Inputs", cell, value, "INACTIVE_LEGACY" if cell in {"C24", "C25"} else "USER_INPUT", "domain_inputs.financial_exposure") for cell, value in mapping.items()]
    route_rows = _row_index(wb["IT 04 - Exposure Adjustments"], 1, 16, 115)
    for route in inputs.get("routes") or []:
        route = _mapping(route, "assessment.domain_inputs.routes[]")
        route_fields = {"route_id", "organisation_applicable", "model_feasible", "opportunity_factor", "stage_factors", "impact_factor", "confidence", "recommendation_id", "evidence_date", "rationale"}
        _reject_unknown(route, route_fields, "assessment.domain_inputs.routes[]")
        _require(route, {"route_id", "organisation_applicable", "model_feasible", "opportunity_factor", "stage_factors", "impact_factor"}, "assessment.domain_inputs.routes[]")
        rid = route.get("route_id")
        if rid not in route_rows:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Unknown IT route identifier.", route_id=rid)
        row = route_rows[rid]
        if route["organisation_applicable"] not in {"Yes", "No", "Unknown"} or route["model_feasible"] not in {"Yes", "No", "Unknown"}:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "IT route state must be Yes, No, or Unknown.", route_id=rid)
        stage_factors = _mapping(route["stage_factors"], f"assessment.domain_inputs.routes.{rid}.stage_factors")
        _reject_unknown(stage_factors, {"S1", "S2", "S3", "S4", "S5"}, f"assessment.domain_inputs.routes.{rid}.stage_factors")
        _require(stage_factors, {"S1", "S2", "S3", "S4", "S5"}, f"assessment.domain_inputs.routes.{rid}.stage_factors")
        for key, value in {"opportunity_factor": route["opportunity_factor"], "impact_factor": route["impact_factor"], **stage_factors}.items():
            _nonnegative_or_none(value, f"assessment.domain_inputs.routes.{rid}.{key}", allow_none=False)
        values = {"C": route["organisation_applicable"], "D": route["model_feasible"], "E": route["opportunity_factor"], "K": route["impact_factor"], "L": route.get("confidence"), "M": route.get("recommendation_id"), "N": route.get("evidence_date"), "O": route.get("rationale")}
        values.update({chr(ord("F") + index): stage_factors[f"S{index + 1}"] for index in range(5)})
        records.extend(_record("IT 04 - Exposure Adjustments", f"{col}{row}", value, "EVIDENCE_ONLY" if col in "LMNO" else "PERMITTED_OVERRIDE", f"domain_inputs.routes.{rid}") for col, value in values.items())
    control_rows = _row_index(wb["IT 05 - Control Assessment"], 1, 16, 115)
    for control in inputs.get("controls") or []:
        control = _validate_control(control, "IT")
        cid = control.get("control_id")
        if cid not in control_rows:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Unknown IT control identifier.", control_id=cid)
        row = control_rows[cid]
        records.extend((_record("IT 05 - Control Assessment", f"E{row}", control["maturity"], "USER_INPUT", f"domain_inputs.controls.{cid}.maturity"), _record("IT 05 - Control Assessment", f"F{row}", control["coverage"], "USER_INPUT", f"domain_inputs.controls.{cid}.coverage")))
    impact_rows = {(str(wb["IT 06 - Impact & BIA Overrides"].cell(r, 1).value or ""), str(wb["IT 06 - Impact & BIA Overrides"].cell(r, 2).value or ""), str(wb["IT 06 - Impact & BIA Overrides"].cell(r, 4).value or "")): r for r in range(16, 251)}
    for override in inputs.get("impact_overrides") or []:
        override = _mapping(override, "assessment.domain_inputs.impact_overrides[]")
        _reject_unknown(override, {"scenario_id", "parameter_id", "p50", "p99", "evidence_id"}, "assessment.domain_inputs.impact_overrides[]")
        _require(override, {"parameter_id"}, "assessment.domain_inputs.impact_overrides[]")
        for key in ("p50", "p99"):
            _nonnegative_or_none(override.get(key), f"assessment.domain_inputs.impact_overrides.{key}")
        key = ("PARAM" if override.get("scenario_id") else "RATE", str(override.get("scenario_id") or ""), str(override["parameter_id"]))
        if key not in impact_rows:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Unknown IT impact override.", parameter_id=override["parameter_id"])
        row = impact_rows[key]
        records.extend((_record("IT 06 - Impact & BIA Overrides", f"I{row}", override.get("p50"), "PERMITTED_OVERRIDE", "domain_inputs.impact_overrides"), _record("IT 06 - Impact & BIA Overrides", f"J{row}", override.get("p99"), "PERMITTED_OVERRIDE", "domain_inputs.impact_overrides")))
    freq_rows = _row_index(wb["IT 07 - Assessment Adjustments"], 1, 16, 22)
    for key, value in (inputs.get("frequency_adjustments") or {}).items():
        _nonnegative_or_none(value, f"assessment.domain_inputs.frequency_adjustments.{key}", allow_none=False)
        if key not in freq_rows:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Unknown IT frequency adjustment.", parameter=key)
        records.append(_record("IT 07 - Assessment Adjustments", f"B{freq_rows[key]}", value, "PERMITTED_OVERRIDE", f"domain_inputs.frequency_adjustments.{key}"))
    records.append(_record("IT 08 - Reporting Settings", "B16", identity["reporting_basis"], "USER_INPUT", "assessment.reporting_basis"))
    return records


def _ot_cells(identity, scope, runtime, inputs, wb) -> list[dict[str, Any]]:
    _reject_unknown(inputs, {"domain", "facility", "topology", "financial_exposure", "controls", "impact_drivers", "assessment_adjustments"}, "assessment.domain_inputs")
    _require(inputs, {"domain", "facility", "topology", "financial_exposure", "controls", "impact_drivers", "assessment_adjustments"}, "assessment.domain_inputs")
    if inputs.get("domain") != "OT" or inputs.get("impact_drivers") not in (None, []):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "OT governed impact drivers cannot be supplied by the caller.")
    facility = _mapping(inputs.get("facility"), "assessment.domain_inputs.facility")
    _reject_unknown(facility, {"facility_name", "critical_process"}, "assessment.domain_inputs.facility")
    topology = _mapping(inputs.get("topology"), "assessment.domain_inputs.topology")
    topology_fields = {"it_ot_connectivity", "remote_access", "internet_ot", "transient_assets", "supply_chain_route", "wireless_ot", "safety_system"}
    _reject_unknown(topology, topology_fields, "assessment.domain_inputs.topology")
    values = {
        "C16": identity["assessment_id"], "C17": identity.get("organisation"), "C18": facility.get("facility_name") or scope.get("facility"), "C19": scope.get("country_region"), "C22": facility.get("critical_process") or scope.get("critical_process"),
        "C23": topology.get("it_ot_connectivity"), "C24": topology.get("remote_access"), "C25": topology.get("internet_ot"), "C26": topology.get("transient_assets"), "C27": topology.get("supply_chain_route"), "C28": topology.get("wireless_ot"), "C29": topology.get("safety_system"),
    }
    records = [_record("OT 03 - Facility Inputs", cell, value, "USER_INPUT", "assessment/scope/domain_inputs.topology") for cell, value in values.items()]
    control_rows = _row_index(wb["OT 04 - Control Assessment"], 1, 16, 115)
    for control in inputs.get("controls") or []:
        control = _validate_control(control, "OT")
        cid = control.get("control_id")
        if cid not in control_rows:
            raise public_error(ErrorCode.INVALID_ASSESSMENT, "Unknown OT control identifier.", control_id=cid)
        row = control_rows[cid]
        for col, key in (("D", "maturity"), ("E", "coverage"), ("F", "tested"), ("G", "test_result")):
            value = control.get(key)
            if key == "tested":
                value = "Yes" if value else "No"
            records.append(_record("OT 04 - Control Assessment", f"{col}{row}", value, "EVIDENCE_ONLY" if col in "FG" else "USER_INPUT", f"domain_inputs.controls.{cid}.{key}"))
    fin = _mapping(inputs["financial_exposure"], "assessment.domain_inputs.financial_exposure")
    financial_fields = {"annual_revenue_at_risk", "bi_loss_factor", "employees", "critical_ot_endpoints", "critical_ot_servers"}
    _reject_unknown(fin, financial_fields, "assessment.domain_inputs.financial_exposure")
    _require(fin, financial_fields - {"employees"}, "assessment.domain_inputs.financial_exposure")
    for key, value in fin.items():
        _nonnegative_or_none(value, f"assessment.domain_inputs.financial_exposure.{key}")
    for cell, key in (("C16", "annual_revenue_at_risk"), ("C17", "bi_loss_factor"), ("C18", "employees"), ("C19", "critical_ot_endpoints"), ("C20", "critical_ot_servers")):
        records.append(_record("OT 05 - Impact & BIA", cell, fin.get(key), "INACTIVE_LEGACY" if key == "employees" else "USER_INPUT", f"domain_inputs.financial_exposure.{key}"))
    adjustments = _mapping(inputs.get("assessment_adjustments"), "assessment.domain_inputs.assessment_adjustments")
    _reject_unknown(adjustments, {"prudence_factor"}, "assessment.domain_inputs.assessment_adjustments")
    prudence = adjustments.get("prudence_factor")
    _nonnegative_or_none(prudence, "assessment.domain_inputs.assessment_adjustments.prudence_factor")
    records.extend((_record("OT 06 - Assessment Adjustments", "D17", prudence, "PERMITTED_OVERRIDE", "domain_inputs.assessment_adjustments.prudence_factor"), _record("OT 06 - Assessment Adjustments", "D18", identity["reporting_basis"], "USER_INPUT", "assessment.reporting_basis"), _record("OT 06 - Assessment Adjustments", "D19", runtime["simulation_count"], "USER_INPUT", "runtime.simulation_count"), _record("OT 06 - Assessment Adjustments", "D20", runtime["random_seed"], "USER_INPUT", "runtime.random_seed")))
    return records


def _inactive(payload: dict[str, Any]) -> dict[str, Any]:
    fin = payload["domain_inputs"]["financial_exposure"]
    result = {"employees": fin.get("employees")}
    if payload["assessment"]["domain"] == "IT":
        result["customers"] = fin.get("customers")
    return result


def _validate_control(value: Any, domain: str) -> dict[str, Any]:
    control = _mapping(value, f"assessment.domain_inputs.controls[{domain}]")
    fields = {"control_id", "maturity", "coverage", "tested", "test_result", "evidence_ids"}
    _reject_unknown(control, fields, "assessment.domain_inputs.controls[]")
    _require(control, {"control_id", "maturity", "coverage"}, "assessment.domain_inputs.controls[]")
    if not isinstance(control["control_id"], str) or not isinstance(control["maturity"], str):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Control identifiers and maturity values must be strings.")
    _nonnegative_or_none(control["coverage"], "assessment.domain_inputs.controls[].coverage", allow_none=False)
    if control["coverage"] > 1:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Control coverage must not exceed 1.")
    if control.get("tested") is not None and type(control["tested"]) is not bool:
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Control tested state must be boolean or null.")
    if not isinstance(control.get("evidence_ids", []), list):
        raise public_error(ErrorCode.INVALID_ASSESSMENT, "Control evidence_ids must be an array.")
    return control
