"""Read-only Excel-to-CRQAssessment adapter and temporary reverse adapter."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.cell.cell import MergedCell

from crq.application.models import CRQAssessment, json_value
from crq.sheet_names import COMMON_RUN_SETUP
from it_ot_crq.reporting.inputs import read_appetite_inputs, read_insurance_programme


def _kv(ws, start: int, end: int, key_col: int = 1, value_col: int = 3) -> dict[str, Any]:
    return {
        str(ws.cell(r, key_col).value).strip(): json_value(ws.cell(r, value_col).value)
        for r in range(start, end + 1)
        if ws.cell(r, key_col).value not in (None, "")
    }


def _cell_records(ws, cells: Iterable[str], classification: str, canonical_path: str) -> list[dict[str, Any]]:
    records = []
    for coordinate in cells:
        cell = ws[coordinate]
        if isinstance(cell, MergedCell):
            continue
        value = cell.value
        if isinstance(value, str) and value.startswith("="):
            continue
        records.append({
            "sheet": ws.title,
            "cell": coordinate,
            "value": json_value(value),
            "classification": classification,
            "canonical_path": canonical_path,
        })
    return records


def _coords(columns: str, first: int, last: int) -> list[str]:
    return [f"{col}{row}" for row in range(first, last + 1) for col in columns]


def extract_assessment(workbook: str | Path, *, run_whatifs: bool = False) -> CRQAssessment:
    """Extract client-owned inputs without retaining a workbook path or governed pack values."""
    source = Path(workbook).expanduser().resolve()
    if not source.is_file():
        raise FileNotFoundError(source)
    wb = openpyxl.load_workbook(source, data_only=False)
    try:
        setup = wb[COMMON_RUN_SETUP]
        domain = str(setup["C6"].value or "").strip()
        sector = str(setup["C7"].value or "").strip()
        if domain not in {"IT", "OT"}:
            raise ValueError("Run Setup domain must be IT or OT.")
        asset = str(setup["C8"].value or "").strip()
        basis = str(setup["C9"].value or "Prudent").strip()
        outside = str(setup["C10"].value or "No").strip()
        appetite_raw = read_appetite_inputs(wb)
        insurance_raw = read_insurance_programme(wb)
        records = _cell_records(setup, _coords("C", 6, 10), "USER_INPUT", "assessment/outside_in")
        records += _cell_records(setup, ["C11"], "INACTIVE_LEGACY", "outside_in.legacy_file_path")
        ai = wb["06 Appetite & Insurance"]
        records += _cell_records(ai, _coords("B", 6, 15), "USER_INPUT", "risk_appetite/insurance")
        records += _cell_records(ai, _coords("ABCD", 18, 21), "USER_INPUT", "insurance.layers")

        if domain == "IT":
            data, domain_records, evidence, inactive = _extract_it(wb)
        else:
            data, domain_records, evidence, inactive = _extract_ot(wb)
        records.extend(domain_records)
        runtime = data.pop("runtime")
        payload = {
            "schema_version": "1.0.0-draft",
            "assessment": {
                "assessment_id": data.pop("assessment_id"),
                "assessment_version": 1,
                "organisation": data.pop("organisation", None),
                "domain": domain,
                "sector": sector,
                "asset_type": asset,
                "reporting_basis": basis,
                "currency": data.pop("currency", None),
            },
            "scope": data.pop("scope"),
            "runtime": {
                **runtime,
                "run_whatifs": bool(run_whatifs),
                "run_packages": bool(run_whatifs),
                "run_sensitivity": False,
                "legacy_simulation_mode": "unchanged engine-native paired/common-random-number semantics",
            },
            "risk_appetite": {
                "annual_loss_tolerance": appetite_raw.get("ANNUAL_LOSS_TOLERANCE"),
                "max_acceptable_event_probability": appetite_raw.get("MAX_ACCEPTABLE_EVENT_PROBABILITY"),
                "max_acceptable_tvar95": appetite_raw.get("MAX_ACCEPTABLE_TVAR_95"),
                "max_acceptable_tvar99": appetite_raw.get("MAX_ACCEPTABLE_TVAR_99"),
                "max_acceptable_downtime_days": appetite_raw.get("MAX_ACCEPTABLE_DOWNTIME_DAYS"),
            },
            "insurance": {
                "retention": float(insurance_raw.get("INSURANCE_RETENTION") or 0),
                "primary_limit": insurance_raw.get("PRIMARY_LIMIT"),
                "aggregate_programme_limit": insurance_raw.get("AGGREGATE_PROGRAMME_LIMIT"),
                "layers": json_value(insurance_raw.get("layers") or []),
            },
            "outside_in": {"apply": outside == "Yes", "evidence_set_id": None, "source_file_hash": None, "approved_adjustments": []},
            "domain_inputs": data,
            "requested_model_bundle_id": None,
            "governed_overrides": [],
            "evidence": evidence,
            "compatibility": {
                "workbook_cells": records,
                "inactive_legacy": inactive,
                "source_format": "combined-xlsx",
                "ownership_rule": "Only USER_INPUT, PERMITTED_OVERRIDE, EVIDENCE_ONLY and INACTIVE_LEGACY cells are carried; governed pack inputs are excluded.",
            },
        }
        return CRQAssessment.from_dict(payload)
    finally:
        wb.close()


def _extract_it(wb) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    org_ws = wb["IT 03 - Organisation Inputs"]
    org = _kv(org_ws, 16, 31)
    routes_ws = wb["IT 04 - Exposure Adjustments"]
    routes, evidence = [], []
    for r in range(16, 22):
        rid = str(routes_ws.cell(r, 1).value or "").strip()
        routes.append({"route_id": rid, "organisation_applicable": routes_ws.cell(r, 3).value, "model_feasible": routes_ws.cell(r, 4).value, "opportunity_factor": routes_ws.cell(r, 5).value, "stage_factors": {f"S{i}": routes_ws.cell(r, 5 + i).value for i in range(1, 6)}, "impact_factor": routes_ws.cell(r, 11).value, "confidence": routes_ws.cell(r, 12).value, "recommendation_id": routes_ws.cell(r, 13).value, "evidence_date": json_value(routes_ws.cell(r, 14).value), "rationale": routes_ws.cell(r, 15).value})
        if any(routes_ws.cell(r, c).value not in (None, "") for c in range(12, 16)):
            evidence.append({"evidence_id": f"IT-ROUTE-{rid}", "description": str(routes_ws.cell(r, 15).value or "Route evidence"), "source": str(routes_ws.cell(r, 13).value or "workbook"), "observed_at": None, "hash": None})
    ctl_ws = wb["IT 05 - Control Assessment"]
    controls = []
    for r in range(16, 116):
        cid = str(ctl_ws.cell(r, 1).value or "").strip()
        if cid:
            controls.append({"control_id": cid, "maturity": ctl_ws.cell(r, 5).value, "coverage": ctl_ws.cell(r, 6).value, "tested": None, "test_result": None, "evidence_ids": [f"IT-CONTROL-{cid}"] if ctl_ws.cell(r, 9).value else []})
    impact_ws = wb["IT 06 - Impact & BIA Overrides"]
    overrides = []
    for r in range(16, 251):
        kind, sid, pid = (impact_ws.cell(r, c).value for c in (1, 2, 4))
        if kind in {"RATE", "PARAM"} and pid:
            overrides.append({"scenario_id": sid, "parameter_id": str(pid), "p50": impact_ws.cell(r, 9).value, "p99": impact_ws.cell(r, 10).value, "evidence_id": None})
    freq = _kv(wb["IT 07 - Assessment Adjustments"], 16, 22, 1, 2)
    records = _cell_records(org_ws, _coords("C", 16, 23) + _coords("C", 26, 31), "USER_INPUT", "assessment/domain_inputs/runtime")
    records += _cell_records(org_ws, _coords("C", 24, 25), "INACTIVE_LEGACY", "domain_inputs.financial_exposure")
    records += _cell_records(routes_ws, _coords("CDEFGHIJK", 16, 21), "PERMITTED_OVERRIDE", "domain_inputs.routes")
    records += _cell_records(routes_ws, _coords("LMNO", 16, 21), "EVIDENCE_ONLY", "domain_inputs.routes.evidence")
    records += _cell_records(ctl_ws, _coords("EF", 16, 115), "USER_INPUT", "domain_inputs.controls")
    records += _cell_records(ctl_ws, _coords("I", 16, 115), "EVIDENCE_ONLY", "domain_inputs.controls.evidence")
    records += _cell_records(impact_ws, _coords("IJ", 16, 250), "PERMITTED_OVERRIDE", "domain_inputs.impact_overrides")
    records += _cell_records(wb["IT 07 - Assessment Adjustments"], _coords("B", 16, 22), "PERMITTED_OVERRIDE", "domain_inputs.frequency_adjustments")
    records += _cell_records(wb["IT 08 - Reporting Settings"], _coords("B", 16, 19), "USER_INPUT", "runtime/reporting")
    return ({
        "domain": "IT",
        "assessment_id": str(org.get("ASSESSMENT_ID") or "IT-assessment"),
        "organisation": org.get("ORGANISATION"),
        "currency": org.get("CURRENCY"),
        "scope": {"country_region": org.get("COUNTRY_REGION"), "facility": None, "critical_process": None},
        "runtime": {"simulation_count": int(org.get("SIMULATION_YEARS") or 500000), "random_seed": int(org.get("RANDOM_SEED") or 20260821)},
        "financial_exposure": {"annual_revenue_at_risk": org.get("ANNUAL_REVENUE_AT_RISK"), "annual_payment_value": org.get("ANNUAL_PAYMENT_VALUE"), "bi_loss_factor": org.get("BI_LOSS_FACTOR"), "employees": org.get("EMPLOYEES"), "customers": org.get("CUSTOMERS"), "sensitive_records": org.get("SENSITIVE_RECORDS"), "critical_endpoints": org.get("CRITICAL_ENDPOINTS"), "critical_servers": org.get("CRITICAL_SERVERS")},
        "exposure_model_enabled": org.get("USE_EXPOSURE_MODEL") == "Yes",
        "routes": routes, "controls": controls, "impact_overrides": overrides, "frequency_adjustments": freq,
    }, records, evidence, {"employees": org.get("EMPLOYEES"), "customers": org.get("CUSTOMERS")})


def _extract_ot(wb) -> tuple[dict[str, Any], list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    fac_ws = wb["OT 03 - Facility Inputs"]
    fac = _kv(fac_ws, 16, 29)
    ctl_ws = wb["OT 04 - Control Assessment"]
    controls, evidence = [], []
    for r in range(16, 116):
        cid = str(ctl_ws.cell(r, 1).value or "").strip()
        if cid:
            ev_id = f"OT-CONTROL-{cid}"
            controls.append({"control_id": cid, "maturity": ctl_ws.cell(r, 4).value, "coverage": ctl_ws.cell(r, 5).value or 0.75, "tested": str(ctl_ws.cell(r, 6).value or "").lower() == "yes", "test_result": ctl_ws.cell(r, 7).value, "evidence_ids": [ev_id] if ctl_ws.cell(r, 8).value else []})
            if any(ctl_ws.cell(r, c).value not in (None, "") for c in (6, 7, 8)):
                evidence.append({"evidence_id": ev_id, "description": str(ctl_ws.cell(r, 7).value or "Control evidence"), "source": str(ctl_ws.cell(r, 8).value or "workbook"), "observed_at": None, "hash": None})
    fin = _kv(wb["OT 05 - Impact & BIA"], 16, 20)
    adj = _kv(wb["OT 06 - Assessment Adjustments"], 16, 20, 1, 4)
    records = _cell_records(fac_ws, _coords("C", 16, 29), "USER_INPUT", "assessment/scope/domain_inputs.topology")
    records += _cell_records(ctl_ws, _coords("DE", 16, 115), "USER_INPUT", "domain_inputs.controls")
    records += _cell_records(ctl_ws, _coords("FGH", 16, 115), "EVIDENCE_ONLY", "domain_inputs.controls.evidence")
    impact_ws = wb["OT 05 - Impact & BIA"]
    records += _cell_records(impact_ws, ["C16", "C17", "C19", "C20"], "USER_INPUT", "domain_inputs.financial_exposure")
    records += _cell_records(impact_ws, ["C18"], "INACTIVE_LEGACY", "domain_inputs.financial_exposure.employees")
    records += _cell_records(wb["OT 06 - Assessment Adjustments"], ["D17"], "PERMITTED_OVERRIDE", "domain_inputs.assessment_adjustments.prudence_factor")
    records += _cell_records(wb["OT 06 - Assessment Adjustments"], ["D18", "D19", "D20"], "USER_INPUT", "runtime/reporting")
    return ({
        "domain": "OT",
        "assessment_id": str(fac.get("ASSESSMENT_ID") or "OT-assessment"),
        "organisation": fac.get("ORGANISATION"),
        "currency": None,
        "scope": {"country_region": fac.get("COUNTRY_REGION"), "facility": fac.get("FACILITY"), "critical_process": fac.get("CRITICAL_PROCESS")},
        "runtime": {"simulation_count": int(adj.get("SIMULATIONS") or 500000), "random_seed": int(adj.get("RANDOM_SEED") or 20260813)},
        "facility": {"facility_name": fac.get("FACILITY"), "critical_process": fac.get("CRITICAL_PROCESS")},
        "topology": {k.lower(): fac.get(k) for k in ("IT_OT_CONNECTIVITY", "REMOTE_ACCESS", "INTERNET_OT", "TRANSIENT_ASSETS", "SUPPLY_CHAIN_ROUTE", "WIRELESS_OT", "SAFETY_SYSTEM")},
        "financial_exposure": {"annual_revenue_at_risk": fin.get("ANNUAL_REVENUE_AT_RISK"), "bi_loss_factor": fin.get("BI_LOSS_FACTOR"), "employees": fin.get("EMPLOYEES"), "critical_ot_endpoints": fin.get("CRITICAL_OT_ENDPOINTS"), "critical_ot_servers": fin.get("CRITICAL_OT_SERVERS")},
        "controls": controls,
        "impact_drivers": [],
        "assessment_adjustments": {"prudence_factor": adj.get("PRUDENCE_FACTOR")},
    }, records, evidence, {"employees": fin.get("EMPLOYEES")})


def materialize_compatibility_workbook(assessment: CRQAssessment, template: Path, destination: Path) -> None:
    """Create a temporary engine workbook from structured client-owned cells."""
    shutil.copy2(template, destination)
    wb = openpyxl.load_workbook(destination, data_only=False)
    try:
        for record in assessment.to_dict()["compatibility"]["workbook_cells"]:
            if record["classification"] in {"USER_INPUT", "PERMITTED_OVERRIDE", "EVIDENCE_ONLY", "INACTIVE_LEGACY"}:
                wb[record["sheet"]][record["cell"]] = record["value"]
        wb.save(destination)
    finally:
        wb.close()
