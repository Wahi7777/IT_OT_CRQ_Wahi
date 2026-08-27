"""Single workbook styling source of truth.

Sheet ROLE (not engine track) selects title/header fill. Font is Arial everywhere.
"""

from __future__ import annotations

from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

FONT_NAME = "Arial"

SIZE_TITLE = 14
SIZE_SECTION = 10
SIZE_BODY = 9

COLOR_WHITE = "FFFFFF"
COLOR_INPUT = "0000FF"
COLOR_FORMULA = "000000"
COLOR_LINK = "008000"

# Role → title / column-header fill.
ROLE_DASHBOARD = "dashboard"
ROLE_NAV = "navigation"
ROLE_INPUT = "user_input"
ROLE_CALC = "engine_calc"
ROLE_LOGIC = "mapping_logic"
ROLE_PACK = "sector_pack"
ROLE_AUDIT = "validation_audit"

HEADER_FILL = {
    ROLE_DASHBOARD: "17365D",
    ROLE_NAV: "17365D",
    ROLE_INPUT: "2F75B5",
    ROLE_CALC: "0F6B78",
    ROLE_LOGIC: "7030A0",
    ROLE_PACK: "D99400",
    ROLE_AUDIT: "736A23",
}

FMT_CURRENCY = r'\$#,##0;[RED]("$"#,##0);\-'
FMT_PERCENT = r"0.0%;[RED]\(0.0%\);\-"
FMT_MULTIPLE = "0.00x"

INPUT_FILL = "FFF2CC"

_FREEZE_ROW_THRESHOLD = 40

SHEET_ROLE_BY_TITLE_PREFIX = (
    ("00 Dashboard", ROLE_DASHBOARD),
    ("00 Risk", ROLE_DASHBOARD),
    ("00 Loss", ROLE_DASHBOARD),
    ("00 What", ROLE_DASHBOARD),
    ("00 Risk Transfer", ROLE_DASHBOARD),
    ("1 Executive", ROLE_DASHBOARD),
    ("2 User", ROLE_NAV),
    ("00 COMMON - Run Setup", ROLE_NAV),
    ("00 COMMON - Outside-In", ROLE_INPUT),
    ("00 COMMON - Output", ROLE_DASHBOARD),
    ("00 COMMON - Engine", ROLE_AUDIT),
    ("3 Assumption", ROLE_NAV),
    ("IT 02", ROLE_DASHBOARD),
    ("OT 02", ROLE_DASHBOARD),
    ("IT 03", ROLE_INPUT),
    ("IT 04", ROLE_INPUT),
    ("IT 05", ROLE_INPUT),
    ("IT 06", ROLE_INPUT),
    ("IT 07", ROLE_INPUT),
    ("IT 08", ROLE_INPUT),
    ("OT 03", ROLE_INPUT),
    ("OT 04", ROLE_INPUT),
    ("OT 05", ROLE_INPUT),
    ("OT 06", ROLE_INPUT),
    ("IT CALC", ROLE_CALC),
    ("OT CALC", ROLE_CALC),
    ("IT CORE", ROLE_LOGIC),
    ("OT CORE", ROLE_LOGIC),
    ("IT 09", ROLE_PACK),
    ("OT PACK", ROLE_PACK),
    ("00 Pack", ROLE_PACK),
    ("01 Pack", ROLE_PACK),
    ("IT 22", ROLE_AUDIT),
    ("18 Validation", ROLE_AUDIT),
)


def role_for_sheet(title: str) -> str:
    for prefix, role in SHEET_ROLE_BY_TITLE_PREFIX:
        if title == prefix or title.startswith(prefix):
            return role
    if "CALC" in title:
        return ROLE_CALC
    if "PACK" in title or title[:2].isdigit() and "Pack" in title:
        return ROLE_PACK
    return ROLE_NAV


def fill_header(role: str) -> PatternFill:
    return PatternFill("solid", fgColor=HEADER_FILL[role])


def font_title() -> Font:
    return Font(name=FONT_NAME, bold=True, size=SIZE_TITLE, color=COLOR_WHITE)


def font_section(*, color: str = COLOR_WHITE) -> Font:
    return Font(name=FONT_NAME, bold=True, size=SIZE_SECTION, color=color)


def font_body(*, bold: bool = False, color: str = COLOR_FORMULA, italic: bool = False) -> Font:
    return Font(name=FONT_NAME, bold=bold, size=SIZE_BODY, color=color, italic=italic)


def font_input() -> Font:
    return font_body(color=COLOR_INPUT)


def font_link() -> Font:
    return font_body(color=COLOR_LINK)


def apply_title_banner(cell, role: str) -> None:
    cell.font = font_title()
    cell.fill = fill_header(role)


def apply_column_headers(ws: Worksheet, row: int, cols: int, role: str | None = None) -> None:
    """Paint an existing header row. Does not change cell values."""
    role = role or role_for_sheet(ws.title)
    fill = fill_header(role)
    font = font_section()
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.fill = fill
        cell.font = font


def apply_freeze_at_headers(ws: Worksheet, header_row: int, *, min_data_rows: int = _FREEZE_ROW_THRESHOLD) -> None:
    """Freeze so title/purpose and the header row stay visible (first unfrozen cell is the first data row)."""
    data_rows = (ws.max_row or 0) - header_row
    if data_rows < min_data_rows:
        return
    ws.freeze_panes = f"A{header_row + 1}"


def style_input_cell(cell) -> None:
    cell.font = font_input()
    cell.fill = PatternFill("solid", fgColor=INPUT_FILL)


def style_formula_cell(cell) -> None:
    cell.font = font_body(color=COLOR_FORMULA)


def style_link_cell(cell) -> None:
    cell.font = font_link()


LEGEND_INPUT = (
    "Colour convention: blue text = hardcoded / override input; black text = formula; "
    "green text = cross-sheet link."
)
