"""Purpose and model-use blurbs for every sector-pack worksheet (IT and OT)."""

from __future__ import annotations

from openpyxl.styles import Alignment

from crq.sheet_theme import font_body

# sheet title → (purpose, how used in the model)
IT_PACK_BLURBS = {
    "00 Pack Guide": (
        "Orientation for the pack owner. Lists what this Financial Services pack contains and what belongs in the assessment workbook instead.",
        "Not loaded as calibration. Python reads the numbered tables on sheets 01–18 after validating 01 Pack Metadata.",
    ),
    "01 Pack Metadata": (
        "Identity, schema, minimum engine version, status and unit of analysis for this pack.",
        "Validated fail-closed before any other pack table is used. Mismatched PACK_ID, SECTOR, DOMAIN or SCHEMA_VERSION stops the run.",
    ),
    "02 Actor Weights": (
        "Reference mix of material campaigns across the four IT threat actors for a broad Financial Services organisation.",
        "it_crq uses these weights as the actor campaign mix (they should sum to 1). Assessment overlays on IT 07 can scale activity but do not replace this mix unless an override is provided.",
    ),
    "03 Actor Scenario": (
        "For each actor, the conditional share of campaigns that lead to each of the four IT loss scenarios.",
        "Combined with actor weights and route propensities to build actor–scenario–route cells that the Monte Carlo simulates.",
    ),
    "04 Route Definitions": (
        "Stable identifiers and descriptions of the composite IT attack routes (identity, endpoint, exploit, supply chain, DDoS, insider).",
        "Route IDs key sheets 05–07 and 13. The engine does not invent routes that are not listed here.",
    ),
    "05 Actor Route": (
        "Reference distribution of each actor’s campaigns across composite routes.",
        "Multiplied with scenario-route propensity and normalised so each actor–scenario pair has a coherent route mix.",
    ),
    "06 Scenario Route": (
        "Reference distribution of each loss scenario across composite routes.",
        "Combined with actor-route propensity; the product is the route mix used for TTP and control calculations.",
    ),
    "07 Stage Requirements": (
        "Which of the five composite stages are required for each scenario–route pair (ALL rows, with optional actor overrides).",
        "Gates the attack path: a required stage with no viable TTP fails the path. Not an assessment input sheet.",
    ),
    "08 Stage Priors": (
        "Reference stage-through probabilities after the pack’s reference control profile.",
        "Used as the prior; assessed control maturity on IT 05 adjusts through-probabilities relative to this reference, not relative to zero controls.",
    ),
    "09 Scenario Parameters": (
        "P50/P99 scenario quantities (shares, days, direct amounts) with evidence grades.",
        "Scaled by organisation size from IT 03 to produce event-loss distributions. Blank assessment overrides on IT 06 keep these pack values.",
    ),
    "10 Driver Rates": (
        "Unit-cost and rate assumptions (notification, forensics, restoration, and so on) at P50/P99.",
        "Applied to the applicable loss blocks for each actor/scenario. Organisation overrides on IT 06 replace individual rates when populated.",
    ),
    "11 Impact Driver Matrix": (
        "Which loss blocks and drivers apply to each actor and scenario.",
        "Turns organisation scale and pack rates into the three IT loss blocks (duration, exposure, direct). Zeros here mean that driver is not modelled for that pair.",
    ),
    "12 TTP Catalogue": (
        "Enterprise ATT&CK techniques in scope for this pack, with actor and scenario applicability.",
        "Techniques that are not applicable are excluded from the route path. Weights and barriers still come from sheets 13 and 15.",
    ),
    "13 Route TTP Map": (
        "Which catalogue techniques belong to each composite route, with relevance weights.",
        "Builds the TTP set for each route. Combined with control efficacy (sheet 15) to compute stage barriers.",
    ),
    "14 Control Reference": (
        "Reference maturity and coverage for each control in the FS control set.",
        "The assessed organisation on IT 05 is measured relative to this profile so the model does not treat ‘no assessment’ as ‘no control’.",
    ),
    "15 Control TTP Map": (
        "How each control affects each TTP (prevent, contain, or reduce a loss block) and base efficacy.",
        "Converts IT 05 maturity into TTP barriers and duration/exposure/consequence factors used in frequency and severity.",
    ),
    "16 Sources": (
        "Bibliography and limitations for pack assumptions.",
        "Traceability only. The engine does not read frequencies or losses from this sheet.",
    ),
    "17 Evidence Register": (
        "The evidence-grade scale (A–D) applied to material pack assumptions.",
        "Governance labels only. Grades do not rescale Python results.",
    ),
    "18 Validation": (
        "Spreadsheet checks (IDs present, weights, row counts).",
        "Advisory in Excel. Python repeats identity, schema and structural checks and fails closed independently.",
    ),
}

OT_PACK_BLURBS = {
    "00 Pack Guide": (
        "Orientation for the pack owner. States that this file is sector calibration for the common facility-level OT engine.",
        "Python validates 01 Pack Metadata, then loads 02 (intensity), 03 (actor–scenario propensity), 05 (asset types), 09 (scenario downtime/capacity) and 12 (TTP catalogue). Placeholders 04/06–08/13–15 are format-only; 10–11 remain reference copies until separately promoted.",
    ),
    "01 Pack Metadata": (
        "Identity, OT domain, schema CRQ-OT-PACK-1.0, minimum engine 1.7.1, status and facility unit of analysis.",
        "Resolved from config/sector_pack_registry.json and checked fail-closed before any OT pack number is applied.",
    ),
    "02 Actor Weights": (
        "Sector actor-intensity multipliers (not a mix that must sum to 1) plus asset-type frequency overlays.",
        "Multiplies core OT campaign intensity for Nation State, Cybercriminal and Malicious Insider after the selected asset overlay. Assessment OT 06 still holds λ method, shares and facility exposure.",
    ),
    "03 Actor Scenario": (
        "Conditional share of each actor’s campaigns across the five OT adverse scenarios.",
        "Authoritative sector input. The OT engine reads this table from the pack (not OT 06). Rows must sum to 1.0. Facility λ, actor shares and exposure remain on OT 06.",
    ),
    "04 Route Definitions": (
        "Placeholder kept so OT packs use the same worksheet names as the FS IT pack.",
        "Not used. OT has no IT composite routes; stages S1–S5 and TTP gates stay in OT CORE.",
    ),
    "05 Actor Route": (
        "Registered plant / asset archetypes for this sector (tab name kept for format; content is asset types, not IT routes).",
        "Fail-closed lookup of the Run Setup asset type. Pack ID on the matched row must equal this workbook’s PACK_ID.",
    ),
    "06 Scenario Route": (
        "Placeholder for IT scenario-to-route propensity.",
        "Not used. OT scenario structure is core methodology, not an IT-style route mix.",
    ),
    "07 Stage Requirements": (
        "Placeholder for IT scenario-aware stage masks.",
        "Not used. Five-stage requirements are defined on OT CORE, not in the pack.",
    ),
    "08 Stage Priors": (
        "Placeholder for IT reference stage-through priors.",
        "Not used. Stage-through math is core OT; sector intensity is sheet 02.",
    ),
    "09 Scenario Parameters": (
        "Row-oriented sector P50/P99 downtime and capacity parameters (DOWNTIME_DAYS, CAPACITY_AFFECTED).",
        "Authoritative sector input. The OT engine loads this table from the pack (not OT 05 A58:G62). IT enterprise parameters are not used. Facility revenue/BI factor and driver rates remain on OT 05.",
    ),
    "10 Driver Rates": (
        "Pack P50/P99 rates for the common OT BIA drivers.",
        "Reference copy. Facility quantities and overrides stay on OT 05. Python does not currently reload these rates from the external pack file.",
    ),
    "11 Impact Driver Matrix": (
        "Pack driver-to-scenario applicability (1 = applies).",
        "Reference copy of the governed OT matrix. Runtime applicability is still the OT 05 / core BIA path.",
    ),
    "12 TTP Catalogue": (
        "ATT&CK for ICS techniques for this sector: default applicable Yes/No, facility gate, rationale and guidance URL.",
        "Every governed TTP on OT CORE - Scenario-TTP Map must have a row. Applicability and gates feed OT path feasibility; missing rows fail closed.",
    ),
    "13 Route TTP Map": (
        "Placeholder for IT route-to-TTP weights.",
        "Not used. OT technique-to-scenario mapping is OT CORE - Scenario-TTP Map; sector Yes/No is sheet 12.",
    ),
    "14 Control Reference": (
        "Placeholder for an IT-style reference control profile.",
        "Not used. OT control maturity and reference-relative adjustment are core OT (OT 04 / OT CORE).",
    ),
    "15 Control TTP Map": (
        "Placeholder for IT control-to-TTP efficacy.",
        "Not used. OT TTP-control efficacy is OT CORE - TTP-Control Map.",
    ),
    "16 Sources": (
        "Bibliography and limitations for this sector pack.",
        "Traceability only. The engine does not take λ or loss numbers from this sheet.",
    ),
    "17 Evidence Register": (
        "Evidence-grade scale used to label pack assumptions.",
        "Governance labels only. Grades do not rescale OT Monte Carlo results.",
    ),
    "18 Validation": (
        "Spreadsheet checks (PACK_ID, DOMAIN=OT, TTP and asset row counts).",
        "Advisory in Excel. Python repeats pack identity, TTP coverage and multiplier presence and fails closed independently.",
    ),
}


def apply_pack_sheet_blurb(ws, domain: str) -> None:
    """Write A2 purpose and A3 how-used. Does not change A1, A4 sheet-type, or table rows."""
    table = IT_PACK_BLURBS if domain == "IT" else OT_PACK_BLURBS
    blurb = table.get(ws.title)
    if not blurb:
        return
    purpose, how = blurb
    ws["A2"] = purpose
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A2"].font = font_body()
    ws.row_dimensions[2].height = 36
    if ws.title == "00 Pack Guide":
        return
    ws["A3"] = "How used in the model: " + how
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws["A3"].font = font_body()
    ws.row_dimensions[3].height = 48


def stamp_pack_workbook(path, domain: str) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    try:
        for name in wb.sheetnames:
            apply_pack_sheet_blurb(wb[name], domain)
        wb.save(path)
    finally:
        wb.close()
