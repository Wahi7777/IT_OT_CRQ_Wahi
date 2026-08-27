"""Rebuild the combined workbook: names, domain-first setup, navigation, pack audit labels."""

from __future__ import annotations

from pathlib import Path
import shutil

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from crq.sheet_names import (
    COMMON_OUTSIDE_IN,
    COMMON_RUN_SETUP,
    DIV_ADVANCED,
    DIV_DASHBOARDS,
    DIV_FILL_IT,
    DIV_FILL_OT,
    DIV_NAV,
    DIV_OWNERSHIP,
    DIV_RESULTS,
    IT_FILL,
    OLD_TO_NEW,
    OT_FILL,
    rewrite_formula_text,
)
from it_ot_crq.navigation import apply_workbook_navigation

COLOR_COMMON = "1F4E79"
COLOR_IT = "00B0F0"
COLOR_OT = "7030A0"
COLOR_CORE = "ED7D31"
COLOR_PACK = "C00080"
COLOR_CALC = "7F7F7F"
COLOR_RESULT = "548235"

GOVERNANCE = {
    "ASSESSMENT": "ASSESSMENT INPUT - USER EDITABLE",
    "PACK": "SECTOR PACK CALIBRATION - PACK OWNER ONLY",
    "CORE": "CORE MODEL METHODOLOGY - MODEL OWNER ONLY",
    "CALC": "CALCULATED OUTPUT - DO NOT EDIT",
    "DOCS": "MODEL DOCUMENTATION",
    "AUDIT": "VALIDATION / AUDIT",
}


def _rename_sheets(wb) -> None:
    for old, new in OLD_TO_NEW.items():
        if old in wb.sheetnames and new not in wb.sheetnames:
            wb[old].title = new


def _rewrite_formulas(wb) -> None:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = rewrite_formula_text(cell.value)
        for dv in ws.data_validations.dataValidation:
            if dv.formula1:
                dv.formula1 = rewrite_formula_text("=" + str(dv.formula1).lstrip("="))[1:]
            if dv.formula2:
                dv.formula2 = rewrite_formula_text("=" + str(dv.formula2).lstrip("="))[1:]
    for name in list(wb.defined_names.values()):
        if name.attr_text:
            name.attr_text = rewrite_formula_text("=" + str(name.attr_text).lstrip("="))[1:] if False else _rewrite_ref(name.attr_text)


def _rewrite_ref(text: str) -> str:
    out = str(text)
    for old, new in sorted(OLD_TO_NEW.items(), key=lambda kv: len(kv[0]), reverse=True):
        out = out.replace(f"'{old}'", f"'{new}'")
    return out


def _apply_domain_first_setup(wb) -> None:
    setup = wb[COMMON_RUN_SETUP]
    old_c6 = str(setup["C6"].value or "").strip()
    old_c7 = str(setup["C7"].value or "").strip()
    if old_c6 in {"IT", "OT"}:
        domain, sector = old_c6, old_c7 if not old_c7.startswith("=") else ""
    elif old_c6 == "Financial Services":
        domain, sector = "IT", old_c6
    elif old_c6 in {"Power Generation", "Energy Assets", "Manufacturing"}:
        domain, sector = "OT", old_c6
    else:
        domain, sector = "OT", "Power Generation"

    setup["A6"] = "Model domain"
    setup["B6"] = "Explicit user selection. Filters the Sector list."
    setup["C6"] = domain
    setup["A7"] = "Sector"
    setup["B7"] = "Filtered by Model domain. Resolves the external sector pack."
    setup["C7"] = sector
    setup["B8"] = "Filtered by Sector: Organisation for IT; facility archetype for OT."
    setup["D8"] = '=IF(C6="IT",IF(C8="Organisation","","IT requires Asset type = Organisation"),"")'
    setup["C12"] = '=IF(C6="IT","it_crq v1.1.1","ot_crq v1.7.1")'
    setup["C13"] = '=IF(C6="IT","Organisation","Facility")'
    setup["C11"] = (
        '="evidence/outside_in/"&IF(C6="OT",'
        "'OT 03 - Facility Inputs'!C16,'IT 03 - Organisation Inputs'!C16)&\"_outside_in.csv\""
    )
    setup["A50"] = "SHEETS TO COMPLETE — driven by Model domain (C6)"
    setup["A51"] = (
        '=IF(C6="IT","Complete the IT assessment section only. Ignore OT input tabs.",'
        'IF(C6="OT","Complete the OT assessment section only. Ignore IT input tabs.",'
        '"Select Model domain IT or OT."))'
    )

    setup["Z5"] = "IT"
    setup["Z6"] = "OT"
    setup["AA5"] = "Financial Services"
    setup["AB5"] = "Power Generation"
    setup["AB6"] = "Energy Assets"
    setup["AB7"] = "Manufacturing"
    setup["Z8"] = "Organisation"

    existing = {str(n) for n in wb.defined_names}
    named = {
        "IT_Sectors": f"'{COMMON_RUN_SETUP}'!$AA$5",
        "OT_Sectors": f"'{COMMON_RUN_SETUP}'!$AB$5:$AB$7",
        "Financial_Services": f"'{COMMON_RUN_SETUP}'!$Z$8",
        "Power_Generation": "'OT PACK - Loaded Registry'!$B$25:$B$30",
        "Energy_Assets": "'OT PACK - Loaded Registry'!$B$31:$B$35",
        "Manufacturing": "'OT PACK - Loaded Registry'!$B$36:$B$39",
    }
    # Asset lists still live on the loaded-registry audit view until packs are displayed; keep original rows.
    if "OT PACK - Loaded Registry" not in wb.sheetnames:
        named["Power_Generation"] = "'28 Sector Pack Registry'!$B$25:$B$30"
        named["Energy_Assets"] = "'28 Sector Pack Registry'!$B$31:$B$35"
        named["Manufacturing"] = "'28 Sector Pack Registry'!$B$36:$B$39"
    for name, attr in named.items():
        if name in existing:
            del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name=name, attr_text=attr))

    for dv in list(setup.data_validations.dataValidation):
        setup.data_validations.dataValidation.remove(dv)
    domain_dv = DataValidation(type="list", formula1='"IT,OT"', allow_blank=False, showErrorMessage=True)
    domain_dv.add("C6")
    sector_dv = DataValidation(
        type="list",
        formula1='INDIRECT($C$6&"_Sectors")',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Sector not valid for domain",
        error="Select a sector that belongs to the Model domain in C6.",
        promptTitle="Sector",
        prompt="This list is filtered by Model domain.",
        showInputMessage=True,
    )
    sector_dv.add("C7")
    asset_dv = DataValidation(
        type="list",
        formula1='INDIRECT(SUBSTITUTE($C$7," ","_"))',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="Asset type not valid for sector",
        error="Select an asset type that belongs to the Sector in C7.",
        promptTitle="Asset type / scope",
        prompt="This list is filtered by Sector.",
        showInputMessage=True,
    )
    asset_dv.add("C8")
    setup.add_data_validation(domain_dv)
    setup.add_data_validation(sector_dv)
    setup.add_data_validation(asset_dv)


def _label_sheets(wb) -> None:
    assessment = set(IT_FILL) | set(OT_FILL) | {COMMON_RUN_SETUP, COMMON_OUTSIDE_IN}
    pack = {n for n in wb.sheetnames if "PACK" in n or "Loaded" in n}
    calc = {n for n in wb.sheetnames if "CALC" in n or n.startswith("00 Dashboard") or n.startswith("00 Risk") or n.startswith("00 Loss") or n.startswith("00 What") or n.startswith("00 COMMON - Output") or n.startswith("00 COMMON - Engine")}
    for name in wb.sheetnames:
        ws = wb[name]
        if name in assessment and name not in {COMMON_RUN_SETUP, COMMON_OUTSIDE_IN}:
            tag, color = GOVERNANCE["ASSESSMENT"], COLOR_IT if name.startswith("IT ") else COLOR_OT
        elif name in {COMMON_RUN_SETUP, COMMON_OUTSIDE_IN}:
            tag, color = GOVERNANCE["ASSESSMENT"], COLOR_COMMON
        elif name in pack or name.startswith("IT 09"):
            tag, color = GOVERNANCE["PACK"] + " / ENGINE-LOADED SECTOR-PACK VALUES - DO NOT EDIT", COLOR_PACK
        elif "CORE" in name:
            tag, color = GOVERNANCE["CORE"], COLOR_CORE
        elif "CALC" in name or name.startswith("00 Dashboard") or name.startswith("00 Risk") or name.startswith("00 Loss") or name.startswith("00 What") or "Bridge" in name or "Engine Results" in name:
            tag, color = GOVERNANCE["CALC"], COLOR_RESULT if name.startswith("00 ") else COLOR_CALC
        elif "Validation" in name or "Audit" in name or "Sources" in name:
            tag, color = GOVERNANCE["AUDIT"], "1F4E79"
        elif name.startswith(("1 ", "2 ", "3 ", "4 ", "5 ", "FILL ")):
            tag, color = GOVERNANCE["DOCS"], COLOR_COMMON
        else:
            tag, color = GOVERNANCE["DOCS"], COLOR_CALC
        ws.sheet_properties.tabColor = color
        if ws["A2"].value in (None, "") or str(ws["A2"].value).startswith("ASSESSMENT INPUT") or str(ws["A2"].value).startswith("SECTOR PACK") or str(ws["A2"].value).startswith("CORE MODEL") or str(ws["A2"].value).startswith("CALCULATED") or str(ws["A2"].value).startswith("MODEL DOCUMENTATION") or str(ws["A2"].value).startswith("VALIDATION"):
            ws["A2"] = tag
            ws["A2"].font = Font(bold=True, size=9, color="FFFFFF")
            ws["A2"].fill = PatternFill("solid", fgColor=color)


def _ownership_index(wb) -> None:
    name = DIV_OWNERSHIP
    if name not in wb.sheetnames:
        ws = wb.create_sheet(name)
    else:
        ws = wb[name]
    headers = ["Group", "Domain", "Sector", "Assumption layer", "Sheet or file", "Purpose", "Owner", "Editable during assessment?", "Authoritative location", "Required evidence", "Calibration status"]
    ws["A1"] = "3 Assumption Ownership"
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, h)
        ws.cell(3, c).font = Font(bold=True, color="FFFFFF")
        ws.cell(3, c).fill = PatternFill("solid", fgColor=COLOR_COMMON)
    rows = [
        ("1 Common setup", "Common", "All", "Assessment setup", COMMON_RUN_SETUP, "Domain, sector, pack, basis, outside-in", "Assessment Team", "Yes", COMMON_RUN_SETUP, "Assessment ID / sector selection", "n/a"),
        ("1 Common setup", "Common", "All", "Optional evidence", COMMON_OUTSIDE_IN, "Approved outside-in findings", "Assessment Team", "Optional", COMMON_OUTSIDE_IN, "Approved CSV", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Assessment input", "IT 03 - Organisation Inputs", "Organisation scale and run inputs", "Assessment Team", "Yes", "IT 03 - Organisation Inputs", "Client evidence", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Optional override", "IT 04 - Exposure Adjustments", "Exposure-model recommendations", "Assessment Team", "Optional", "IT 04 - Exposure Adjustments", "Exposure model", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Assessment input", "IT 05 - Control Assessment", "Actual IT control maturity", "Assessment Team", "Yes", "IT 05 - Control Assessment", "Control evidence", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Optional override", "IT 06 - Impact & BIA Overrides", "Org-specific BIA overrides", "Assessment Team", "Optional", "IT 06; blank keeps pack", "Org BIA", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Optional override", "IT 07 - Assessment Adjustments", "Org frequency/threat overlays", "Assessment Team", "Optional", "IT 07; pack owns reference λ", "Threat intel", "n/a"),
        ("2 IT assessment", "IT", "Financial Services", "Assessment input", "IT 08 - Reporting Settings", "Basis, retention, tail metric", "Assessment Team", "Yes", "IT 08 - Reporting Settings", "Reporting policy", "n/a"),
        ("3 Core IT", "IT", "All IT", "Core methodology", "IT CORE / IT CALC sheets", "IT calculation method", "Core IT Model Owner", "No", "src/it_crq + IT CORE sheets", "n/a", "Governed"),
        ("4 FS pack", "IT", "Financial Services", "Sector pack", "sector_packs/IT/IT_CRQ_Sector_Pack_Financial_Services_v1_1_1.xlsx", "FS reference calibration", "Sector Pack Owner", "No", "External FS pack", "Pack evidence register", "Working priors"),
        ("5 OT assessment", "OT", "Selected OT sector", "Assessment input", "OT 03 - Facility Inputs", "Facility identity and architecture", "Assessment Team", "Yes", "OT 03 - Facility Inputs", "Facility evidence", "n/a"),
        ("5 OT assessment", "OT", "Selected OT sector", "Assessment input", "OT 04 - Control Assessment", "Actual OT control maturity", "Assessment Team", "Yes", "OT 04 - Control Assessment", "Control evidence", "n/a"),
        ("5 OT assessment", "OT", "Selected OT sector", "Facility override", "OT 05 - Impact & BIA", "Facility BIA and overrides", "Assessment Team", "Yes", "OT 05; pack owns sector BIA refs", "Facility BIA", "n/a"),
        ("5 OT assessment", "OT", "Selected OT sector", "Facility adjustment", "OT 06 - Assessment Adjustments", "Prudence, geography, reporting", "Assessment Team", "Yes", "OT 06; pack owns sector λ overlays", "Facility threat context", "n/a"),
        ("6 Core OT", "OT", "All OT", "Core methodology", "OT CORE sheets", "Five-stage path, controls, floors", "Core OT Model Owner", "No", "src/ot_crq + OT CORE sheets", "n/a", "Governed"),
        ("7 PG pack", "OT", "Power Generation", "Sector pack", "sector_packs/OT/OT_CRQ_Sector_Pack_Power_Generation_v1_6.xlsx", "PG calibration", "Sector Pack Owner", "No", "External PG pack", "Pack evidence", "PG-v1.6"),
        ("8 EA pack", "OT", "Energy Assets", "Sector pack", "sector_packs/OT/OT_CRQ_Sector_Pack_Energy_Assets_v1_0.xlsx", "EA calibration", "Sector Pack Owner", "No", "External EA pack", "Pack evidence", "EA-v1.0"),
        ("9 MF pack", "OT", "Manufacturing", "Sector pack", "sector_packs/OT/OT_CRQ_Sector_Pack_Manufacturing_v1_0.xlsx", "MF calibration", "Sector Pack Owner", "No", "External MF pack", "Pack evidence", "MF-v1.0"),
        ("10 Outputs", "Common", "All", "Calculated", "00 dashboards / Engine Results", "AAL, LEC, validation", "Python engine", "No", "00 COMMON - Engine Results", "n/a", "Run-specific"),
    ]
    for i, row in enumerate(rows, 4):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    ws.sheet_properties.tabColor = COLOR_COMMON


def _bia_instructions(wb) -> None:
    if "IT 06 - Impact & BIA Overrides" in wb.sheetnames:
        ws = wb["IT 06 - Impact & BIA Overrides"]
        ws["A4"] = (
            "Enter an override only where organisation-specific evidence supports departure from the "
            "Financial Services pack baseline. A blank override retains the pack value. "
            "Pack baseline is read-only; Used P50/P99 are calculated."
        )
    if "OT 05 - Impact & BIA" in wb.sheetnames:
        ws = wb["OT 05 - Impact & BIA"]
        ws["A4"] = (
            "Enter facility-specific values or overrides only where facility evidence supports departure "
            "from the selected OT sector-pack baseline. Pack-owned rows are read-only."
        )
    if "OT 06 - Assessment Adjustments" in wb.sheetnames:
        ws = wb["OT 06 - Assessment Adjustments"]
        ws["A4"] = (
            "Assessment adjustments only. Effective λ = Core reference × sector-pack actor multiplier × "
            "asset-type overlay × facility exposure × geography × client actor activity × prudence. "
            "Do not edit pack multipliers here."
        )
    if "IT 07 - Assessment Adjustments" in wb.sheetnames:
        ws = wb["IT 07 - Assessment Adjustments"]
        ws["A4"] = (
            "Organisation-level overlays only. Pack-owned campaign rate, actor weights and route priors are read-only."
        )
    for name in ("OT PACK PG - Loaded Values", "OT PACK EA - Loaded Values", "OT PACK MF - Loaded Values"):
        if name in wb.sheetnames:
            wb[name]["A1"] = "ENGINE-LOADED SECTOR-PACK VALUES - DO NOT EDIT"


def rebuild_combined_workbook(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    try:
        _rename_sheets(wb)
        _rewrite_formulas(wb)
        _apply_domain_first_setup(wb)
        _ownership_index(wb)
        _bia_instructions(wb)
        _label_sheets(wb)
        apply_workbook_navigation(wb, domain=None)
        wb.save(path)
    finally:
        wb.close()


def _stamp_it_pack(path: Path) -> None:
    wb = openpyxl.load_workbook(path)
    try:
        if "00 Pack Guide" in wb.sheetnames:
            ws = wb["00 Pack Guide"]
            ws["A1"] = "IT Sector Pack - Financial Services - v1.1.1"
            ws["A3"] = (
                "This workbook contains Financial Services reference calibration. "
                "Do not enter client-specific or assessment-specific data in this workbook."
            )
        meta = wb["01 Pack Metadata"]
        found = False
        for r in range(16, 80):
            if str(meta.cell(r, 1).value or "").strip() == "DOMAIN":
                meta.cell(r, 2).value = "IT"
                found = True
                break
        if not found:
            meta["A14"] = "DOMAIN"
            meta["B14"] = "IT"
        if "09 Sector Pack Registry" in wb.sheetnames:
            pass
        from crq.pack_sheet_guide import apply_pack_sheet_blurb

        for name in wb.sheetnames:
            apply_pack_sheet_blurb(wb[name], "IT")
        wb.save(path)
    finally:
        wb.close()


def _update_it_registry_path(path: Path, cell: str = "D16") -> None:
    wb = openpyxl.load_workbook(path)
    try:
        if "09 Sector Pack Registry" in wb.sheetnames:
            wb["09 Sector Pack Registry"][cell] = "sector_packs/IT/IT_CRQ_Sector_Pack_Financial_Services_v1_1_1.xlsx"
        if "IT 09 - Sector Pack Registry" in wb.sheetnames:
            ws = wb["IT 09 - Sector Pack Registry"]
            for r in range(15, 20):
                for c in range(1, 8):
                    val = str(ws.cell(r, c).value or "")
                    if "IT_CRQ_Sector_Pack_Financial_Services" in val:
                        ws.cell(r, c).value = "sector_packs/IT/IT_CRQ_Sector_Pack_Financial_Services_v1_1_1.xlsx"
        wb.save(path)
    finally:
        wb.close()


def rebuild_release(project_root: Path) -> None:
    from ot_crq.packs import extract_ot_packs

    combined = project_root / "model" / "Guided_IT_OT_CRQ_Model_v1_0.xlsx"
    if not combined.is_file():
        combined = project_root / "model" / "Guided_IT_OT_CRQ_Combined_Model_v0_3.xlsx"
    extract_ot_packs(combined, project_root)
    src = project_root / "sector_packs" / "IT_CRQ_Sector_Pack_Financial_Services_v1_1_1.xlsx"
    dest_dir = project_root / "sector_packs" / "IT"
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / "IT_CRQ_Sector_Pack_Financial_Services_v1_1_1.xlsx"
    if src.is_file():
        shutil.copy2(src, dest)
    elif dest.is_file():
        pass
    else:
        raise FileNotFoundError("Financial Services pack not found to copy into sector_packs/IT/")
    _stamp_it_pack(dest)
    rebuild_combined_workbook(combined)
    _update_it_registry_path(combined)
    native = project_root / "model" / "it" / "Guided_IT_CRQ_Model_v1_1_1_Dashboard.xlsx"
    if native.is_file():
        _update_it_registry_path(native)
    v03 = project_root / "model" / "Guided_IT_OT_CRQ_Combined_Model_v0_3.xlsx"
    if v03.is_file() and v03.resolve() != combined.resolve():
        shutil.copy2(combined, v03)

