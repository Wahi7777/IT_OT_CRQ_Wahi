from __future__ import annotations

import csv
import os
import shutil
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.workbook.properties import CalcProperties

from crq import COMBINED_MODEL_VERSION, ROUTER_VERSION
from crq.io_safety import (
    MAX_CSV_ROWS,
    MAX_FINDINGS,
    assert_not_canonical,
    enforce_csv_limits,
    sanitize_untrusted_text,
    sha256_file,
    validate_user_path,
)
from crq.pack_registry import resolve_pack, validate_pack_file
from crq.sheet_names import (
    COMMON_BRIDGE,
    COMMON_ENGINE,
    COMMON_OUTSIDE_IN,
    COMMON_RUN_SETUP,
    IT_BIA,
    IT_CONTROLS,
    IT_EXPOSURE,
    IT_ORG,
    OT_FACILITY,
)
from it_crq.engine import ENGINE_VERSION as IT_ENGINE_VERSION
from it_ot_crq.navigation import apply_workbook_navigation
from ot_crq.engine import ENGINE_VERSION as OT_ENGINE_VERSION

SECTOR_DOMAIN = {
    "Financial Services": "IT",
    "Power Generation": "OT",
    "Energy Assets": "OT",
    "Manufacturing": "OT",
}

CANONICAL_COMBINED_NAMES = (
    "model/Guided_IT_OT_CRQ_Model_v1_0.xlsx",
    "model/Guided_IT_OT_CRQ_Combined_Model_v0_3.xlsx",
)
IT_TEMPLATE = "model/it/Guided_IT_CRQ_Model_v1_1_1_Dashboard.xlsx"

IT_SHEET_MAP = {
    "IT CORE - Model Guide": "01 Model Guide",
    "IT 02 - Executive Summary": "02 Executive Summary",
    "IT 03 - Organisation Inputs": "03 Organisation Inputs",
    "IT 04 - Exposure Adjustments": "04 Exposure Adjustments",
    "IT 05 - Control Assessment": "05 Control Assessment",
    "IT 06 - Impact & BIA Overrides": "06 Impact and BIA",
    "IT 07 - Assessment Adjustments": "07 Frequency Assumptions",
    "IT 08 - Reporting Settings": "08 Tail Risk Settings",
    "IT 09 - Sector Pack Registry": "09 Sector Pack Registry",
    "IT CALC - TTP Relevance": "10 TTP Relevance Calc",
    "IT CALC - Attack Path": "11 Attack Path Calc",
    "IT CALC - Frequency and Success": "12 Frequency and Success",
    "IT CALC - Consequence": "13 Consequence Calc",
    "IT CALC - Actor Scenario Matrix": "14 Actor Scenario Matrix",
    "IT CALC - Tail Risk Metrics": "15 Tail Risk Metrics",
    "IT CALC - Control What-If": "16 Control What-If",
    "IT CALC - Aggregate LECs": "17 Aggregate LECs",
    "IT CALC - Scenario LECs": "18 Scenario LECs",
    "IT CALC - Actor LECs": "19 Actor LECs",
    "IT CALC - Risk Charts": "20 Risk Charts",
    "IT 21 - Sources and Evidence": "21 Sources and Evidence",
    "IT 22 - Validation Tests": "22 Validation Tests",
}

IT_INPUT_SHEETS = {
    "IT 03 - Organisation Inputs",
    "IT 04 - Exposure Adjustments",
    "IT 05 - Control Assessment",
    "IT 06 - Impact & BIA Overrides",
    "IT 07 - Assessment Adjustments",
    "IT 08 - Reporting Settings",
    "IT 09 - Sector Pack Registry",
}

POSITIVE_CONDITIONING = {
    ("INTERNET_OT", "Yes"),
    ("REMOTE_ACCESS", "Yes"),
    ("REMOTE_ACCESS", "Controlled"),
    ("IT_OT_CONNECTIVITY", "Yes"),
}


def _copy_values(src, dst) -> None:
    """Copy values and merge ranges. Do not materialise merged-title text into every cell."""
    from copy import copy
    from openpyxl.utils import range_boundaries

    for rng in list(dst.merged_cells.ranges):
        dst.unmerge_cells(str(rng))
    rows = max(src.max_row or 1, dst.max_row or 1)
    cols = max(src.max_column or 1, dst.max_column or 1)
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            sc = src.cell(r, c)
            if sc.__class__.__name__ == "MergedCell":
                continue
            target = dst.cell(r, c)
            target.value = sc.value
            if sc.has_style:
                target._style = copy(sc._style)
    for rng in src.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(rng))
        dst.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _find_key_row(ws, key: str, key_col: int = 1) -> int | None:
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, key_col).value or "").strip() == key:
            return r
    return None


def _sector_from_workbook(path: Path) -> str:
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        ws = wb[COMMON_RUN_SETUP]
        domain = str(ws["C6"].value or "").strip()
        sector = str(ws["C7"].value or "").strip()
        if sector not in SECTOR_DOMAIN and domain in SECTOR_DOMAIN:
            # Pre-rebuild workbooks stored sector in C6.
            sector, domain = domain, SECTOR_DOMAIN.get(domain, domain)
    finally:
        wb.close()
    if sector not in SECTOR_DOMAIN:
        raise ValueError(
            f"Unsupported sector {sector!r}. Expected one of: {', '.join(SECTOR_DOMAIN)}"
        )
    if domain and domain != SECTOR_DOMAIN[sector]:
        raise ValueError(
            f"Model domain {domain!r} does not match sector {sector!r} ({SECTOR_DOMAIN[sector]})."
        )
    return sector


def _resolve_workbook_outside_in(path: Path, project_root: Path, domain: str) -> Path | None:
    wb = openpyxl.load_workbook(path, read_only=False, data_only=False)
    try:
        setup = wb[COMMON_RUN_SETUP]
        apply_scan = str(setup["C10"].value or "No").strip().lower()
        if apply_scan not in {"yes", "true", "1"}:
            return None

        raw = str(setup["C11"].value or "").strip()
        if raw and not raw.startswith("="):
            candidate = Path(raw).expanduser()
            return candidate if candidate.is_absolute() else (project_root / candidate)

        input_sheet = IT_ORG if domain == "IT" else OT_FACILITY
        rr = _find_key_row(wb[input_sheet], "ASSESSMENT_ID")
        assessment_id = str(wb[input_sheet].cell(rr, 3).value or "assessment").strip() if rr else "assessment"
        return project_root / "evidence" / "outside_in" / f"{assessment_id}_outside_in.csv"
    finally:
        wb.close()


def _default_output(model: Path) -> Path:
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return Path("outputs") / f"{model.stem}_results_{stamp}.xlsx"


def _snapshot_outside_in(wb, csv_path: Path) -> tuple[str, list[dict[str, str]]]:
    enforce_csv_limits(csv_path)
    with csv_path.open("r", encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))
    if len(rows) > MAX_CSV_ROWS:
        raise ValueError(f"Outside-in CSV has {len(rows)} rows; maximum is {MAX_CSV_ROWS}.")

    ws = wb[COMMON_OUTSIDE_IN]
    start = 22
    headers = list(rows[0].keys()) if rows else []
    if headers:
        for c, h in enumerate(headers, 1):
            ws.cell(start, c).value = sanitize_untrusted_text(h)
        for rr, row in enumerate(rows, start + 1):
            for c, h in enumerate(headers, 1):
                ws.cell(rr, c).value = sanitize_untrusted_text(row.get(h))

    approved = [
        r for r in rows
        if str(r.get("approved") or "").strip().lower() in {"yes", "true", "1", "approved"}
        and str(r.get("quantitative_target") or "").strip()
    ]
    if len(approved) > MAX_FINDINGS:
        raise ValueError(f"Too many approved findings ({len(approved)}); maximum is {MAX_FINDINGS}.")
    return f"Loaded {len(rows)} finding(s); {len(approved)} approved quantitative adjustment(s)", approved


def _polarity_allows(baseline, conditioned, row: dict[str, str]) -> None:
    """Positive evidence may contradict assumed absence; absence does not prove non-existence."""
    polarity = str(row.get("evidence_polarity") or "positive").strip().lower()
    b = str(baseline or "").strip().lower()
    c = str(conditioned or "").strip().lower()
    absence_credit = str(row.get("absence_credit_approved") or "").strip().lower() in {
        "yes", "true", "1",
    }
    if polarity in {"negative", "absence", "non_observation", "non-observation"}:
        if b in {"yes", "true", "1", "controlled"} and c in {"no", "false", "0", "none"}:
            if not absence_credit:
                raise ValueError(
                    "Failure to observe an exposure externally cannot automatically condition "
                    "a Yes/Controlled assumption to No unless absence_credit_approved is set."
                )


OI_SCHEMA_VERSION = "1.0.0"
SUFFICIENT_ATTRIBUTION = {"high", "medium", "confirmed", "certain"}
EXPOSURE_TARGETS = {"facility_input", "organisation_input", "exposure_route"}
EXPLOIT_TARGETS = {"ttp_prevent_barrier"}


def _attribution_ok(row: dict[str, str]) -> bool:
    conf = str(row.get("attribution_confidence") or "").strip().lower()
    override = str(row.get("attribution_override") or "").strip().lower() in {"yes", "true", "1"}
    return conf in SUFFICIENT_ATTRIBUTION or override


def _ensure_overlay_sheet(wb):
    if "00 OI Engine Overlay" not in wb.sheetnames:
        ws = wb.create_sheet("00 OI Engine Overlay")
        ws["A1"] = "TTP ID"
        ws["B1"] = "Prevent barrier multiplier (1=unchanged; <1 weakens barrier)"
        ws["C1"] = "Classification"
        ws["C2"] = "ENGINE_WRITTEN overlay — not a second P(success) model"
        return ws
    return wb["00 OI Engine Overlay"]


def _write_overlay_multiplier(wb, tid: str, multiplier: float) -> None:
    ws = _ensure_overlay_sheet(wb)
    for r in range(2, ws.max_row + 2):
        existing = str(ws.cell(r, 1).value or "").strip()
        if existing in {"", tid}:
            ws.cell(r, 1).value = tid
            ws.cell(r, 2).value = multiplier
            return
    raise ValueError("Unable to write OI overlay row.")


def _apply_approved_adjustments(wb, domain: str, rows: list[dict[str, str]], run_id: str) -> list[dict]:
    rationale = wb["00 COMMON - OI Audit Trail"]
    out_row = max(16, rationale.max_row + 1)
    seen_fields: set[tuple[str, str]] = set()
    finding_roles: dict[str, set[str]] = {}
    applied_records: list[dict] = []

    header = [
        "Finding ID", "Domain", "Sector", "Facility / organisation", "Actor", "Scenario",
        "Route / TTP / stage", "Affected parameter", "Baseline value", "Conditioned value",
        "Direction of change", "Evidence", "Attribution confidence", "Rationale",
        "Double-counting assessment", "Approval", "Reviewer", "Run ID", "Timestamp",
    ]
    for c, h in enumerate(header, 1):
        if rationale.cell(15, c).value in (None, ""):
            rationale.cell(15, c).value = h

    for row in rows:
        target = str(row.get("quantitative_target") or "").strip()
        field = str(row.get("model_field") or "").strip()
        conditioned = sanitize_untrusted_text(row.get("conditioned_value"))
        if conditioned.startswith("'"):
            conditioned = conditioned[1:]
        finding_id = sanitize_untrusted_text(row.get("finding_id") or "")
        reason = sanitize_untrusted_text(row.get("adjustment_rationale") or "")
        approver = sanitize_untrusted_text(row.get("approver") or "")
        assessment = sanitize_untrusted_text(
            row.get("double_counting_assessment") or "Single approved parameter change"
        )
        sector = sanitize_untrusted_text(row.get("sector") or "")

        def write_rationale(status, baseline_v, target_disp, direction, extra_reason=""):
            nonlocal out_row
            values = [
                finding_id, domain, sector,
                sanitize_untrusted_text(row.get("organisation_attribution") or row.get("facility_attribution") or ""),
                sanitize_untrusted_text(row.get("facility_attribution") or ""),
                sanitize_untrusted_text(row.get("actor") or ""),
                sanitize_untrusted_text(row.get("scenario") or ""),
                sanitize_untrusted_text(row.get("mapped_ttp") or row.get("model_stage") or field),
                target_disp, baseline_v, conditioned, direction,
                sanitize_untrusted_text(row.get("evidence") or row.get("source_scanner") or row.get("cve") or ""),
                sanitize_untrusted_text(row.get("attribution_confidence") or ""),
                extra_reason or reason, assessment, status, approver, run_id,
                datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            ]
            for c, value in enumerate(values, 1):
                rationale.cell(out_row, c).value = value
            out_row += 1

        if not _attribution_ok(row):
            write_rationale(
                "not applied — insufficient attribution",
                None, f"{target}::{field}", "unchanged",
                "Attribution confidence below High/Medium; no model change.",
            )
            continue

        key = (target, field)
        if key in seen_fields:
            raise ValueError(
                f"Two approved outside-in rows target the same parameter {field!r}; "
                "refusing possible double counting."
            )
        seen_fields.add(key)
        roles = finding_roles.setdefault(finding_id, set())
        if roles & EXPOSURE_TARGETS and target in EXPLOIT_TARGETS:
            if "exposure vs exploitability" not in assessment.lower():
                raise ValueError(
                    f"Finding {finding_id} would change both exposure/feasibility and a TTP barrier. "
                    "Record double_counting_assessment including 'exposure vs exploitability' "
                    "or split into separate findings."
                )
        if roles & EXPLOIT_TARGETS and target in EXPOSURE_TARGETS:
            if "exposure vs exploitability" not in assessment.lower():
                raise ValueError(
                    f"Finding {finding_id} would change both a TTP barrier and exposure/feasibility."
                )
        roles.add(target)

        baseline = None
        applied = False
        target_display = target

        if domain == "OT" and target == "facility_input":
            ws = wb[OT_FACILITY]
            rr = _find_key_row(ws, field)
            if rr is None:
                raise ValueError(f"Outside-in target not found in OT Facility Inputs: {field}")
            baseline = ws.cell(rr, 3).value
            _polarity_allows(baseline, conditioned, row)
            ws.cell(rr, 3).value = conditioned
            target_display = f"03 Facility Inputs::{field}"
            applied = True

        elif domain == "OT" and target == "ttp_prevent_barrier":
            try:
                mult = float(conditioned)
            except (TypeError, ValueError) as exc:
                raise ValueError("ttp_prevent_barrier conditioned_value must be numeric in (0, 1].") from exc
            if not 0.0 < mult <= 1.0:
                raise ValueError("ttp_prevent_barrier multiplier must be in (0, 1].")
            baseline = 1.0
            _write_overlay_multiplier(wb, field, mult)
            target_display = f"Prevent barrier multiplier::{field}"
            applied = True

        elif domain == "IT" and target == "organisation_input":
            ws = wb[IT_ORG]
            rr = _find_key_row(ws, field)
            if rr is None:
                raise ValueError(f"Outside-in target not found in IT Organisation Inputs: {field}")
            baseline = ws.cell(rr, 3).value
            _polarity_allows(baseline, conditioned, row)
            ws.cell(rr, 3).value = conditioned
            target_display = f"IT 03 Organisation Inputs::{field}"
            applied = True

        elif domain == "IT" and target == "exposure_route":
            parts = field.split(".", 1)
            if len(parts) != 2:
                raise ValueError("IT exposure_route model_field must be ROUTE_ID.FIELD")
            rid, attr = parts[0].strip(), parts[1].strip().upper()
            ws = wb[IT_EXPOSURE]
            rr = _find_key_row(ws, rid)
            if rr is None:
                raise ValueError(f"Outside-in route not found in IT Exposure Adjustments: {rid}")
            col_map = {"OPPORTUNITY": 5, "S1": 6, "S2": 7, "S3": 8, "S4": 9, "S5": 10, "IMPACT": 11}
            if attr not in col_map:
                raise ValueError(f"Unsupported IT exposure-route field: {attr}")
            cc = col_map[attr]
            baseline = ws.cell(rr, cc).value
            try:
                value = float(conditioned)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"Conditioned value for {field} must be numeric") from exc
            ws.cell(rr, cc).value = value
            target_display = f"IT 04 Exposure Adjustments::{field}"
            applied = True

        if not applied:
            raise ValueError(
                f"Approved outside-in target {target!r} is not supported for domain {domain}. "
                "Use an existing governed model input rather than a direct risk-score multiplier."
            )

        try:
            bnum, cnum = float(baseline), float(conditioned)
            direction = "increase" if cnum > bnum else ("decrease" if cnum < bnum else "unchanged")
        except (TypeError, ValueError):
            direction = "changed" if str(baseline) != str(conditioned) else "unchanged"

        write_rationale("approved", baseline, target_display, direction)
        applied_records.append({
            "finding_id": finding_id, "target": target, "field": field,
            "baseline": baseline, "conditioned": conditioned,
        })
    return applied_records


def _write_cell(ws, row: int, col: int, value) -> None:
    cell = ws.cell(row, col)
    if cell.__class__.__name__ == "MergedCell":
        return
    cell.value = value


OT_ONLY_SHEETS = {
    OT_FACILITY,
    "04 Control Assessment",
    "05 Impact and BIA",
    "02 Executive Summary",
    "OT 06 - Assessment Adjustments",
    "07 Control Assumptions",
    "08 Scenario Definitions",
    "09 Scenario-TTP Map",
    "10 Mapping Change Log",
    "11 Actor-TTP Applicability",
    "12 Facility Feasibility",
    "13 TTP-Control Map",
    "14 TTP Relevance Calc",
    "15 Attack Path Calc",
    "16 Threat and Loss Frequency",
    "17 Consequence Calc",
    "18 Actor-Scenario Matrix",
    "19 Tail Risk Metrics",
    "20 Control What-If",
    "21 Aggregate LECs",
    "22 Scenario LECs",
    "23 Actor LECs",
    "24 Actor-Scenario LEC Data",
    "25 Risk Charts",
    "26 Sources and Evidence",
    "27 Validation Tests",
    "28 Sector Pack Registry",
    "29 Sector TTP Rationale",
    "30 Power Generation Pack",
    "31 Energy Assets Pack",
    "32 Manufacturing Pack",
    "33 Sector Pack Comparison",
}

IT_ONLY_PREFIX = "IT "


def apply_run_setup_asset_dropdowns(wb) -> None:
    """Domain C6 filters Sector C7; Sector filters Asset C8."""
    from it_ot_crq.rebuild import _apply_domain_first_setup

    _apply_domain_first_setup(wb)


def apply_output_sheet_visibility(wb, domain: str) -> None:
    """Hide pack/engine tabs and the domain the operator is not assessing."""
    apply_workbook_navigation(wb, domain=domain or None)


def _write_python_snapshot(wb, result: dict, meta: dict) -> None:
    """Write engine metrics as values. Dashboards may only select Best Estimate vs Prudent."""
    setup = wb[COMMON_RUN_SETUP]
    view = str(setup["C9"].value or "Prudent").strip()
    if view not in {"Prudent", "Best Estimate"}:
        setup["C9"] = "Prudent"
        view = "Prudent"
    setup["A9"] = "Basis"
    meta["reporting_view"] = view
    ws = wb[COMMON_BRIDGE]
    if COMMON_ENGINE in wb.sheetnames:
        er_ws = wb[COMMON_ENGINE]
    else:
        er_ws = wb.create_sheet(COMMON_ENGINE)
    er_ws["A1"] = "ENGINE_WRITTEN authoritative metrics (Python). Do not re-implement CRQ math here."
    er_ws["A2"] = "Metric"
    er_ws["B2"] = "Best Estimate"
    er_ws["C2"] = "Prudent"
    er_ws["D2"] = "Field class"

    ws["A2"] = (
        "Authoritative last-run metrics are ENGINE_WRITTEN values in B/C and H–J, "
        "duplicated on 00 COMMON - Engine Results. Dashboard formulas may only select Best Estimate vs Prudent."
    )
    ws["K13"] = "Field class"
    metrics = [
        (15, "AAL", "best_aal", "prudent_aal"),
        (16, "VaR 95", "best_var95", "prudent_var95"),
        (17, "TVaR 95", "best_tvar95", "prudent_tvar95"),
        (18, "VaR 99", "best_var99", "prudent_var99"),
        (19, "TVaR 99", "best_tvar99", "prudent_tvar99"),
        (20, "P(any successful event)", "best_pany", "prudent_pany"),
        (21, "Successful-event frequency / yr", "best_event_frequency", "prudent_event_frequency"),
        (22, "Attempt / campaign frequency / yr", "best_attempt_frequency", "prudent_attempt_frequency"),
        (23, "P(annual loss exceeds tolerance)", "best_p_exceed_tolerance", "prudent_p_exceed_tolerance"),
    ]
    _write_cell(ws, 14, 1, "Metric")
    _write_cell(ws, 14, 2, "Best Estimate")
    _write_cell(ws, 14, 3, "Prudent")
    for i, (row, name, bk, pk) in enumerate(metrics, start=3):
        _write_cell(ws, row, 1, name)
        _write_cell(ws, row, 2, result.get(bk))
        _write_cell(ws, row, 3, result.get(pk))
        _write_cell(ws, row, 11, "ENGINE_WRITTEN")
        er_ws.cell(i, 1).value = name
        er_ws.cell(i, 2).value = result.get(bk)
        er_ws.cell(i, 3).value = result.get(pk)
        er_ws.cell(i, 4).value = "ENGINE_WRITTEN"
    _write_cell(ws, 24, 1, "Risk-appetite status")
    _write_cell(ws, 24, 2, result.get("appetite_status") or "Tolerance not set")
    _write_cell(ws, 24, 3, result.get("appetite_status") or "Tolerance not set")
    _write_cell(ws, 24, 11, "ENGINE_WRITTEN")
    if result.get("executive_narrative"):
        _write_cell(ws, 12, 1, "Executive Risk Story")
        _write_cell(ws, 12, 2, result.get("executive_narrative"))
        _write_cell(ws, 12, 11, "ENGINE_WRITTEN")


    _write_cell(ws, 25, 1, "ACTOR CONTRIBUTION")
    _write_cell(ws, 26, 1, "Actor")
    _write_cell(ws, 26, 2, "Best Estimate")
    _write_cell(ws, 26, 3, "Prudent")
    _write_cell(ws, 26, 11, "ENGINE_WRITTEN")
    r = 27
    er_ws["A12"] = "Actor"
    er_ws["B12"] = "Best Estimate"
    er_ws["C12"] = "Prudent"
    er_r = 13
    actors_be = result.get("actor_aal_best") or result.get("actor_aal") or {}
    actors_pr = result.get("actor_aal_prudent") or result.get("actor_aal") or {}
    for name in actors_be:
        _write_cell(ws, r, 1, name)
        _write_cell(ws, r, 2, actors_be.get(name))
        _write_cell(ws, r, 3, actors_pr.get(name))
        _write_cell(ws, r, 11, "ENGINE_WRITTEN")
        er_ws.cell(er_r, 1).value = name
        er_ws.cell(er_r, 2).value = actors_be.get(name)
        er_ws.cell(er_r, 3).value = actors_pr.get(name)
        er_ws.cell(er_r, 4).value = "ENGINE_WRITTEN"
        r += 1
        er_r += 1

    _write_cell(ws, 32, 1, "SCENARIO CONTRIBUTION")
    _write_cell(ws, 33, 1, "Scenario")
    _write_cell(ws, 33, 2, "Best Estimate")
    _write_cell(ws, 33, 3, "Prudent")
    _write_cell(ws, 33, 11, "ENGINE_WRITTEN")
    r = 34
    er_ws["A20"] = "Scenario"
    er_ws["B20"] = "Best Estimate"
    er_ws["C20"] = "Prudent"
    er_r = 21
    scen_be = result.get("scenario_aal_best") or result.get("scenario_aal") or {}
    scen_pr = result.get("scenario_aal_prudent") or result.get("scenario_aal") or {}
    for name in scen_be:
        _write_cell(ws, r, 1, name)
        _write_cell(ws, r, 2, scen_be.get(name))
        _write_cell(ws, r, 3, scen_pr.get(name))
        _write_cell(ws, r, 11, "ENGINE_WRITTEN")
        er_ws.cell(er_r, 1).value = name
        er_ws.cell(er_r, 2).value = scen_be.get(name)
        er_ws.cell(er_r, 3).value = scen_pr.get(name)
        er_ws.cell(er_r, 4).value = "ENGINE_WRITTEN"
        r += 1
        er_r += 1

    identity = [
        (6, "Selected sector", meta["sector"], "STATIC_METADATA"),
        (7, "Model domain", meta["domain"], "STATIC_METADATA"),
        (8, "Asset type / scope", meta.get("asset_type"), "STATIC_METADATA"),
        (9, "Pack ID", meta["pack_id"], "STATIC_METADATA"),
        (10, "Basis", meta.get("reporting_view") or "Prudent", "STATIC_METADATA"),
    ]
    for row, label, value, klass in identity:
        _write_cell(ws, row, 1, label)
        _write_cell(ws, row, 2, value)
        _write_cell(ws, row, 11, klass)

    ws["H4"] = "PYTHON ENGINE SNAPSHOT"
    for row, label, value in [
        (6, "Selected sector", meta["sector"]),
        (7, "Model domain", meta["domain"]),
        (8, "Engine version", meta.get("engine_version")),
        (9, "Pack ID", meta["pack_id"]),
        (10, "Pack ID", meta["pack_id"]),
        (11, "Engine version", meta.get("engine_version")),
        (12, "Outside-in applied", meta["outside_in_applied"]),
        (13, "Validation", result.get("validation") or result.get("overall")),
    ]:
        _write_cell(ws, row, 8, label)
        _write_cell(ws, row, 9, value)
    ws["H14"] = "Metric"
    ws["I14"] = "Best Estimate"
    ws["J14"] = "Prudent"
    for row, name, bk, pk in metrics:
        _write_cell(ws, row, 8, name)
        _write_cell(ws, row, 9, result.get(bk))
        _write_cell(ws, row, 10, result.get(pk))

    ws["H25"] = "Actor AAL (reporting view)"
    r = 26
    for name, value in (result.get("actor_aal") or {}).items():
        _write_cell(ws, r, 8, name)
        _write_cell(ws, r, 9, value)
        r += 1
    ws["H32"] = "Scenario AAL (reporting view)"
    r = 33
    for name, value in (result.get("scenario_aal") or {}).items():
        _write_cell(ws, r, 8, name)
        _write_cell(ws, r, 9, value)
        r += 1

    ws["H40"] = "RUN METADATA (Python)"
    rows = [
        (41, "Combined model version", COMBINED_MODEL_VERSION),
        (42, "Router version", ROUTER_VERSION),
        (43, "IT engine version", IT_ENGINE_VERSION),
        (44, "OT engine version", OT_ENGINE_VERSION),
        (45, "Sector", meta["sector"]),
        (46, "Domain", meta["domain"]),
        (47, "Asset type", meta.get("asset_type")),
        (48, "Sector-pack ID", meta["pack_id"]),
        (49, "Sector-pack version / status", meta["pack_status"]),
        (50, "Outside-in schema version", meta.get("oi_schema", OI_SCHEMA_VERSION)),
        (51, "Outside-in applied", meta["outside_in_applied"]),
        (52, "Input workbook SHA-256", meta["input_hash"]),
        (53, "Outside-in file SHA-256", meta.get("oi_hash") or "n/a"),
        (54, "Random seed", result.get("random_seed")),
        (55, "Simulation years", result.get("simulation_years")),
        (56, "Run timestamp UTC", meta["timestamp"]),
        (57, "Validation status", result.get("validation") or result.get("overall")),
        (58, "Run ID", meta["run_id"]),
    ]
    for row, label, value in rows:
        _write_cell(ws, row, 8, label)
        _write_cell(ws, row, 9, value)
        _write_cell(ws, row, 11, "STATIC_METADATA")
        er_ws.cell(row - 20, 6).value = label
        er_ws.cell(row - 20, 7).value = value
        er_ws.cell(row - 20, 8).value = "STATIC_METADATA"

    # Selected-view cache for Python recon (Dashboard A8/D8/G8 remain presentation formulas).
    view = str(meta.get("reporting_view") or "Prudent").strip()
    pick = "best" if view == "Best Estimate" else "prudent"
    dash = wb["00 Dashboard"]
    _write_cell(dash, 70, 1, "ENGINE_WRITTEN selected-view cache (Python)")
    _write_cell(dash, 71, 1, "AAL")
    _write_cell(dash, 71, 2, result.get(f"{pick}_aal"))
    _write_cell(dash, 72, 1, "TVaR 95")
    _write_cell(dash, 72, 2, result.get(f"{pick}_tvar95"))
    _write_cell(dash, 73, 1, "TVaR 99")
    _write_cell(dash, 73, 2, result.get(f"{pick}_tvar99"))
    _write_cell(dash, 74, 1, "VaR 95")
    _write_cell(dash, 74, 2, result.get(f"{pick}_var95"))
    _write_cell(dash, 75, 1, "VaR 99")
    _write_cell(dash, 75, 2, result.get(f"{pick}_var99"))
    _write_cell(dash, 76, 1, "P(any)")
    _write_cell(dash, 76, 2, result.get(f"{pick}_pany"))
    _write_cell(dash, 77, 1, "Attempt frequency")
    _write_cell(dash, 77, 2, result.get(f"{pick}_attempt_frequency"))
    _write_cell(dash, 78, 1, "Successful-event frequency")
    _write_cell(dash, 78, 2, result.get(f"{pick}_event_frequency"))

    er_ws["A40"] = "CONTROL WHAT-IF (ENGINE_WRITTEN; ranked by AAL reduction on the Prudent basis for freeze comparability)"
    headers = [
        "Control ID", "Control name", "Channel", "Current maturity", "What-If maturity",
        "Baseline AAL", "What-If AAL", "AAL Reduction", "AAL Reduction %",
        "Baseline AAL", "What-If AAL", "AAL Reduction", "AAL Reduction %",
        "Baseline TVaR 95", "What-If TVaR 95", "Baseline TVaR 99", "What-If TVaR 99",
        "Baseline TVaR 95", "What-If TVaR 95", "Baseline TVaR 99", "What-If TVaR 99",
        "Baseline successful-event frequency", "What-If successful-event frequency",
        "Baseline successful-event frequency", "What-If successful-event frequency",
        "Mapped",
    ]
    for c, h in enumerate(headers, 1):
        er_ws.cell(41, c).value = h
    er_ws["F40"] = "Best Estimate"
    er_ws["J40"] = "Prudent"
    whatifs = list(result.get("whatifs") or [])
    whatifs.sort(key=lambda x: float((x.get("prudent") or {}).get("reduction") or x.get("reduction") or 0), reverse=True)
    for i, x in enumerate(whatifs, 42):
        be = x.get("be") or {}
        pr = x.get("prudent") or {}
        vals = [
            x.get("cid"), x.get("name"), x.get("channel"), x.get("current"), x.get("whatif") or x.get("next"),
            be.get("baseline"), be.get("aal"), be.get("reduction"), be.get("pct"),
            pr.get("baseline"), pr.get("aal"), pr.get("reduction"), pr.get("pct"),
            be.get("tvar95_base"), be.get("tvar95"), be.get("tvar99_base"), be.get("tvar99"),
            pr.get("tvar95_base"), pr.get("tvar95"), pr.get("tvar99_base"), pr.get("tvar99"),
            be.get("freq_base"), be.get("event_freq"), pr.get("freq_base"), pr.get("event_freq"),
            "Yes" if x.get("mapped", True) else "No",
        ]
        for c, v in enumerate(vals, 1):
            er_ws.cell(i, c).value = v

    whatif_dash = wb["00 What-If"]
    _write_cell(whatif_dash, 4, 1, "Basis")
    _write_cell(whatif_dash, 4, 2, "='00 COMMON - Run Setup'!C9")
    _write_cell(whatif_dash, 8, 6, "AAL Reduction")
    _write_cell(whatif_dash, 8, 7, "AAL Reduction %")
    _write_cell(whatif_dash, 8, 8, "Baseline AAL")
    _write_cell(whatif_dash, 8, 9, "What-If AAL")
    _write_cell(whatif_dash, 8, 10, "TVaR 99 Reduction")
    for idx in range(12):
        row = 9 + idx
        src = 42 + idx
        _write_cell(whatif_dash, row, 1, f"='00 COMMON - Engine Results'!A{src}")
        _write_cell(whatif_dash, row, 2, f"='00 COMMON - Engine Results'!B{src}")
        _write_cell(whatif_dash, row, 3, f"='00 COMMON - Engine Results'!C{src}")
        _write_cell(whatif_dash, row, 4, f"='00 COMMON - Engine Results'!D{src}")
        _write_cell(whatif_dash, row, 5, f"='00 COMMON - Engine Results'!E{src}")
        _write_cell(whatif_dash, row, 6, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Engine Results'!H{src},'00 COMMON - Engine Results'!L{src})")
        _write_cell(whatif_dash, row, 7, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Engine Results'!I{src},'00 COMMON - Engine Results'!M{src})")
        _write_cell(whatif_dash, row, 8, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Engine Results'!F{src},'00 COMMON - Engine Results'!J{src})")
        _write_cell(whatif_dash, row, 9, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Engine Results'!G{src},'00 COMMON - Engine Results'!K{src})")
        _write_cell(whatif_dash, row, 10, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",IF('00 COMMON - Engine Results'!P{src}=\"\",\"\",'00 COMMON - Engine Results'!P{src}-'00 COMMON - Engine Results'!Q{src}),IF('00 COMMON - Engine Results'!T{src}=\"\",\"\",'00 COMMON - Engine Results'!T{src}-'00 COMMON - Engine Results'!U{src}))")

    _write_cell(dash, 4, 7, "Basis")
    for drow, brow in ((31, 27), (32, 28), (33, 29), (34, 30)):
        _write_cell(dash, drow, 2, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Output Bridge'!B{brow},'00 COMMON - Output Bridge'!C{brow})")
        _write_cell(dash, drow, 10, f"=IF('00 COMMON - Run Setup'!C9=\"Best Estimate\",'00 COMMON - Output Bridge'!B{brow + 7},'00 COMMON - Output Bridge'!C{brow + 7})")

    transfer = wb["00 Risk Transfer"]
    _write_cell(transfer, 4, 1, "Basis")
    _write_cell(transfer, 4, 2, "='00 COMMON - Run Setup'!C9")
    from it_ot_crq.dashboards import populate_common_dashboards
    from it_ot_crq.reporting.populate import populate_reporting_suite

    populate_common_dashboards(wb, result, meta)
    populate_reporting_suite(wb, result, meta)
    apply_run_setup_asset_dropdowns(wb)
    apply_output_sheet_visibility(wb, str(meta.get("domain") or ""))


def _run_it(combined_path: Path, sector_pack_dir: Path, work_dir: Path, project_root: Path, run_whatifs: bool = True) -> dict:
    from it_crq.engine import refresh as it_refresh
    from it_ot_crq.reporting.inputs import read_appetite_inputs, read_insurance_programme

    if not sector_pack_dir.is_dir():
        raise FileNotFoundError(f"IT sector-pack directory not found: {sector_pack_dir}")
    packs = list(sector_pack_dir.glob("*.xlsx"))
    if not packs:
        raise FileNotFoundError(f"No IT sector packs found in {sector_pack_dir}")

    template = project_root / IT_TEMPLATE
    if not template.is_file():
        raise FileNotFoundError(f"Governed IT template not found: {template}")
    work_dir.mkdir(parents=True, exist_ok=True)
    it_input = work_dir / "it_input.xlsx"
    it_result = work_dir / "it_result.xlsx"
    shutil.copy2(template, it_input)

    combined = openpyxl.load_workbook(combined_path, data_only=False)
    staged = openpyxl.load_workbook(it_input, data_only=False)
    try:
        for combined_name in IT_INPUT_SHEETS:
            _copy_values(combined[combined_name], staged[IT_SHEET_MAP[combined_name]])
        staged["03 Organisation Inputs"]["C19"] = str(combined[COMMON_RUN_SETUP]["C7"].value or "Financial Services").strip()
        org = staged["03 Organisation Inputs"]
        bia = staged["06 Impact and BIA"]
        for dest, src_row in ((7, 21), (8, 22), (9, 23), (10, 25), (11, 26), (12, 27), (13, 28)):
            bia.cell(dest, 2).value = org.cell(src_row, 3).value
        it08_view = str(combined["IT 08 - Reporting Settings"]["B16"].value or "").strip()
        c9 = str(combined[COMMON_RUN_SETUP]["C9"].value or "Prudent").strip()
        if it08_view in {"Best Estimate", "Prudent", "Both"}:
            reporting = it08_view
            combined[COMMON_RUN_SETUP]["C9"] = reporting
        elif c9 in {"Best Estimate", "Prudent", "Both"}:
            reporting = c9
        else:
            reporting = "Prudent"
        tail_ws = staged["08 Tail Risk Settings"]
        rr = _find_key_row(tail_ws, "REPORTING_VIEW")
        if rr is not None:
            tail_ws.cell(rr, 2).value = reporting
        combined.save(combined_path)
        staged.save(it_input)
    finally:
        combined.close()
        staged.close()

    wb_inputs = openpyxl.load_workbook(combined_path, data_only=False)
    try:
        appetite = read_appetite_inputs(wb_inputs)
        programme = read_insurance_programme(wb_inputs)
    finally:
        wb_inputs.close()

    result = it_refresh(
        it_input,
        it_result,
        sector_pack_dir=sector_pack_dir,
        appetite_inputs=appetite,
        insurance_programme=programme,
        run_packages=run_whatifs or os.environ.get("CRQ_RUN_PACKAGES") == "1",
        run_sensitivity=bool(os.environ.get("CRQ_RUN_SENSITIVITY")),
        run_whatifs=run_whatifs,
    )

    combined = openpyxl.load_workbook(combined_path, data_only=False)
    it_wb = openpyxl.load_workbook(it_result, data_only=False)
    try:
        for combined_name, source_name in IT_SHEET_MAP.items():
            _copy_values(it_wb[source_name], combined[combined_name])
        combined.calculation = combined.calculation or CalcProperties()
        combined.calculation.calcMode = "auto"
        combined.calculation.fullCalcOnLoad = True
        combined.calculation.forceFullCalc = True
        combined.save(combined_path)
    finally:
        combined.close()
        it_wb.close()
    result["engine_label"] = f"it_crq {result.get('pack_id', IT_ENGINE_VERSION)}"
    return result


def _run_ot(combined_path: Path, run_whatifs: bool = True, project_root: Path | None = None) -> dict:
    os.environ.setdefault("OT_CRQ_USE_XLSX_BACKEND", "1")
    try:
        from ot_crq.engine import refresh as ot_refresh
    except ImportError as exc:
        raise RuntimeError(
            "OT sector selected but the governed ot_crq runtime is not installed in src/ot_crq."
        ) from exc

    wb = openpyxl.load_workbook(combined_path, data_only=False)
    try:
        setup = wb[COMMON_RUN_SETUP]
        domain = str(setup["C6"].value or "").strip()
        sector = str(setup["C7"].value or "").strip()
        if sector not in SECTOR_DOMAIN and domain in SECTOR_DOMAIN:
            sector, domain = domain, SECTOR_DOMAIN[domain]
        if SECTOR_DOMAIN.get(sector) != "OT" or domain not in {"", "OT"}:
            raise RuntimeError(f"OT engine invoked for non-OT sector {sector!r}.")
        asset_type = str(setup["C8"].value or "").strip()
        if not asset_type:
            raise ValueError("Asset type must be populated for OT sectors.")
        reporting = str(setup["C9"].value or "Prudent").strip()
        wb[OT_FACILITY]["C20"] = sector
        wb[OT_FACILITY]["C21"] = asset_type
        freq_ws = wb["OT 06 - Assessment Adjustments"]
        rr = _find_key_row(freq_ws, "REPORTING_VIEW")
        if rr is not None:
            freq_ws.cell(rr, 4).value = reporting
        wb.save(combined_path)
    finally:
        wb.close()

    rec = resolve_pack("OT", sector, project_root)
    from it_ot_crq.reporting.inputs import read_appetite_inputs, read_insurance_programme

    wb_inputs = openpyxl.load_workbook(combined_path, data_only=False)
    try:
        appetite = read_appetite_inputs(wb_inputs)
        programme = read_insurance_programme(wb_inputs)
    finally:
        wb_inputs.close()

    result = ot_refresh(
        str(combined_path),
        str(combined_path),
        run_whatifs=run_whatifs,
        sector_pack_path=rec.path(project_root),
        project_root=project_root,
        appetite_inputs=appetite,
        insurance_programme=programme,
        run_packages=run_whatifs or os.environ.get("CRQ_RUN_PACKAGES") == "1",
        run_sensitivity=bool(os.environ.get("CRQ_RUN_SENSITIVITY")),
    )
    result["engine_label"] = f"ot_crq {result.get('sector_pack_id', OT_ENGINE_VERSION)}"
    result["pack_id"] = result.get("sector_pack_id")
    result["pack_status"] = result.get("sector_pack_status")
    result["validation"] = result.get("overall")
    return result


def run_combined(
    model: Path,
    output: Path | None = None,
    sector_pack_dir: Path = Path("sector_packs"),
    outside_in_file: Path | None = None,
    work_dir: Path = Path("_work"),
    run_whatifs: bool = True,
) -> dict[str, object]:
    from crq.pack_registry import project_root_from

    model = validate_user_path(model, must_exist=True)
    project_root = project_root_from(model)
    output = (output or _default_output(model)).expanduser()
    if ".." in Path(output).parts:
        raise ValueError(f"Path traversal is not permitted: {output}")
    output = output.resolve()
    for name in CANONICAL_COMBINED_NAMES:
        candidate = project_root / name
        if candidate.is_file():
            assert_not_canonical(output, candidate)
    if output == model:
        raise ValueError("Refusing to overwrite the canonical combined input template.")
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(model, output)

    sector = _sector_from_workbook(output)
    domain = SECTOR_DOMAIN[sector]
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    input_hash = sha256_file(model)

    outside_status = "Not used"
    oi_applied = "No"
    oi_hash = None
    resolved_outside = None
    if outside_in_file is not None:
        resolved_outside = validate_user_path(outside_in_file, must_exist=True)
    else:
        candidate = _resolve_workbook_outside_in(output, project_root, domain)
        if candidate is not None:
            resolved_outside = validate_user_path(candidate, must_exist=True)

    if resolved_outside is not None:
        oi_hash = sha256_file(resolved_outside)
        wb = openpyxl.load_workbook(output, data_only=False)
        try:
            outside_status, approved = _snapshot_outside_in(wb, resolved_outside)
            applied_rows = _apply_approved_adjustments(wb, domain, approved, run_id)
            oi_applied = "Yes" if applied_rows else "Scan supplied, not approved"
            wb.calculation = wb.calculation or CalcProperties()
            wb.calculation.calcMode = "auto"
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.save(output)
        finally:
            wb.close()

    pack_dir = sector_pack_dir.expanduser().resolve()
    it_pack_dir = project_root / "sector_packs" / "IT"
    if it_pack_dir.is_dir():
        pack_dir = it_pack_dir if domain == "IT" else pack_dir
    work = work_dir.expanduser().resolve()
    if domain == "IT":
        engine_result = _run_it(output, pack_dir, work, project_root, run_whatifs=run_whatifs)
    else:
        engine_result = _run_ot(output, run_whatifs=run_whatifs, project_root=project_root)

    meta = {
        "sector": sector,
        "domain": domain,
        "engine_version": IT_ENGINE_VERSION if domain == "IT" else OT_ENGINE_VERSION,
        "pack_id": engine_result.get("pack_id") or engine_result.get("sector_pack_id"),
        "pack_status": engine_result.get("pack_status") or engine_result.get("sector_pack_status"),
        "outside_in_applied": oi_applied,
        "input_hash": input_hash,
        "oi_hash": oi_hash,
        "timestamp": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
        "run_id": run_id,
        "oi_schema": OI_SCHEMA_VERSION,
        "reporting_view": None,
        "asset_type": engine_result.get("asset_type"),
    }
    setup_wb = openpyxl.load_workbook(output, data_only=False)
    try:
        meta["reporting_view"] = str(setup_wb[COMMON_RUN_SETUP]["C9"].value or "Prudent").strip()
        meta["asset_type"] = meta["asset_type"] or str(setup_wb[COMMON_RUN_SETUP]["C8"].value or "").strip()
        _write_python_snapshot(setup_wb, engine_result, meta)
        setup_wb.calculation = setup_wb.calculation or CalcProperties()
        setup_wb.calculation.calcMode = "auto"
        setup_wb.calculation.fullCalcOnLoad = True
        setup_wb.save(output)
    finally:
        setup_wb.close()

    return {
        "output": str(output),
        "sector": sector,
        "domain": domain,
        "engine": engine_result.get("engine_label"),
        "engine_version": meta["engine_version"],
        "pack_id": meta["pack_id"],
        "pack_status": meta["pack_status"],
        "outside_in": outside_status,
        "outside_in_applied": oi_applied,
        "validation": engine_result.get("validation") or engine_result.get("overall"),
        "aal": engine_result.get("prudent_aal") if domain == "OT" else engine_result.get("prudent_aal"),
        "best_aal": engine_result.get("best_aal"),
        "prudent_aal": engine_result.get("prudent_aal"),
        "tvar99": engine_result.get("prudent_tvar99"),
        "engine_result": engine_result,
        "input_hash": input_hash,
        "run_id": run_id,
    }
