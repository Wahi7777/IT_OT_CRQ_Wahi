"""Create and theme the seven presentation sheets."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from it_ot_crq.reporting import REPORTING_SHEETS

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
SECTION_FILL = PatternFill("solid", fgColor="2E75B6")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CURRENCY = r"$#,##0"
PCT = "0.0%"
THIN = Border(
    left=Side(style="thin", color="D0D7D0"),
    right=Side(style="thin", color="D0D7D0"),
    top=Side(style="thin", color="D0D7D0"),
    bottom=Side(style="thin", color="D0D7D0"),
)


def ensure_reporting_sheets(wb) -> None:
    existing = set(wb.sheetnames)
    # Insert after dashboards divider if present
    anchor = 0
    if "1 Executive dashboards" in wb.sheetnames:
        anchor = wb.sheetnames.index("1 Executive dashboards") + 1
    for i, name in enumerate(REPORTING_SHEETS):
        if name not in existing:
            wb.create_sheet(name, anchor + i)
        ws = wb[name]
        ws.sheet_view.showGridLines = False
        for col in range(1, 16):
            ws.column_dimensions[get_column_letter(col)].width = 16
        ws.column_dimensions["A"].width = 36
        ws.column_dimensions["B"].width = 18


def _unmerge_all(ws) -> None:
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))


def _title(ws, text: str) -> None:
    _unmerge_all(ws)
    ws["A1"] = text
    ws["A1"].font = Font(name="Carlito", bold=True, size=16, color="1F4E79")
    ws.merge_cells("A1:L1")
    ws.row_dimensions[1].height = 28


def _subtitle(ws, text: str) -> None:
    ws["A2"] = text
    ws["A2"].font = Font(name="Carlito", size=10, color="5B6B7A")
    ws["A2"].alignment = Alignment(wrap_text=True)
    try:
        ws.merge_cells("A2:L2")
    except ValueError:
        pass
    ws.row_dimensions[2].height = 36


def write_cell(ws, row, col, value, fmt=None, *, input_cell=False, bold=False):
    cell = ws.cell(row, col, value)
    cell.font = Font(name="Carlito", size=10, bold=bold)
    cell.border = THIN
    if fmt:
        cell.number_format = fmt
    if input_cell:
        cell.fill = INPUT_FILL
    return cell


def section_header(ws, row, text, cols=12):
    ws.cell(row, 1, text)
    ws.cell(row, 1).fill = SECTION_FILL
    ws.cell(row, 1).font = Font(name="Carlito", bold=True, color="FFFFFF", size=11)
    for c in range(1, cols + 1):
        ws.cell(row, c).fill = SECTION_FILL
    return row + 1
