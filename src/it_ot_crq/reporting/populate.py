"""Populate the seven presentation views from an engine result dict."""

from __future__ import annotations

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font

from crq.metrics import multi_year_event_probability
from it_ot_crq.reporting import PACKAGE_NAMES, REPORTING_SHEETS
from it_ot_crq.reporting.inputs import (
    ensure_appetite_inputs_on_run_setup,
    ensure_insurance_inputs,
    read_appetite_inputs,
    read_insurance_programme,
    read_revenue,
)
from it_ot_crq.reporting.layout import (
    CURRENCY,
    PCT,
    _subtitle,
    _title,
    ensure_reporting_sheets,
    section_header,
    write_cell,
)

STAGE_NAMES = {
    "S1": "Initial access",
    "S2": "Establish presence",
    "S3": "Lateral movement / process access",
    "S4": "Understand the process",
    "S5": "Achieve adverse outcome",
}


def _pct_of_revenue(value, revenue):
    if value is None or revenue in (None, 0):
        return None
    return float(value) / float(revenue)


def _clear_body(ws, start_row=3, max_row=120, max_col=14):
    from it_ot_crq.reporting.layout import _unmerge_all
    _unmerge_all(ws)
    for r in range(start_row, max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c).value = None


def populate_executive(ws, result, view: str):
    _clear_body(ws)
    _title(ws, "01 — Executive Risk Story")
    _subtitle(ws, "Headline annual-aggregate financial risk from the selected run. Metrics use AAL / VaR 95 / TVaR 95 / VaR 99 / TVaR 99.")
    pick = "best" if view == "Best Estimate" else "prudent"
    cards = [
        ("P(at least one successful material event/year)", result.get(f"{pick}_pany"), PCT),
        ("AAL", result.get(f"{pick}_aal"), CURRENCY),
        ("VaR 95", result.get(f"{pick}_var95"), CURRENCY),
        ("TVaR 95", result.get(f"{pick}_tvar95"), CURRENCY),
        ("VaR 99", result.get(f"{pick}_var99"), CURRENCY),
        ("TVaR 99", result.get(f"{pick}_tvar99"), CURRENCY),
        ("P(annual aggregate loss exceeds tolerance)", result.get(f"{pick}_p_exceed_tolerance"), PCT),
        ("Risk-appetite status", result.get("appetite_status") or "Tolerance not set", None),
    ]
    section_header(ws, 4, "HEADLINE CARDS")
    for i, (label, val, fmt) in enumerate(cards):
        row = 5 + i
        write_cell(ws, row, 1, label, bold=True)
        write_cell(ws, row, 2, val, fmt)

    revenue = result.get("annual_revenue_at_risk")
    section_header(ws, 14, "SECONDARY CONTEXT")
    p3 = result.get(f"{pick}_pany_3y")
    p5 = result.get(f"{pick}_pany_5y")
    if p3 is None and result.get(f"{pick}_pany") is not None:
        p3 = multi_year_event_probability(result[f"{pick}_pany"], 3)
        p5 = multi_year_event_probability(result[f"{pick}_pany"], 5)
    write_cell(ws, 15, 1, "3-year successful-event probability")
    write_cell(ws, 15, 2, p3, PCT)
    write_cell(ws, 16, 1, "5-year successful-event probability")
    write_cell(ws, 16, 2, p5, PCT)
    write_cell(ws, 17, 1, "Annual revenue at risk")
    write_cell(ws, 17, 2, revenue, CURRENCY)
    r = 18
    for label, key in (
        ("AAL % of revenue", f"{pick}_aal"),
        ("VaR 95 % of revenue", f"{pick}_var95"),
        ("TVaR 95 % of revenue", f"{pick}_tvar95"),
        ("VaR 99 % of revenue", f"{pick}_var99"),
        ("TVaR 99 % of revenue", f"{pick}_tvar99"),
    ):
        write_cell(ws, r, 1, label)
        write_cell(ws, r, 2, _pct_of_revenue(result.get(key), revenue), PCT)
        r += 1
    write_cell(ws, r, 1, "Largest AAL-contributing scenario")
    write_cell(ws, r, 2, result.get("top_aal_contributor"))
    write_cell(ws, r + 1, 1, "Largest TVaR 99-contributing scenario")
    write_cell(ws, r + 1, 2, result.get("top_tvar99_contributor"))
    write_cell(ws, r + 2, 1, "Largest TVaR 99-contributing loss component")
    write_cell(
        ws,
        r + 2,
        2,
        result.get("top_tvar99_component")
        or result.get("top_tvar99_component_name")
        or "Not available",
    )
    pkgs = result.get("control_packages") or []
    best_pkg = max(pkgs, key=lambda p: p.get("aal_reduction") or 0) if pkgs else None
    whatifs = result.get("whatifs") or []
    best_ctrl = max(whatifs, key=lambda w: w.get("reduction") or 0) if whatifs else None
    write_cell(ws, r + 3, 1, "Highest-value modelled control improvement")
    write_cell(ws, r + 3, 2, (best_ctrl or {}).get("name") or (best_pkg or {}).get("name"))

    section_header(ws, r + 5, "CURRENT → TARGET CONTROLS → POST-INSURANCE")
    write_cell(ws, r + 6, 1, "Position", bold=True)
    for c, h in enumerate(("AAL", "VaR 95", "TVaR 95", "VaR 99", "TVaR 99"), 2):
        write_cell(ws, r + 6, c, h, bold=True)
    target = next((p for p in pkgs if p.get("name") == "Target"), None)
    ins = result.get("insurance_analysis") or {}
    mres = ins.get("metrics_residual") or {}
    rows = [
        ("Current ground-up", result.get(f"{pick}_aal"), result.get(f"{pick}_var95"), result.get(f"{pick}_tvar95"), result.get(f"{pick}_var99"), result.get(f"{pick}_tvar99")),
        ("After target controls", (target or {}).get("aal"), (target or {}).get("var95"), (target or {}).get("tvar95"), (target or {}).get("var99"), (target or {}).get("tvar99")),
        ("After insurance (residual)", mres.get("AAL"), mres.get("VaR95"), mres.get("TVaR95"), mres.get("VaR99"), mres.get("TVaR99")),
    ]
    for i, row in enumerate(rows):
        write_cell(ws, r + 7 + i, 1, row[0])
        for c, v in enumerate(row[1:], 2):
            write_cell(ws, r + 7 + i, c, v, CURRENCY)

    section_header(ws, r + 11, "EXECUTIVE NARRATIVE")
    write_cell(ws, r + 12, 1, result.get("executive_narrative") or "Run the model to generate the narrative.")
    ws.cell(r + 12, 1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r + 12, start_column=1, end_row=r + 14, end_column=8)


def populate_formation(ws, result, domain: str):
    _clear_body(ws)
    _title(ws, "02 — Risk Formation")
    _subtitle(ws, "Frequency funnel from campaigns to annual average loss. Conditional success is frequency-weighted.")
    form = result.get("formation") or {}
    section_header(ws, 4, "FUNNEL")
    headers = ["Stage", "Value", "Formula / source", "Interpretation"]
    for c, h in enumerate(headers, 1):
        write_cell(ws, 5, c, h, bold=True)
    rows = [
        ("Campaigns / year", form.get("campaigns_per_year"), "Simulated Poisson campaign rate after modifiers", "Material campaigns arriving at the facility/organisation"),
        ("Applicable campaigns / year", form.get("applicable_campaigns_per_year"), "Campaigns that enter the attack-path model", "Same as campaigns when all are facility-applicable"),
        ("Conditional attack-path success probability", form.get("conditional_success_probability"), "Successful-event frequency ÷ applicable campaign frequency", "Frequency-weighted path success — not an unweighted mean of route probabilities"),
        ("Successful events / year", form.get("successful_events_per_year"), "Mean successful campaigns per year", "Expected annual event count"),
        ("P(at least one successful event / year)", form.get("p_any_successful_event"), "Share of years with annual aggregate loss > 0", "Annual event probability"),
        ("AAL", form.get("aal"), "Mean annual aggregate loss including zero-loss years", "Expected annual financial loss"),
    ]
    for i, (a, b, c, d) in enumerate(rows):
        write_cell(ws, 6 + i, 1, a)
        write_cell(ws, 6 + i, 2, b, PCT if "probability" in a.lower() or a.startswith("P(") else ("0.000" if i < 4 else CURRENCY))
        write_cell(ws, 6 + i, 3, c)
        write_cell(ws, 6 + i, 4, d)

    as_rows = list(result.get("actor_scenario_rows") or [])
    has_route = any(r.get("route") for r in as_rows)
    section_header(ws, 14, "ACTOR × SCENARIO RECONCILIATION" + (" (by route)" if has_route else ""))
    cols = ["Actor", "Scenario"]
    if has_route:
        cols.append("Route")
    cols.extend(
        ["Campaign freq", "Applicable freq", "Success probability", "Successful-event freq", "Annual event probability", "AAL"]
    )
    for c, h in enumerate(cols, 1):
        write_cell(ws, 15, c, h, bold=True)
    for i, row in enumerate(as_rows):
        r = 16 + i
        col = 1
        write_cell(ws, r, col, row.get("actor")); col += 1
        write_cell(ws, r, col, row.get("scenario")); col += 1
        if has_route:
            write_cell(ws, r, col, row.get("route")); col += 1
        write_cell(ws, r, col, row.get("campaign_frequency"), "0.000"); col += 1
        write_cell(ws, r, col, row.get("applicable_frequency"), "0.000"); col += 1
        write_cell(ws, r, col, row.get("success_probability"), PCT); col += 1
        write_cell(ws, r, col, row.get("successful_event_frequency"), "0.000"); col += 1
        write_cell(ws, r, col, row.get("annual_event_probability"), PCT); col += 1
        write_cell(ws, r, col, row.get("aal"), CURRENCY)

    # Place stage diagnostic AFTER the reconciliation table (fixed row 34 collided with IT route rows).
    stage_row = 16 + len(as_rows) + 2
    section_header(ws, stage_row, "STAGE DIAGNOSTIC (business-readable)")
    write_cell(ws, stage_row + 1, 1, "Stage code", bold=True)
    write_cell(ws, stage_row + 1, 2, "Business name", bold=True)
    for i, (code, name) in enumerate(STAGE_NAMES.items()):
        write_cell(ws, stage_row + 2 + i, 1, code)
        write_cell(ws, stage_row + 2 + i, 2, name)
    note_row = stage_row + 2 + len(STAGE_NAMES)
    write_cell(ws, note_row, 1, "Domain note")
    write_cell(
        ws,
        note_row,
        2,
        "OT uses the five ICS composite stages above. IT uses organisation attack routes — each row above is one actor×scenario×route cell; AAL sums to the funnel AAL."
        if domain == "IT"
        else "Selected actor×scenario stage barriers are assessed on OT CALC - Attack Path.",
    )
    if as_rows and form.get("aal") is not None:
        write_cell(ws, note_row + 1, 1, "Reconciliation")
        write_cell(
            ws,
            note_row + 1,
            2,
            f"Sum of table AAL = {sum(float(r.get('aal') or 0) for r in as_rows):,.2f}; funnel AAL = {float(form.get('aal')):,.2f}",
        )


def populate_scenarios(ws, result):
    _clear_body(ws)
    _title(ws, "03 — Scenario Analysis")
    _subtitle(
        ws,
        "Standalone scenario VaR/TVaR are from each scenario’s own annual aggregate distribution. "
        "Portfolio TVaR contribution columns allocate aggregate TVaR across scenarios on the same "
        "worst-tail trials — they are not standalone scenario TVaR and must sum to aggregate TVaR.",
    )
    section_header(ws, 4, "SCENARIO TABLE")
    headers = [
        "Scenario", "Campaign freq", "Successful-event freq", "Annual event probability",
        "AAL", "VaR 95 (standalone)", "TVaR 95 (standalone)", "VaR 99 (standalone)", "TVaR 99 (standalone)",
        "% of total AAL", "Portfolio TVaR 95 contribution", "Portfolio TVaR 99 contribution",
        "Primary operational driver", "Primary financial driver",
    ]
    for c, h in enumerate(headers, 1):
        write_cell(ws, 5, c, h, bold=True)
    for i, row in enumerate(result.get("scenario_analysis") or []):
        vals = [
            row.get("name"), row.get("campaign_frequency"), row.get("successful_event_frequency"),
            row.get("annual_event_probability"), row.get("aal"), row.get("var95"), row.get("tvar95"),
            row.get("var99"), row.get("tvar99"), row.get("pct_total_aal"), row.get("contrib_tvar95"),
            row.get("contrib_tvar99"), row.get("primary_operational_driver"), row.get("primary_financial_driver"),
        ]
        for c, v in enumerate(vals, 1):
            fmt = CURRENCY if c in (5, 6, 7, 8, 9, 11, 12) else (PCT if c in (4, 10) else ("0.000" if c in (2, 3) else None))
            write_cell(ws, 6 + i, c, v, fmt)
    n_scen = len(result.get("scenario_analysis") or [])
    if n_scen:
        write_cell(ws, 6 + n_scen, 1, "Sum of portfolio TVaR 99 contributions", bold=True)
        write_cell(
            ws,
            6 + n_scen,
            12,
            sum(float(r.get("contrib_tvar99") or 0) for r in result["scenario_analysis"]),
            CURRENCY,
        )
        write_cell(ws, 7 + n_scen, 1, "Aggregate TVaR 99 (must match sum above)")
        pick = "prudent" if str(result.get("reporting_view") or result.get("detailed_view") or "").startswith("P") else "best"
        write_cell(ws, 7 + n_scen, 12, result.get(f"{pick}_tvar99") or result.get("TVaR99"), CURRENCY)

    actor_start = 10 + n_scen
    section_header(ws, actor_start, "ACTOR TABLE")
    ah = [
        "Actor", "Campaign freq", "Successful-event freq", "Annual event probability",
        "AAL", "VaR 95 (standalone)", "TVaR 95 (standalone)", "VaR 99 (standalone)", "TVaR 99 (standalone)",
        "% AAL", "Portfolio TVaR 95 contribution", "Portfolio TVaR 99 contribution",
    ]
    for c, h in enumerate(ah, 1):
        write_cell(ws, actor_start + 1, c, h, bold=True)
    for i, row in enumerate(result.get("actor_analysis") or []):
        vals = [
            row.get("name"), row.get("campaign_frequency"), row.get("successful_event_frequency"),
            row.get("annual_event_probability"), row.get("aal"), row.get("var95"), row.get("tvar95"),
            row.get("var99"), row.get("tvar99"), row.get("pct_total_aal"), row.get("contrib_tvar95"), row.get("contrib_tvar99"),
        ]
        for c, v in enumerate(vals, 1):
            fmt = CURRENCY if c in (5, 6, 7, 8, 9, 11, 12) else (PCT if c in (4, 10) else ("0.000" if c in (2, 3) else None))
            write_cell(ws, actor_start + 2 + i, c, v, fmt)

    # Charts from AAL / portfolio TVaR99 contribution columns — place below actor table
    chart_row = actor_start + 2 + len(result.get("actor_analysis") or []) + 2
    if result.get("scenario_analysis"):
        write_cell(ws, chart_row, 1, "Scenario")
        write_cell(ws, chart_row, 2, "AAL")
        write_cell(ws, chart_row, 3, "Portfolio TVaR 99 contribution")
        for i, row in enumerate(result["scenario_analysis"]):
            write_cell(ws, chart_row + 1 + i, 1, row.get("name"))
            write_cell(ws, chart_row + 1 + i, 2, row.get("aal"), CURRENCY)
            write_cell(ws, chart_row + 1 + i, 3, row.get("contrib_tvar99"), CURRENCY)
        ws._charts = []
        bar = BarChart()
        bar.type = "bar"
        bar.title = "AAL contribution by scenario"
        last = chart_row + len(result["scenario_analysis"])
        bar.add_data(Reference(ws, min_col=2, min_row=chart_row, max_row=last), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=chart_row + 1, max_row=last))
        ws.add_chart(bar, "E" + str(chart_row))
        bar2 = BarChart()
        bar2.type = "bar"
        bar2.title = "Portfolio TVaR 99 contribution by scenario"
        bar2.add_data(Reference(ws, min_col=3, min_row=chart_row, max_row=last), titles_from_data=True)
        bar2.set_categories(Reference(ws, min_col=1, min_row=chart_row + 1, max_row=last))
        ws.add_chart(bar2, "E" + str(chart_row + 16))

    interp_row = chart_row + 20
    section_header(ws, interp_row, "INTERPRETATION")
    write_cell(
        ws,
        interp_row + 1,
        1,
        "Expected-risk concentration follows AAL share. Tail-risk concentration follows portfolio TVaR 99 contributions — "
        "do not add standalone scenario TVaR figures to obtain aggregate TVaR.",
    )


def populate_impact(ws, result, domain: str):
    _clear_body(ws)
    _title(ws, "04 — Business Impact")
    _subtitle(
        ws,
        "Operational percentiles are not financial VaR. Downtime/capacity are annual aggregates "
        "across all simulated years (most years are zero). Recovery duration is shown only when "
        "separately modelled — OT currently reports downtime hours, not a distinct recovery clock.",
    )
    section_header(ws, 4, "OPERATIONAL CONSEQUENCE")
    op = result.get("operational_diagnostics") or {}
    labels = [
        ("Downtime P50 (days)", "downtime_p50"),
        ("Downtime P95 (days)", "downtime_p95"),
        ("Downtime P99 (days)", "downtime_p99"),
        ("Recovery duration P50 (days)", "recovery_p50"),
        ("Recovery duration P95 (days)", "recovery_p95"),
        ("Recovery duration P99 (days)", "recovery_p99"),
        ("Capacity affected P50", "capacity_p50"),
        ("Capacity affected P95", "capacity_p95"),
        ("Capacity affected P99", "capacity_p99"),
        ("Records affected P50", "records_p50"),
        ("Records affected P95", "records_p95"),
        ("Records affected P99", "records_p99"),
        ("Endpoints affected P50", "endpoints_p50"),
        ("Endpoints affected P95", "endpoints_p95"),
        ("Endpoints affected P99", "endpoints_p99"),
        ("Services affected P50", "services_p50"),
        ("Services affected P95", "services_p95"),
        ("Services affected P99", "services_p99"),
        ("Revenue at risk per day", "revenue_per_day"),
        ("P(downtime > 7 days)", "p_down_7"),
        ("P(downtime > 15 days)", "p_down_15"),
        ("P(downtime > 30 days)", "p_down_30"),
        ("P(downtime > 60 days)", "p_down_60"),
        ("P(capacity > 25%)", "p_cap_25"),
        ("P(capacity > 50%)", "p_cap_50"),
        ("P(capacity > 75%)", "p_cap_75"),
    ]
    write_cell(ws, 5, 1, "Metric", bold=True)
    write_cell(ws, 5, 2, "Value", bold=True)
    row = 6
    written = 0
    for lab, key in labels:
        val = op.get(key)
        if val is None:
            continue
        write_cell(ws, row, 1, lab)
        fmt = PCT if lab.startswith("P(") or "Capacity affected P" in lab else ("0.00" if "days" in lab.lower() else ("#,##0" if "affected P" in lab else CURRENCY))
        write_cell(ws, row, 2, val, fmt)
        row += 1
        written += 1
    if not written:
        write_cell(ws, 6, 1, "Operational diagnostics")
        write_cell(ws, 6, 2, "No applicable operational metrics for this sector/scenario pack.")

    section_header(ws, 34, "FINANCIAL LOSS CATEGORIES (Balbix major categories)")
    comps = result.get("loss_components") or []
    ch = ["Category", "AAL", "Portfolio TVaR 95 contribution", "Portfolio TVaR 99 contribution", "% of AAL", "% of TVaR 99"]
    for c, h in enumerate(ch, 1):
        write_cell(ws, 35, c, h, bold=True)
    if not comps:
        write_cell(ws, 36, 1, "No applicable category decomposition for this run.")
    for i, crow in enumerate(comps):
        write_cell(ws, 36 + i, 1, crow.get("name"))
        write_cell(ws, 36 + i, 2, crow.get("aal"), CURRENCY)
        write_cell(ws, 36 + i, 3, crow.get("contrib_tvar95"), CURRENCY)
        write_cell(ws, 36 + i, 4, crow.get("contrib_tvar99"), CURRENCY)
        write_cell(ws, 36 + i, 5, crow.get("pct_aal"), PCT)
        write_cell(ws, 36 + i, 6, crow.get("pct_tvar99"), PCT)

    section_header(ws, 48, "APPLICABLE IMPACT DRIVERS")
    drivers = result.get("loss_drivers") or []
    dh = ["Driver", "Category", "AAL", "Portfolio TVaR 99 contribution", "% of AAL", "% of TVaR 99"]
    for c, h in enumerate(dh, 1):
        write_cell(ws, 49, c, h, bold=True)
    if not drivers:
        write_cell(ws, 50, 1, "No applicable drivers for this sector pack / scenario selection.")
    for i, drow in enumerate(drivers[:30]):
        write_cell(ws, 50 + i, 1, drow.get("name"))
        write_cell(ws, 50 + i, 2, drow.get("category"))
        write_cell(ws, 50 + i, 3, drow.get("aal"), CURRENCY)
        write_cell(ws, 50 + i, 4, drow.get("contrib_tvar99"), CURRENCY)
        write_cell(ws, 50 + i, 5, drow.get("pct_aal"), PCT)
        write_cell(ws, 50 + i, 6, drow.get("pct_tvar99"), PCT)

    section_header(ws, 82, "Successful-Event Consequence Diagnostics")
    write_cell(
        ws,
        83,
        1,
        "These describe loss severity when an event occurs. They are not annual aggregate VaR or TVaR.",
    )
    write_cell(ws, 84, 1, "See IT CALC - Frequency and Success / OT impact drivers for event-level P50/P99 severity.")
    gaps = result.get("balbix_gaps") or []
    if gaps:
        section_header(ws, 86, "BALBIX CATALOGUE GAPS (flagged — not invented)")
        write_cell(ws, 87, 1, "Driver / topic", bold=True)
        write_cell(ws, 87, 2, "Gap", bold=True)
        for i, g in enumerate(gaps[:20]):
            if isinstance(g, dict):
                write_cell(ws, 88 + i, 1, g.get("id") or g.get("driver_id"))
                write_cell(ws, 88 + i, 2, g.get("reason") or g.get("gap"))
            else:
                write_cell(ws, 88 + i, 1, str(g))


def populate_treatment(ws, result):
    _clear_body(ws)
    _title(ws, "05 — Risk Treatment")
    _subtitle(ws, "Individual control uplifts and combined packages. Packages are re-simulated — reductions are not summed.")
    section_header(ws, 4, "INDIVIDUAL CONTROL IMPROVEMENTS")
    headers = [
        "Control", "Current", "Target", "Baseline P(event)", "Residual P(event)",
        "Baseline AAL", "Residual AAL", "Baseline VaR 95", "Residual VaR 95",
        "Baseline TVaR 95", "Residual TVaR 95", "Baseline VaR 99", "Residual VaR 99",
        "Baseline TVaR 99", "Residual TVaR 99", "AAL Δ%", "TVaR 99 Δ%",
        "Cost status", "Net benefit", "BCR", "Payback (years)",
    ]
    for c, h in enumerate(headers, 1):
        write_cell(ws, 5, c, h, bold=True)
    whatifs = result.get("whatifs") or []
    pick = "prudent" if str(result.get("reporting_view") or result.get("detailed_view") or "").lower().startswith("p") else "be"
    for i, w in enumerate(whatifs[:40]):
        be = w.get(pick) if isinstance(w.get(pick), dict) else (w.get("prudent") or w.get("be") or {})
        if not isinstance(be, dict):
            be = {}
        costs = w.get("costs") or {}
        write_cell(ws, 6 + i, 1, w.get("name"))
        write_cell(ws, 6 + i, 2, w.get("current"))
        write_cell(ws, 6 + i, 3, w.get("whatif") or w.get("next"))
        write_cell(ws, 6 + i, 4, be.get("freq_base"))
        write_cell(ws, 6 + i, 5, be.get("event_freq"))
        write_cell(ws, 6 + i, 6, be.get("baseline") or w.get("current_aal"), CURRENCY)
        write_cell(ws, 6 + i, 7, be.get("aal") or w.get("whatif_aal"), CURRENCY)
        write_cell(ws, 6 + i, 8, be.get("var95_base"), CURRENCY)
        write_cell(ws, 6 + i, 9, be.get("var95") or be.get("p95"), CURRENCY)
        write_cell(ws, 6 + i, 10, be.get("tvar95_base"), CURRENCY)
        write_cell(ws, 6 + i, 11, be.get("tvar95"), CURRENCY)
        write_cell(ws, 6 + i, 12, be.get("var99_base"), CURRENCY)
        write_cell(ws, 6 + i, 13, be.get("var99") or be.get("p99"), CURRENCY)
        write_cell(ws, 6 + i, 14, be.get("tvar99_base"), CURRENCY)
        write_cell(ws, 6 + i, 15, be.get("tvar99"), CURRENCY)
        write_cell(ws, 6 + i, 16, be.get("pct") or w.get("reduction_pct"), PCT)
        tbase = be.get("tvar99_base")
        tres = be.get("tvar99")
        tpct = ((tbase - tres) / tbase) if tbase and tres is not None else None
        write_cell(ws, 6 + i, 17, tpct, PCT)
        write_cell(ws, 6 + i, 18, costs.get("cost_status") or "Cost not provided")
        write_cell(ws, 6 + i, 19, costs.get("net_benefit"), CURRENCY)
        write_cell(ws, 6 + i, 20, costs.get("benefit_cost_ratio"))
        write_cell(ws, 6 + i, 21, costs.get("payback_years"))

    section_header(ws, 50, "COMBINED CONTROL PACKAGES (re-simulated)")
    ph = ["Package", "Controls", "Baseline AAL", "Package AAL", "VaR 95", "TVaR 95", "VaR 99", "TVaR 99", "AAL reduction", "TVaR 99 reduction", "Simulated?", "Cost status"]
    for c, h in enumerate(ph, 1):
        write_cell(ws, 51, c, h, bold=True)
    for i, p in enumerate(result.get("control_packages") or []):
        write_cell(ws, 52 + i, 1, p.get("name"))
        write_cell(ws, 52 + i, 2, p.get("n_controls"))
        write_cell(ws, 52 + i, 3, p.get("baseline_aal"), CURRENCY)
        write_cell(ws, 52 + i, 4, p.get("aal"), CURRENCY)
        write_cell(ws, 52 + i, 5, p.get("var95"), CURRENCY)
        write_cell(ws, 52 + i, 6, p.get("tvar95"), CURRENCY)
        write_cell(ws, 52 + i, 7, p.get("var99"), CURRENCY)
        write_cell(ws, 52 + i, 8, p.get("tvar99"), CURRENCY)
        write_cell(ws, 52 + i, 9, p.get("aal_reduction"), CURRENCY)
        write_cell(ws, 52 + i, 10, p.get("tvar99_reduction"), CURRENCY)
        write_cell(ws, 52 + i, 11, "Yes" if p.get("simulated") else "No")
        write_cell(ws, 52 + i, 12, (p.get("costs") or {}).get("cost_status") or "Cost not provided")
    if not result.get("control_packages"):
        write_cell(ws, 52, 1, "No packages simulated for this run.")


def populate_appetite_insurance(ws, result):
    # Preserve input block rows 1–22 seeded by ensure_insurance_inputs
    for r in range(24, 90):
        for c in range(1, 12):
            ws.cell(r, c).value = None
    if ws["A1"].value in (None, ""):
        _title(ws, "06 — Appetite & Insurance")
    section_header(ws, 24, "APPETITE STATUS (independent of insurance retention)")
    write_cell(ws, 25, 1, "Risk-appetite status", bold=True)
    write_cell(ws, 25, 2, result.get("appetite_status") or "Tolerance not set")
    write_cell(ws, 26, 1, "P(annual loss exceeds ANNUAL_LOSS_TOLERANCE)")
    write_cell(ws, 26, 2, result.get("prudent_p_exceed_tolerance") if result.get("prudent_p_exceed_tolerance") is not None else result.get("best_p_exceed_tolerance"), PCT)
    breaches = result.get("appetite_breaches") or {}
    write_cell(ws, 27, 1, "Breached thresholds")
    write_cell(ws, 27, 2, ", ".join(k for k, v in breaches.items() if v) or "None")

    section_header(ws, 29, "RISK-FINANCING ANALYSIS (not insurance pricing)")
    ins = result.get("insurance_analysis")
    if not ins:
        write_cell(ws, 30, 1, "No active insurance programme (set retention/layers above and re-run).")
        return
    write_cell(ws, 30, 1, "Metric", bold=True)
    write_cell(ws, 30, 2, "Ground-up", bold=True)
    write_cell(ws, 30, 3, "Residual after insurance", bold=True)
    gu = ins.get("metrics_ground_up") or {}
    res = ins.get("metrics_residual") or {}
    for i, (lab, key) in enumerate((("AAL", "AAL"), ("VaR 95", "VaR95"), ("TVaR 95", "TVaR95"), ("VaR 99", "VaR99"), ("TVaR 99", "TVaR99"))):
        write_cell(ws, 31 + i, 1, lab)
        write_cell(ws, 31 + i, 2, gu.get(key), CURRENCY)
        write_cell(ws, 31 + i, 3, res.get(key), CURRENCY)
    write_cell(ws, 37, 1, "P(retention exceeded)")
    write_cell(ws, 37, 2, ins.get("p_retention_exceeded"), PCT)
    write_cell(ws, 38, 1, "Programme exhaustion probability")
    write_cell(ws, 38, 2, ins.get("p_programme_exhaustion"), PCT)
    write_cell(ws, 39, 1, "Expected insured recovery")
    write_cell(ws, 39, 2, ins.get("expected_insured_recovery"), CURRENCY)
    write_cell(ws, 40, 1, "Expected retained / residual loss")
    write_cell(ws, 40, 2, ins.get("expected_retained_loss"), CURRENCY)
    write_cell(ws, 41, 1, "Expected uninsured above programme")
    write_cell(ws, 41, 2, ins.get("expected_uninsured_above_programme"), CURRENCY)
    write_cell(ws, 42, 1, "Trial reconcile (GU = residual + insured)")
    write_cell(ws, 42, 2, "PASS" if ins.get("reconcile_ok") else "FAIL")
    write_cell(ws, 44, 1, "Layer", bold=True)
    write_cell(ws, 44, 2, "P(attach)", bold=True)
    write_cell(ws, 44, 3, "P(exhaust)", bold=True)
    write_cell(ws, 44, 4, "Expected loss to layer", bold=True)
    write_cell(ws, 44, 5, "% of total expected recovery", bold=True)
    total_rec = float(ins.get("expected_insured_recovery") or 0)
    for i, name in enumerate((ins.get("p_layer_attaches") or {}).keys()):
        el = (ins.get("expected_loss_to_layer") or {}).get(name)
        write_cell(ws, 45 + i, 1, name)
        write_cell(ws, 45 + i, 2, (ins.get("p_layer_attaches") or {}).get(name), PCT)
        write_cell(ws, 45 + i, 3, (ins.get("p_layer_exhausts") or {}).get(name), PCT)
        write_cell(ws, 45 + i, 4, el, CURRENCY)
        write_cell(ws, 45 + i, 5, (None if el is None or not total_rec else float(el) / total_rec), PCT)


def populate_uncertainty(ws, result):
    _clear_body(ws)
    _title(ws, "07 — Uncertainty & Evidence")
    _subtitle(ws, "Best Estimate vs Prudent use identical metric definitions. Sensitivity rows are engine reruns where populated.")
    section_header(ws, 4, "BEST ESTIMATE → PRUDENT")
    write_cell(ws, 5, 1, "Metric", bold=True)
    write_cell(ws, 5, 2, "Best Estimate", bold=True)
    write_cell(ws, 5, 3, "Prudent", bold=True)
    write_cell(ws, 5, 4, "Difference", bold=True)
    for i, (lab, bk, pk) in enumerate((
        ("AAL", "best_aal", "prudent_aal"),
        ("VaR 95", "best_var95", "prudent_var95"),
        ("TVaR 95", "best_tvar95", "prudent_tvar95"),
        ("VaR 99", "best_var99", "prudent_var99"),
        ("TVaR 99", "best_tvar99", "prudent_tvar99"),
    )):
        be = result.get(bk)
        pr = result.get(pk)
        write_cell(ws, 6 + i, 1, lab)
        write_cell(ws, 6 + i, 2, be, CURRENCY)
        write_cell(ws, 6 + i, 3, pr, CURRENCY)
        write_cell(ws, 6 + i, 4, (None if be is None or pr is None else float(pr) - float(be)), CURRENCY)
    write_cell(ws, 12, 1, "Bridge limitation")
    write_cell(
        ws,
        12,
        2,
        "Prudence is applied as a combined frequency factor. Individual attribution by sequential prudence adjustments is not available in this engine.",
    )

    section_header(ws, 14, "SENSITIVITY (engine reruns)")
    sh = ["Parameter", "Kind", "Factor", "Years", "AAL", "VaR 95", "TVaR 95", "VaR 99", "TVaR 99", "Δ AAL", "Δ% AAL", "Limitation / detail"]
    for c, h in enumerate(sh, 1):
        write_cell(ws, 15, c, h, bold=True)
    for i, s in enumerate(result.get("sensitivities") or []):
        write_cell(ws, 16 + i, 1, s.get("parameter"))
        write_cell(ws, 16 + i, 2, s.get("kind"))
        write_cell(ws, 16 + i, 3, s.get("factor"))
        write_cell(ws, 16 + i, 4, s.get("simulation_years"))
        write_cell(ws, 16 + i, 5, s.get("aal"), CURRENCY)
        write_cell(ws, 16 + i, 6, s.get("var95"), CURRENCY)
        write_cell(ws, 16 + i, 7, s.get("tvar95"), CURRENCY)
        write_cell(ws, 16 + i, 8, s.get("var99"), CURRENCY)
        write_cell(ws, 16 + i, 9, s.get("tvar99"), CURRENCY)
        write_cell(ws, 16 + i, 10, s.get("delta_aal"), CURRENCY)
        write_cell(ws, 16 + i, 11, s.get("pct_aal"), PCT)
        detail = s.get("limitation") or s.get("rationale") or ""
        if s.get("shocked_actor_weights"):
            detail = f"weights={s['shocked_actor_weights']}"
        if s.get("route_field"):
            detail = f"{s.get('route_field')}: {s.get('baseline_route_value')!r}→{s.get('shocked_route_value')!r}"
        write_cell(ws, 16 + i, 12, detail)
    if not result.get("sensitivities"):
        write_cell(ws, 16, 1, "Sensitivity analysis not run")

    section_header(ws, 30, "EVIDENCE SUMMARY")
    eh = ["Parameter", "Current value/range", "Evidence source", "Interpretation", "Calibration method", "Best Estimate", "Prudence adjustment", "Limitation", "Evidence category", "Last review date"]
    for c, h in enumerate(eh, 1):
        write_cell(ws, 31, c, h, bold=True)
    evidence = result.get("evidence_summary") or [
        {
            "parameter": "Campaign frequency (λ)",
            "value": result.get("attempt_frequency") or result.get("best_attempt_frequency"),
            "source": "Sector pack + OT 06 / IT 07 assessment adjustments",
            "interpretation": "Material campaign arrival rate prior after modifiers",
            "calibration": "Pack baseline with facility overlays",
            "best": result.get("best_attempt_frequency"),
            "prudence": "× PRUDENCE_FACTOR / PRUDENCE_FREQUENCY_FACTOR",
            "limitation": "No opaque evidence score; grades live on pack sheets",
            "category": "Frequency",
            "reviewed": "",
        }
    ]
    for i, e in enumerate(evidence[:20]):
        write_cell(ws, 32 + i, 1, e.get("parameter"))
        write_cell(ws, 32 + i, 2, e.get("value"))
        write_cell(ws, 32 + i, 3, e.get("source"))
        write_cell(ws, 32 + i, 4, e.get("interpretation"))
        write_cell(ws, 32 + i, 5, e.get("calibration"))
        write_cell(ws, 32 + i, 6, e.get("best"))
        write_cell(ws, 32 + i, 7, e.get("prudence"))
        write_cell(ws, 32 + i, 8, e.get("limitation"))
        write_cell(ws, 32 + i, 9, e.get("category"))
        write_cell(ws, 32 + i, 10, e.get("reviewed"))


def populate_reporting_suite(wb, result: dict, meta: dict) -> None:
    """Ensure sheets exist and populate all seven views."""
    ensure_reporting_sheets(wb)
    ensure_appetite_inputs_on_run_setup(wb)
    ensure_insurance_inputs(wb)
    domain = str(meta.get("domain") or "").strip()
    view = str(meta.get("reporting_view") or "Prudent").strip()
    if not result.get("annual_revenue_at_risk"):
        result = dict(result)
        result["annual_revenue_at_risk"] = read_revenue(wb, domain)

    populate_executive(wb["01 Executive Risk Story"], result, view)
    populate_formation(wb["02 Risk Formation"], result, domain)
    populate_scenarios(wb["03 Scenario Analysis"], result)
    populate_impact(wb["04 Business Impact"], result, domain)
    populate_treatment(wb["05 Risk Treatment"], result)
    populate_appetite_insurance(wb["06 Appetite & Insurance"], result)
    populate_uncertainty(wb["07 Uncertainty & Evidence"], result)
