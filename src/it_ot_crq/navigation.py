"""Operator-facing tab order, colours, and fill-vs-ignore visibility."""

from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from crq.sheet_names import (
    COMMON_OUTSIDE_IN,
    COMMON_RUN_SETUP,
    DIV_ADVANCED,
    DIV_DASHBOARDS,
    DIV_FILL_IT,
    DIV_FILL_OT,
    DIV_NAV,
    DIV_OWNERSHIP,
    DIV_RESULTS,
    IT_FILL,
    OT_FILL,
)

COLOR_COMMON = "1F4E79"
COLOR_IT = "00B0F0"
COLOR_OT = "7030A0"
COLOR_DIVIDER = "7F7F7F"

IT_FILL_SHEETS = IT_FILL
OT_FILL_SHEETS = OT_FILL

SHARED_VISIBLE = (
    DIV_DASHBOARDS,
    "01 Executive Risk Story",
    "02 Risk Formation",
    "03 Scenario Analysis",
    "04 Business Impact",
    "05 Risk Treatment",
    "06 Appetite & Insurance",
    "07 Uncertainty & Evidence",
    "00 Dashboard",
    "00 Risk Drivers",
    "00 Loss Analysis",
    "00 What-If",
    "00 Risk Transfer",
    DIV_NAV,
    COMMON_RUN_SETUP,
    COMMON_OUTSIDE_IN,
    DIV_OWNERSHIP,
)

NAV_TAB_ORDER = [
    DIV_DASHBOARDS,
    "01 Executive Risk Story",
    "02 Risk Formation",
    "03 Scenario Analysis",
    "04 Business Impact",
    "05 Risk Treatment",
    "06 Appetite & Insurance",
    "07 Uncertainty & Evidence",
    "00 Dashboard",
    "00 Risk Drivers",
    "00 Loss Analysis",
    "00 What-If",
    "00 Risk Transfer",
    DIV_NAV,
    COMMON_RUN_SETUP,
    COMMON_OUTSIDE_IN,
    DIV_OWNERSHIP,
    DIV_FILL_IT,
    *IT_FILL,
    DIV_FILL_OT,
    *OT_FILL,
    DIV_RESULTS,
    "00 COMMON - Output Bridge",
    "OT 02 - Executive Summary",
    "IT 02 - Executive Summary",
    DIV_ADVANCED,
    "IT CORE - Model Guide",
    "OT CORE - Model Guide",
    "OT CORE - Control Method",
    "OT CORE - Scenario Definitions",
    "OT CORE - Scenario-TTP Map",
    "OT CORE - Actor-TTP Map",
    "OT CORE - Facility Feasibility",
    "OT CORE - TTP-Control Map",
    "OT CORE - Mapping Change Log",
    "OT PACK - Loaded Registry",
    "OT PACK - Loaded TTP Rationale",
    "OT PACK PG - Loaded Values",
    "OT PACK EA - Loaded Values",
    "OT PACK MF - Loaded Values",
    "OT PACK - Loaded Comparison",
    "IT 09 - Sector Pack Registry",
    "OT CALC - TTP Relevance",
    "OT CALC - Attack Path",
    "OT CALC - Threat Frequency",
    "OT CALC - Consequence",
    "OT CALC - Actor-Scenario Matrix",
    "OT CALC - Tail Risk Metrics",
    "OT CALC - Control What-If",
    "OT CALC - Aggregate LECs",
    "OT CALC - Scenario LECs",
    "OT CALC - Actor LECs",
    "OT CALC - Actor-Scenario LEC",
    "OT CALC - Risk Charts",
    "IT CALC - TTP Relevance",
    "IT CALC - Attack Path",
    "IT CALC - Frequency and Success",
    "IT CALC - Consequence",
    "IT CALC - Actor Scenario Matrix",
    "IT CALC - Tail Risk Metrics",
    "IT CALC - Control What-If",
    "IT CALC - Aggregate LECs",
    "IT CALC - Scenario LECs",
    "IT CALC - Actor LECs",
    "IT CALC - Risk Charts",
    "OT 26 - Sources and Evidence",
    "OT 27 - Validation Tests",
    "00 COMMON - OI Audit Trail",
    "IT 21 - Sources and Evidence",
    "IT 22 - Validation Tests",
]


def _tab_color(ws: Worksheet, hex_color: str) -> None:
    ws.sheet_properties.tabColor = hex_color


def _clear_merges(ws: Worksheet) -> None:
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def _ensure_divider(wb, title: str, body_lines: list[str]) -> None:
    ws = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
    ws.sheet_state = "visible"
    _clear_merges(ws)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=COLOR_DIVIDER)
    ws.merge_cells("A1:F1")
    for i, line in enumerate(body_lines, start=3):
        ws.cell(i, 1).value = line
        ws.cell(i, 1).alignment = Alignment(wrap_text=True)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    ws.column_dimensions["A"].width = 28
    ws.row_dimensions[3].height = 48
    _tab_color(ws, COLOR_DIVIDER)


def _reorder(wb) -> None:
    present = [name for name in NAV_TAB_ORDER if name in wb.sheetnames]
    extra = [name for name in wb.sheetnames if name not in present]
    for idx, name in enumerate(present + extra):
        wb.move_sheet(name, offset=idx - wb.sheetnames.index(name))


def _hide_unless(wb, allowed: set[str]) -> None:
    for name in wb.sheetnames:
        ws = wb[name]
        if name in {"00 COMMON - Engine Results", "00 Engine Results"}:
            ws.sheet_state = "visible"
            continue
        ws.sheet_state = "visible" if name in allowed else "hidden"


def apply_fill_banners(wb) -> None:
    it_msg = (
        '=IF(\'00 COMMON - Run Setup\'!C6="IT",'
        '"FILL this sheet — IT assessment.",'
        '"Not used for this run. Complete the OT assessment section only.")'
    )
    ot_msg = (
        '=IF(\'00 COMMON - Run Setup\'!C6="OT",'
        '"FILL this sheet — OT assessment.",'
        '"Not used for this run. Complete the IT assessment section only.")'
    )
    for name in IT_FILL:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        ws["A3"] = it_msg
        ws["A3"].font = Font(bold=True, color="FFFFFF")
        ws["A3"].fill = PatternFill("solid", fgColor=COLOR_IT)
        _tab_color(ws, COLOR_IT)
    for name in OT_FILL:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        ws["A3"] = ot_msg
        ws["A3"].font = Font(bold=True, color="FFFFFF")
        ws["A3"].fill = PatternFill("solid", fgColor=COLOR_OT)
        _tab_color(ws, COLOR_OT)


def apply_run_setup_fill_guide(wb) -> None:
    setup = wb[COMMON_RUN_SETUP]
    setup["A50"] = "Required workflow"
    setup["A51"] = (
        '=IF(C6="IT","Complete the IT input section only",'
        'IF(C6="OT","Complete the OT input section only","Select IT or OT"))'
    )
    setup["A53"] = "Selected domain"
    setup["B53"] = "=C6"
    setup["A54"] = "Selected sector"
    setup["B54"] = "=C7"
    setup["A55"] = "Unit of analysis"
    setup["B55"] = "=C13"


def apply_run_setup_asset_dropdowns(wb) -> None:
    """Kept for tests; domain-first lists are applied by rebuild."""
    apply_run_setup_fill_guide(wb)
    apply_fill_banners(wb)


def apply_workbook_navigation(wb, domain: str | None = None) -> None:
    _ensure_divider(
        wb, DIV_NAV,
        [
            '=IF(\'00 COMMON - Run Setup\'!C6="IT","Complete the IT input section only",'
            'IF(\'00 COMMON - Run Setup\'!C6="OT","Complete the OT input section only","Set Model domain first"))',
            "Common every run: 00 COMMON - Run Setup; 00 COMMON - Outside-In (optional).",
            "IT only: IT 03, optional IT 04, IT 05, IT 06 overrides, IT 07 adjustments, IT 08 reporting.",
            "OT only: OT 03, OT 04, OT 05, OT 06 adjustments.",
            "Governed packs and core methodology are Advanced / Read Only / Not required to complete an assessment.",
        ],
    )
    _ensure_divider(
        wb, DIV_FILL_IT,
        [
            '=IF(\'00 COMMON - Run Setup\'!C6="IT","FILL the cyan IT tabs. Ignore FILL OT.","Not used for this run.")',
        ],
    )
    _ensure_divider(
        wb, DIV_FILL_OT,
        [
            '=IF(\'00 COMMON - Run Setup\'!C6="OT","FILL the purple OT tabs. Ignore FILL IT.","Not used for this run.")',
        ],
    )
    for title, lines in (
        (DIV_DASHBOARDS, ["Read after python -m crq run. Do not enter assessment data here."]),
        (DIV_RESULTS, ["Results and audit trail. Calculated output — do not edit."]),
        (DIV_ADVANCED, ["Advanced / Governed. Read Only. Not required to complete an assessment."]),
    ):
        if title in wb.sheetnames or True:
            _ensure_divider(wb, title, lines)
    apply_fill_banners(wb)
    apply_run_setup_fill_guide(wb)
    apply_it_input_ux(wb)
    _reorder(wb)
    allowed = set(SHARED_VISIBLE) | {DIV_FILL_IT, DIV_FILL_OT, *IT_FILL, *OT_FILL}
    if domain == "IT":
        allowed = set(SHARED_VISIBLE) | {DIV_FILL_IT, *IT_FILL, "IT 02 - Executive Summary", "00 COMMON - Output Bridge"}
    elif domain == "OT":
        allowed = set(SHARED_VISIBLE) | {DIV_FILL_OT, *OT_FILL, "OT 02 - Executive Summary", "00 COMMON - Output Bridge"}
    if "00 COMMON - Engine Results" in wb.sheetnames:
        allowed.add("00 COMMON - Engine Results")
    _hide_unless(wb, allowed)


def apply_it_input_ux(wb) -> None:
    """Link IT 06 scale to IT 03 and restore assessment dropdowns lost in the combined rename."""
    org = "IT 03 - Organisation Inputs"
    bia = "IT 06 - Impact & BIA Overrides"
    if org in wb.sheetnames:
        ws = wb[org]
        ws["A4"] = (
            "Enter organisation identity and scale here. Python uses these Value cells for the IT calculation. "
            "Sector is taken from Run Setup. Currency and Use exposure model are dropdowns."
        )
        ws["C19"] = "='00 COMMON - Run Setup'!C7"
        _replace_list_validations(
            ws,
            [
                ("C20", '"USD,GBP,EUR,AED"'),
                ("C29", '"Yes,No"'),
            ],
        )
    if bia in wb.sheetnames:
        ws = wb[bia]
        ws["A4"] = (
            "Organisation-scale figures below are live links to IT 03 — do not type them here. "
            "On this sheet, fill Override P50/P99 only where you have organisation-specific BIA evidence. "
            "Leave overrides blank to keep the Financial Services pack baseline."
        )
        links = [
            (7, "='IT 03 - Organisation Inputs'!C21"),
            (8, "='IT 03 - Organisation Inputs'!C22"),
            (9, "='IT 03 - Organisation Inputs'!C23"),
            (10, "='IT 03 - Organisation Inputs'!C25"),
            (11, "='IT 03 - Organisation Inputs'!C26"),
            (12, "='IT 03 - Organisation Inputs'!C27"),
            (13, "='IT 03 - Organisation Inputs'!C28"),
        ]
        for row, formula in links:
            ws.cell(row, 2).value = formula
            ws.cell(row, 3).value = "IT 03 - Organisation Inputs (live)"
    if "IT 04 - Exposure Adjustments" in wb.sheetnames:
        ws = wb["IT 04 - Exposure Adjustments"]
        ws["A4"] = (
            "Optional. Leave this sheet at Yes / Unknown / 1.0 unless you set USE_EXPOSURE_MODEL = Yes on IT 03 "
            "and have a governed exposure-model recommendation. Organisation applicable? can still mark a route as out of scope."
        )
        _replace_list_validations(
            ws,
            [
                ("C16:D21", '"Yes,No,Unknown"'),
                ("L16:L21", '"High,Medium,Low,Not assessed"'),
            ],
        )
    if "IT 05 - Control Assessment" in wb.sheetnames:
        _replace_list_validations(
            wb["IT 05 - Control Assessment"],
            [("E16:E115", '"Absent,Initial,Developing,Managed,Optimised,Not Assessed"')],
        )
    if "IT 08 - Reporting Settings" in wb.sheetnames:
        ws = wb["IT 08 - Reporting Settings"]
        ws["A4"] = (
            "Set reporting basis (Best Estimate vs Prudent), TVaR 95 vs TVaR 99, and illustrative insurance retention here. "
            "00 Risk Transfer, the executive dashboard Basis, and the IT engine all read these three cells."
        )
        current = str(ws["B16"].value or "").strip()
        if current.startswith("=") or current not in {"Best Estimate", "Prudent", "Both"}:
            setup_view = "Prudent"
            if COMMON_RUN_SETUP in wb.sheetnames:
                setup_view = str(wb[COMMON_RUN_SETUP]["C9"].value or "Prudent").strip()
                if setup_view.startswith("="):
                    setup_view = "Prudent"
            ws["B16"] = setup_view if setup_view in {"Best Estimate", "Prudent", "Both"} else "Prudent"
        _replace_list_validations(
            ws,
            [
                ("B16", '"Best Estimate,Prudent"'),
                ("B17", '"TVaR 95,TVaR 99"'),
            ],
        )


def _replace_list_validations(ws, specs: list[tuple[str, str]]) -> None:
    ws.data_validations.dataValidation.clear()
    for sqref, formula in specs:
        dv = DataValidation(
            type="list",
            formula1=formula,
            allow_blank=False,
            showDropDown=False,
            showErrorMessage=True,
        )
        dv.error = "Select a value from the list."
        dv.errorTitle = "Invalid selection"
        ws.add_data_validation(dv)
        dv.add(sqref)


def visible_sheet_names(wb) -> list[str]:
    return [name for name in wb.sheetnames if wb[name].sheet_state == "visible"]
