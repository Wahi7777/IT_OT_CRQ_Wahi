"""Minimal openpyxl-backed shim for the Guided OT-CRQ companion engine."""
import openpyxl
from openpyxl.utils import range_boundaries


class Blob:
    def __init__(self, path):
        self.path = path

    @staticmethod
    def load(path):
        return Blob(path)


class _Range:
    def __init__(self, ws, ref):
        self.ws = ws
        self.min_col, self.min_row, self.max_col, self.max_row = range_boundaries(ref)

    @property
    def values(self):
        out = []
        for r in range(self.min_row, self.max_row + 1):
            out.append([self.ws.cell(r, c).value
                        for c in range(self.min_col, self.max_col + 1)])
        return out

    @values.setter
    def values(self, data):
        for i, row in enumerate(data):
            for j, v in enumerate(row):
                self.ws.cell(self.min_row + i, self.min_col + j, v)

    def clear(self, _opts=None):
        for r in range(self.min_row, self.max_row + 1):
            for c in range(self.min_col, self.max_col + 1):
                cell = self.ws.cell(r, c)
                if cell.__class__.__name__ != "MergedCell":
                    cell.value = None


class _Sheet:
    def __init__(self, ws):
        self.ws = ws

    def get_range(self, ref):
        return _Range(self.ws, ref)


class _Worksheets:
    def __init__(self, wb):
        self.wb = wb

    def get_item(self, name):
        return _Sheet(self.wb[name])


class _Exported:
    def __init__(self, wb):
        self.wb = wb

    def save(self, path):
        self.wb.save(path)


class SpreadsheetFile:
    def __init__(self, wb):
        self.wb = wb
        self.worksheets = _Worksheets(wb)

    @staticmethod
    def import_xlsx(blob):
        return SpreadsheetFile(openpyxl.load_workbook(blob.path))

    @staticmethod
    def export_xlsx(sf):
        return _Exported(sf.wb)
