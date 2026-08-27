from __future__ import annotations

import math
from collections import defaultdict
from copy import copy
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.workbook.properties import CalcProperties

from crq.sheet_theme import ROLE_CALC, apply_column_headers, apply_freeze_at_headers
from crq.appetite import evaluate_appetite, parse_appetite_inputs
from crq.insurance import apply_programme, parse_programme
from crq.metrics import (
    LABEL_AAL,
    LABEL_TVAR_95,
    LABEL_TVAR_99,
    LABEL_VAR_95,
    LABEL_VAR_99,
    VAR_QUANTILE_METHOD,
    annual_aggregate_metrics,
    expected_shortfall as _shared_es,
    exceedance_probability,
    executive_risk_narrative,
    multi_year_event_probability,
    normalize_tail_basis,
    tvar_tail_contributions,
)
from crq.impact import (
    assert_pack_drivers_are_balbix,
    category_contribution_rows,
    driver_contribution_rows,
    formulas as impact_formulas,
    reconcile_impact_hierarchy,
)
from crq.reporting_helpers import (
    leading_tvar99_component,
    metric_bundle_from_sim,
    package_maturity_overrides,
    treatment_cost_metrics,
)

ENGINE_VERSION = "1.1.1"
PACK_SCHEMA_VERSION = "CRQ-PACK-1.1"
STAGES = ["S1", "S2", "S3", "S4", "S5"]
RETURN_PERIODS = [2, 5, 10, 20, 25, 50, 100, 200, 250, 500]
GRADES = {"A", "B", "C", "D"}
EFFECTS = {"PREVENT_SUCCESS", "CONTAIN_SUCCESS", "REDUCE_DURATION", "REDUCE_EXPOSURE", "REDUCE_CONSEQUENCE"}
ACTIVITY = {"Nation-state": "NATION_STATE_ACTIVITY", "Cybercriminal": "CYBERCRIMINAL_ACTIVITY", "Malicious insider": "MALICIOUS_INSIDER_ACTIVITY", "Hacktivist": "HACKTIVIST_ACTIVITY"}
MATURITY = {"Absent": 0.0, "Initial": .25, "Developing": .5, "Managed": .75, "Optimised": .95, "Not Assessed": .5}
MATURITY_ORDER = ["Absent", "Initial", "Developing", "Managed", "Optimised"]


def _num(v, label, lo=None, hi=None):
    if v in (None, ""):
        raise ValueError(f"{label} must be populated.")
    try:
        x = float(v)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label} must be numeric; got {v!r}.") from exc
    if not math.isfinite(x) or (lo is not None and x < lo) or (hi is not None and x > hi):
        raise ValueError(f"{label} is outside the permitted range.")
    return x


def _txt(v, label):
    x = str(v or "").strip()
    if not x:
        raise ValueError(f"{label} must be populated.")
    return x


def _version_tuple(value):
    parts = str(value).strip().split(".")
    if not parts or any(not part.isdigit() for part in parts):
        raise ValueError(f"Invalid version number: {value!r}.")
    return tuple(int(part) for part in parts)


def _kv(ws, start, end, key_col=1, val_col=3):
    return {str(ws.cell(r, key_col).value).strip(): ws.cell(r, val_col).value for r in range(start, end + 1) if ws.cell(r, key_col).value not in (None, "")}


def _table(ws, header=15):
    heads = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(header, c).value
        if v in (None, ""):
            if heads:
                break
        else:
            heads.append(str(v).strip())
    if not heads:
        raise ValueError(f"No table header found on {ws.title}.")
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(heads) + 1)]
        if not all(v in (None, "") for v in vals):
            rows.append(dict(zip(heads, vals)))
    return rows


def _sum1(values, label):
    total = float(sum(values))
    if abs(total - 1) > 1e-6:
        raise ValueError(f"{label} must sum to 1.0; got {total}.")


def _grade(row, label):
    g = str(row.get("Evidence grade") or "").strip().upper()
    if g not in GRADES:
        raise ValueError(f"{label} evidence grade must be A, B, C or D.")
    return g


def _expected_shortfall(arr, q):
    """Exact upper-tail mean, including fractional mass at the quantile cut."""
    return _shared_es(arr, q)


def _metrics(arr):
    return annual_aggregate_metrics(arr)


def _reporting_view(value):
    view = str(value or "").strip()
    if view in {"Best Estimate", "Prudent"}:
        return view
    if view == "Both":
        return "Prudent"
    raise ValueError("REPORTING_VIEW must be Best Estimate, Prudent or Both.")


def _lognormal(p50, p99):
    p50 = max(float(p50), 1.0); p99 = max(float(p99), p50 * 1.000001)
    mu = math.log(p50); sigma = (math.log(p99) - mu) / 2.326347874
    return mu, sigma, math.exp(mu + .5 * sigma * sigma)


def _registry(wb, model_path, pack_dir):
    from crq.pack_registry import project_root_from, resolve_pack, validate_pack_file

    sector = _txt(_kv(wb["03 Organisation Inputs"], 16, 40).get("SECTOR"), "SECTOR")
    root = project_root_from(Path(model_path))
    rec = resolve_pack("IT", sector, root)
    validate_pack_file(rec, root, ENGINE_VERSION, "IT")
    path = rec.path(root)
    if pack_dir:
        alt = Path(pack_dir) / Path(rec.relative_file_path).name
        if alt.is_file():
            path = alt.resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Sector pack not found: {path}")
    return sector, {"Pack ID": rec.pack_id, "Relative file path": rec.relative_file_path, "Status": rec.pack_status}, path


def _load_pack(main_wb, model_path, pack_dir=None):
    sector, reg, path = _registry(main_wb, model_path, pack_dir)
    wb = openpyxl.load_workbook(path, data_only=False)
    req = ["01 Pack Metadata","02 Actor Weights","03 Actor Scenario","04 Route Definitions","05 Actor Route","06 Scenario Route","07 Stage Requirements","08 Stage Priors","09 Scenario Parameters","10 Driver Rates","11 Impact Driver Matrix","12 TTP Catalogue","13 Route TTP Map","14 Control Reference","15 Control TTP Map","16 Sources","17 Evidence Register"]
    missing = [s for s in req if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"Sector pack missing required sheets: {missing}")
    meta = _kv(wb["01 Pack Metadata"], 16, 60, 1, 2)
    if _txt(meta.get("SECTOR"), "pack sector") != sector or _txt(meta.get("PACK_ID"), "pack ID") != _txt(reg.get("Pack ID"), "registry pack ID"):
        raise ValueError("Pack metadata does not match registry.")
    if _txt(meta.get("SCHEMA_VERSION"), "schema") != PACK_SCHEMA_VERSION or _version_tuple(meta.get("MIN_ENGINE_VERSION")) > _version_tuple(ENGINE_VERSION):
        raise ValueError("Incompatible pack schema or engine version.")

    actors = []
    for row in _table(wb["02 Actor Weights"]):
        aid = str(row.get("Actor ID") or "").strip()
        if aid and aid != "CHECK":
            actors.append({"id": aid, "name": _txt(row.get("Threat actor"), aid), "weight": _num(row.get("Weight"), f"{aid} weight", 0), "grade": _grade(row, aid)})
    _sum1([a["weight"] for a in actors], "Actor weights")
    aids = [a["id"] for a in actors]
    as_rows = _table(wb["03 Actor Scenario"]); sids = [h for h in as_rows[0] if h.startswith("SC_")]
    actor_scenario = {}
    for row in as_rows:
        aid = str(row.get("Actor ID") or "").strip()
        if aid in aids:
            actor_scenario[aid] = {sid: _num(row.get(sid), f"{aid}/{sid}", 0) for sid in sids}; _sum1(actor_scenario[aid].values(), f"{aid} scenarios")
    routes = []
    for row in _table(wb["04 Route Definitions"]):
        rid = _txt(row.get("Route ID"), "route ID")
        routes.append({"id": rid, "name": _txt(row.get("Attack route"), rid), "grade": _grade(row, rid)})
    rids = [r["id"] for r in routes]

    def matrix(sheet, key_name, cols):
        out = {}
        for row in _table(wb[sheet]):
            key = str(row.get(key_name) or "").strip()
            if key:
                out[key] = {c: _num(row.get(c), f"{sheet}/{key}/{c}", 0) for c in cols}; _sum1(out[key].values(), f"{sheet}/{key}")
        return out
    actor_route = matrix("05 Actor Route", "Actor ID", rids); scenario_route = matrix("06 Scenario Route", "Scenario ID", rids)
    if set(actor_route) != set(aids) or set(scenario_route) != set(sids):
        raise ValueError("Route propensity matrices are incomplete.")
    stage_req = {}
    for row in _table(wb["07 Stage Requirements"]):
        key = (_txt(row.get("Actor ID"), "actor"), _txt(row.get("Scenario ID"), "scenario"), _txt(row.get("Route ID"), "route"))
        stage_req[key] = {s: bool(int(_num(row.get(s), f"{key}/{s}", 0, 1))) for s in STAGES}
    priors = {}
    for row in _table(wb["08 Stage Priors"]):
        rid = str(row.get("Route ID") or "").strip()
        if rid: priors[rid] = {s: _num(row.get(s), f"{rid}/{s}", 0, 1) for s in STAGES}
    names, params, param_ev = {}, {}, {}
    for row in _table(wb["09 Scenario Parameters"]):
        sid = str(row.get("Scenario ID") or "").strip()
        if sid in sids:
            names[sid] = _txt(row.get("Scenario"), sid); param_ev[sid] = (_grade(row, sid), row.get("Source ID"))
            params[sid] = {k: float(v) for k,v in row.items() if k not in {"Scenario ID","Scenario","Evidence grade","Source ID"} and v not in (None,"")}
    rates = {}
    for row in _table(wb["10 Driver Rates"]):
        rid = str(row.get("Rate ID") or "").strip()
        if rid: rates[rid] = {"name":row.get("Driver"),"unit":row.get("Unit"),"p50":_num(row.get("P50"),f"{rid} P50",0),"p99":_num(row.get("P99"),f"{rid} P99",0),"grade":_grade(row,rid),"source":row.get("Source ID")}
    impact = {}
    anames = [a["name"] for a in actors]
    for row in _table(wb["11 Impact Driver Matrix"]):
        did = str(row.get("Driver ID") or "").strip()
        if did:
            block = _txt(row.get("Loss block"), f"{did} block")
            if block not in {"Duration","Exposure","Direct"}: raise ValueError(f"Invalid block for {did}.")
            impact[did] = {"block":block,"actor":{a:bool(int(_num(row.get(a),f"{did}/{a}",0,1))) for a in anames},"scenario":{s:bool(int(_num(row.get(s),f"{did}/{s}",0,1))) for s in sids},"grade":_grade(row,did),"source":row.get("Source ID")}
    ttp = {str(r.get("TTP ID")):r for r in _table(wb["12 TTP Catalogue"]) if r.get("TTP ID")}
    route_ttp = defaultdict(list)
    for row in _table(wb["13 Route TTP Map"]):
        rid,tid = str(row.get("Route ID") or "").strip(),str(row.get("TTP ID") or "").strip()
        if rid not in rids or tid not in ttp: raise ValueError(f"Invalid route/TTP {rid}/{tid}.")
        route_ttp[rid].append((tid,_num(row.get("Relevance weight"),f"{rid}/{tid}",0)))
    control_ref = {}
    for row in _table(wb["14 Control Reference"]):
        cid = str(row.get("Control ID") or "").strip()
        if cid:
            maturity = _txt(row.get("Reference maturity"),cid)
            if maturity not in MATURITY: raise ValueError(f"Unknown maturity {maturity}.")
            control_ref[cid] = {**row,"maturity":maturity,"coverage":_num(row.get("Reference coverage"),f"{cid} coverage",0,1),"grade":_grade(row,cid)}
    control_ttp = defaultdict(list)
    for row in _table(wb["15 Control TTP Map"]):
        cid,tid = str(row.get("Control ID") or "").strip(),str(row.get("TTP ID") or "").strip(); effect = _txt(row.get("Effect"),"effect")
        if cid not in control_ref or tid not in ttp or effect not in EFFECTS: raise ValueError(f"Invalid control/TTP {cid}/{tid}/{effect}.")
        _grade(row,f"{cid}/{tid}"); control_ttp[tid].append((cid,_num(row.get("Base efficacy"),"efficacy",0,1),effect))
    settings = {k:_num(meta.get(k),k) for k in ["CONTROL_AGG_INCREMENT","CONTROL_ADJ_FLOOR","CONTROL_ADJ_CEILING","DURATION_FACTOR_FLOOR","EXPOSURE_FACTOR_FLOOR","CONSEQUENCE_FACTOR_FLOOR","CONSEQUENCE_FACTOR_CEILING","LOSS_BLOCK_RHO"]}
    settings["CONTROL_BARRIER_CAP"] = _num(meta.get("CONTROL_BARRIER_CAP", .95), "CONTROL_BARRIER_CAP", 0, .95)
    return {"path":path,"meta":meta,"settings":settings,"actors":actors,"actor_scenario":actor_scenario,"scenarios":sids,"scenario_names":names,"scenario_params":params,"scenario_evidence":param_ev,"routes":routes,"actor_route":actor_route,"scenario_route":scenario_route,"stage_requirements":stage_req,"stage_priors":priors,"driver_rates":rates,"impact_matrix":impact,"ttp":ttp,"route_ttp":route_ttp,"control_ref":control_ref,"control_ttp":control_ttp}


def _style_row(ws, source, target, cols):
    for c in range(1, cols + 1):
        if ws.cell(source,c).has_style: ws.cell(target,c)._style = copy(ws.cell(source,c)._style)


def _save_safe(wb, output_path):
    output_path=Path(output_path).resolve(); tmp=output_path.with_name(output_path.stem+".__writing__.xlsx")
    wb.save(tmp); tmp.replace(output_path)


def _prepare(wb, pack):
    ws=wb["04 Exposure Adjustments"]
    old={str(ws.cell(r,1).value):[ws.cell(r,c).value for c in range(3,16)] for r in range(16,ws.max_row+1) if ws.cell(r,1).value}
    for r in range(16,116):
        for c in range(1,23): ws.cell(r,c).value=None
    default=["Yes","Unknown",1,1,1,1,1,1,1,"Not assessed","","","No exposure-model recommendation loaded."]
    for r,route in enumerate(pack["routes"],16):
        _style_row(ws,16,r,22)
        for c,v in enumerate([route["id"],route["name"],*old.get(route["id"],default)],1): ws.cell(r,c,v)
    ws=wb["05 Control Assessment"]
    old={str(ws.cell(r,1).value):(ws.cell(r,5).value,ws.cell(r,6).value,ws.cell(r,9).value) for r in range(16,ws.max_row+1) if ws.cell(r,1).value}
    for r in range(16,116):
        for c in range(1,11): ws.cell(r,c).value=None
    for r,(cid,ref) in enumerate(pack["control_ref"].items(),16):
        _style_row(ws,16,r,10); actual=old.get(cid,(ref["maturity"],ref["coverage"],"User assessment required"))
        vals=[cid,ref.get("Control family"),ref.get("NIST CSF 2.0"),ref.get("MITRE mitigations"),actual[0],actual[1],ref["maturity"],ref["coverage"],actual[2],ref.get("Status")]
        for c,v in enumerate(vals,1): ws.cell(r,c,v)
    ws=wb["06 Impact and BIA"]
    old={}
    for r in range(17,251):
        typ,sid,pid=str(ws.cell(r,1).value or ""),str(ws.cell(r,2).value or ""),str(ws.cell(r,4).value or "")
        if typ in {"RATE","PARAM"} and pid: old[(typ,sid,pid)]=(ws.cell(r,9).value,ws.cell(r,10).value)
    for r in range(17,251):
        for c in range(1,15): ws.cell(r,c).value=None
    rows=[]
    for rid,x in pack["driver_rates"].items():
        o50,o99=old.get(("RATE","",rid),(None,None)); rows.append(["RATE","","",rid,x["name"],x["unit"],x["p50"],x["p99"],o50,o99,x["p50"] if o50 in (None,"") else o50,x["p99"] if o99 in (None,"") else o99,x["grade"],x["source"]])
    for sid in pack["scenarios"]:
        keys=sorted(k[:-4] for k in pack["scenario_params"][sid] if k.endswith("_P50"))
        for key in keys:
            o50,o99=old.get(("PARAM",sid,key),(None,None)); p50=pack["scenario_params"][sid][key+"_P50"]; p99=pack["scenario_params"][sid][key+"_P99"]; g,src=pack["scenario_evidence"][sid]
            unit="days" if key.endswith("DAYS") else ("USD" if key.endswith("DIRECT") else "share")
            rows.append(["PARAM",sid,pack["scenario_names"][sid],key,key,unit,p50,p99,o50,o99,p50 if o50 in (None,"") else o50,p99 if o99 in (None,"") else o99,g,src])
    for r,vals in enumerate(rows,17):
        _style_row(ws,17,r,14)
        for c,v in enumerate(vals,1): ws.cell(r,c,v)


def prepare(input_path, output_path=None, sector_pack_dir=None):
    input_path=Path(input_path).resolve(); output_path=Path(output_path or input_path).resolve(); wb=openpyxl.load_workbook(input_path,data_only=False); pack=_load_pack(wb,input_path,sector_pack_dir); _prepare(wb,pack); _save_safe(wb,output_path)
    return {"output":str(output_path),"pack_id":pack["meta"]["PACK_ID"],"routes":len(pack["routes"]),"controls":len(pack["control_ref"])}


def _load_main(wb,pack):
    org=_kv(wb["03 Organisation Inputs"],16,40); freq=_kv(wb["07 Frequency Assumptions"],16,22,1,2); tail=_kv(wb["08 Tail Risk Settings"],16,19,1,2)
    use=_txt(org.get("USE_EXPOSURE_MODEL"),"USE_EXPOSURE_MODEL")
    if use not in {"Yes","No"}: raise ValueError("USE_EXPOSURE_MODEL must be Yes or No.")
    ws=wb["04 Exposure Adjustments"]; exposure={}
    for r in range(16,116):
        rid=str(ws.cell(r,1).value or "").strip()
        if rid: exposure[rid]={"org_applicable":ws.cell(r,3).value,"model_feasible":ws.cell(r,4).value,"opportunity":_num(ws.cell(r,5).value,f"{rid} opportunity",0),"stage":{s:_num(ws.cell(r,6+i).value,f"{rid}/{s}",0) for i,s in enumerate(STAGES)},"impact":_num(ws.cell(r,11).value,f"{rid} impact",0)}
    if set(exposure)!={r["id"] for r in pack["routes"]}: raise ValueError("Exposure routes do not match pack; run prepare.")
    ws=wb["05 Control Assessment"]; controls={}
    for r in range(16,116):
        cid=str(ws.cell(r,1).value or "").strip()
        if cid:
            maturity=_txt(ws.cell(r,5).value,f"{cid} maturity")
            if maturity not in MATURITY: raise ValueError(f"Unknown maturity {maturity}.")
            controls[cid]={"name":ws.cell(r,2).value,"family":pack["control_ref"][cid].get("Control family") or ws.cell(r,2).value,"maturity":maturity,"coverage":_num(ws.cell(r,6).value,f"{cid} coverage",0,1)}
    if set(controls)!=set(pack["control_ref"]): raise ValueError("Controls do not match pack; run prepare.")
    rates={k:{**v,"used_p50":v["p50"],"used_p99":v["p99"]} for k,v in pack["driver_rates"].items()}; params={s:dict(v) for s,v in pack["scenario_params"].items()}; rows={}; ws=wb["06 Impact and BIA"]
    for r in range(17,251):
        typ,sid,pid=str(ws.cell(r,1).value or ""),str(ws.cell(r,2).value or ""),str(ws.cell(r,4).value or "")
        if typ not in {"RATE","PARAM"} or not pid: continue
        o50,o99=ws.cell(r,9).value,ws.cell(r,10).value
        if typ=="RATE": rates[pid]["used_p50"]=rates[pid]["p50"] if o50 in (None,"") else _num(o50,f"{pid} override",0); rates[pid]["used_p99"]=rates[pid]["p99"] if o99 in (None,"") else _num(o99,f"{pid} override",0)
        else: params[sid][pid+"_P50"]=params[sid][pid+"_P50"] if o50 in (None,"") else _num(o50,f"{sid}/{pid}",0); params[sid][pid+"_P99"]=params[sid][pid+"_P99"] if o99 in (None,"") else _num(o99,f"{sid}/{pid}",0)
        rows[(typ,sid,pid)]=r
    return {"org":org,"freq":freq,"tail":tail,"use_exposure":use,"exposure":exposure,"controls":controls,"rates":rates,"params":params,"assumption_rows":rows}


def _agg(values,increment,cap=.95):
    vals=sorted((min(max(v,0),cap) for v in values),reverse=True)
    if not vals:
        return 0.0
    barrier=vals[0]
    for i,efficacy in enumerate(vals[1:],start=1):
        barrier += (1-barrier)*(increment**i)*efficacy
    return min(barrier,cap)


def _barrier(tid,effects,controls,pack):
    actual=[]; reference=[]
    for cid,efficacy,effect in pack["control_ttp"].get(tid,[]):
        if effect in effects:
            actual.append(efficacy*MATURITY[controls[cid]["maturity"]]*controls[cid]["coverage"]); reference.append(efficacy*MATURITY[pack["control_ref"][cid]["maturity"]]*pack["control_ref"][cid]["coverage"])
    inc=pack["settings"]["CONTROL_AGG_INCREMENT"]
    cap=pack["settings"]["CONTROL_BARRIER_CAP"]
    return _agg(actual,inc,cap),_agg(reference,inc,cap)


def _channel(tid,effect,controls,pack,lo,hi):
    a,r=_barrier(tid,{effect},controls,pack)
    return min(max((1-a)/max(1-r,1e-9),lo),hi)


def _path(actor,sid,route,main,pack,override=None,collect=False):
    controls={k:dict(v) for k,v in main["controls"].items()}
    if override:
        if isinstance(override, dict):
            for cid, maturity in override.items():
                if cid in controls:
                    controls[cid]["maturity"]=maturity
        else:
            controls[override[0]]["maturity"]=override[1]
    aid,rid=actor["id"],route["id"]; required=pack["stage_requirements"].get((aid,sid,rid),pack["stage_requirements"].get(("ALL",sid,rid)))
    if required is None: return {"valid":False,"success":0,"opportunity":0,"impact_factor":1,"duration_factor":1,"exposure_factor":1,"consequence_factor":1,"stages":{},"ttp_rows":[]}
    exp=main["exposure"][rid]; use=main["use_exposure"]=="Yes"; feasible=str(exp["org_applicable"] or "").lower()!="no" and (not use or str(exp["model_feasible"] or "").lower()!="no")
    valid=feasible and any(required.values()); stages={}; audit=[]; all_rel=[]
    for stage in STAGES:
        if not required[stage]: stages[stage]={"required":False,"reference":1,"actual_barrier":0,"reference_barrier":0,"exposure":1,"through":1}; continue
        rel=[]
        for tid,w in pack["route_ttp"][rid]:
            row=pack["ttp"][tid]; actor_ok=bool(int(float(row.get(actor["name"]) or 0))); scenario_ok=bool(int(float(row.get(sid) or 0))); stage_ok=str(row.get("Stage") or "")==stage; relevant=actor_ok and scenario_ok and stage_ok
            if collect and stage_ok:
                ab,rb=_barrier(tid,{"PREVENT_SUCCESS","CONTAIN_SUCCESS"},controls,pack); audit.append([actor["name"],pack["scenario_names"][sid],route["name"],tid,row.get("Technique"),stage,actor_ok,scenario_ok,relevant,w,ab,rb,row.get("Platforms"),row.get("MITRE URL")])
            if relevant: rel.append((tid,w)); all_rel.append((tid,w))
        if not rel: valid=False; stages[stage]={"required":True,"reference":pack["stage_priors"][rid][stage],"actual_barrier":0,"reference_barrier":0,"exposure":1,"through":0}; continue
        total=sum(w for _,w in rel); ab=sum(w*_barrier(t,{"PREVENT_SUCCESS","CONTAIN_SUCCESS"},controls,pack)[0] for t,w in rel)/total; rb=sum(w*_barrier(t,{"PREVENT_SUCCESS","CONTAIN_SUCCESS"},controls,pack)[1] for t,w in rel)/total
        ratio=min(max((1-ab)/max(1-rb,1e-9),pack["settings"]["CONTROL_ADJ_FLOOR"]),pack["settings"]["CONTROL_ADJ_CEILING"]); x=exp["stage"][stage] if use else 1; ref=pack["stage_priors"][rid][stage]; through=min(max(ref*ratio*x,0),1)
        stages[stage]={"required":True,"reference":ref,"actual_barrier":ab,"reference_barrier":rb,"exposure":x,"through":through}
    def factor(effect,lo,hi):
        mapped=[(t,w) for t,w in all_rel if any(mapped_effect==effect for _,_,mapped_effect in pack["control_ttp"].get(t,[]))]
        total=sum(w for _,w in mapped)
        return sum(w*_channel(t,effect,controls,pack,lo,hi) for t,w in mapped)/total if total else 1
    duration=factor("REDUCE_DURATION",pack["settings"]["DURATION_FACTOR_FLOOR"],pack["settings"]["CONSEQUENCE_FACTOR_CEILING"]); exposure=factor("REDUCE_EXPOSURE",pack["settings"]["EXPOSURE_FACTOR_FLOOR"],pack["settings"]["CONSEQUENCE_FACTOR_CEILING"]); consequence=factor("REDUCE_CONSEQUENCE",pack["settings"]["CONSEQUENCE_FACTOR_FLOOR"],pack["settings"]["CONSEQUENCE_FACTOR_CEILING"])
    return {"valid":valid,"success":math.prod(v["through"] for v in stages.values()) if valid else 0,"opportunity":(exp["opportunity"] if use else 1) if feasible else 0,"impact_factor":exp["impact"] if use else 1,"duration_factor":duration,"exposure_factor":exposure,"consequence_factor":consequence,"stages":stages,"ttp_rows":audit}


def _severity(main,pack):
    org=main["org"]; revenue=_num(org.get("ANNUAL_REVENUE_AT_RISK"),"revenue",0); payment=_num(org.get("ANNUAL_PAYMENT_VALUE"),"payment value",0); bi=_num(org.get("BI_LOSS_FACTOR"),"BI factor",0); records=_num(org.get("SENSITIVE_RECORDS"),"records",0); endpoints=_num(org.get("CRITICAL_ENDPOINTS"),"endpoints",0); servers=_num(org.get("CRITICAL_SERVERS"),"servers",0)
    unknown=assert_pack_drivers_are_balbix(pack["impact_matrix"])
    if unknown:
        raise ValueError(f"Sector pack contains non-Balbix impact drivers: {unknown}")
    out={}; audit=[]
    for actor in pack["actors"]:
        for sid in pack["scenarios"]:
            driver_p={did:[0.0,0.0] for did in pack["impact_matrix"]}; p=main["params"][sid]
            for qi,q in enumerate(("P50","P99")):
                rate=lambda k:main["rates"][k]["used_p50" if q=="P50" else "used_p99"]; val=lambda k:p[k+"_"+q]
                amounts={
                    "IR_FORENSICS":impact_formulas.daily_rate_x_duration(rate("FORENSIC_DAILY"),val("IR_DAYS")),
                    "EXTERNAL_RESPONSE":impact_formulas.daily_rate_x_duration(rate("EXTERNAL_CYBER_DAILY"),val("CYBER_DAYS")),
                    "LEGAL_RESPONSE":impact_formulas.daily_rate_x_duration(rate("LEGAL_DAILY"),val("LEGAL_DAYS")),
                    "RESTORATION":impact_formulas.daily_rate_x_duration(rate("RESTORATION_DAILY"),val("RESTORE_DAYS")),
                    "BUSINESS_INTERRUPTION":impact_formulas.revenue_per_day_x_duration_x_share_x_margin(revenue,val("DOWNTIME_DAYS"),val("AFFECTED_SERVICE_SHARE"),bi),
                    "ENDPOINT_RECOVERY":impact_formulas.asset_count_x_share_x_unit(endpoints,val("AFFECTED_ENDPOINT_SHARE"),rate("ENDPOINT_REPAIR")),
                    "SERVER_RECOVERY":impact_formulas.asset_count_x_share_x_unit(servers,val("AFFECTED_SERVER_SHARE"),rate("SERVER_REPAIR")),
                    "NOTIFICATION":impact_formulas.records_x_share_x_unit(records,val("AFFECTED_RECORD_SHARE"),rate("NOTIFICATION_PER_RECORD")),
                    "CREDIT_MONITORING":impact_formulas.records_x_share_x_takeup_x_unit(records,val("AFFECTED_RECORD_SHARE"),val("MONITORING_TAKEUP"),rate("CREDIT_MONITORING_PER_RECORD")),
                    "REGULATORY":impact_formulas.revenue_x_fine_percentage(revenue,val("REGULATORY_REVENUE_SHARE")),
                    "CUSTOMER_ATTRITION":impact_formulas.revenue_x_fine_percentage(revenue,val("CHURN_REVENUE_SHARE")),
                    "FRAUD_NET_RECOVERY":impact_formulas.payment_value_x_compromised_x_unrecovered(payment,val("PAYMENT_FLOW_DAYS"),val("DIVERTED_SHARE"),1-val("FRAUD_RECOVERY_RATE")),
                    "EXTORTION":impact_formulas.direct_amount(val("EXTORTION_DIRECT")),
                    "POST_EVENT_UPLIFT":impact_formulas.revenue_x_fine_percentage(revenue,val("POST_EVENT_REVENUE_SHARE")),
                }
                for did,spec in pack["impact_matrix"].items():
                    amount=amounts.get(did,0)
                    applicable=spec["actor"][actor["name"]] and spec["scenario"][sid]; used=max(float(amount),0) if applicable else 0; driver_p[did][qi]=used; audit.append([actor["name"],sid,pack["scenario_names"][sid],spec["block"],did,q,used,applicable,spec["grade"],spec["source"]])
            fitted={}
            for did,(p50,p99) in driver_p.items():
                if p50<=0 and p99<=0: fitted[did]={"p50":0,"p99":0,"mu":None,"sigma":0,"mean":0,"block":pack["impact_matrix"][did]["block"]}
                else:
                    mu,sigma,mean=_lognormal(p50,p99); fitted[did]={"p50":p50,"p99":p99,"mu":mu,"sigma":sigma,"mean":mean,"block":pack["impact_matrix"][did]["block"]}
            out[(actor["id"],sid)]=fitted
    return out,audit


def _cells(main,pack,severity,override=None,collect=False):
    ref=_num(pack["meta"].get("REFERENCE_CAMPAIGNS_PER_YEAR"),"reference campaigns",0); org=_num(main["freq"].get("ORGANISATION_EXPOSURE_MULT"),"organisation multiplier",0); geo=_num(main["freq"].get("GEOGRAPHY_THREAT_MULT"),"geography multiplier",0); weighted=sum(a["weight"]*_num(main["freq"].get(ACTIVITY[a["name"]]),ACTIVITY[a["name"]],0) for a in pack["actors"]); base=ref*org*geo*weighted; shares={a["id"]:a["weight"]*_num(main["freq"].get(ACTIVITY[a["name"]]),ACTIVITY[a["name"]],0)/weighted for a in pack["actors"]}
    cells=[]; ttp=[]
    for actor in pack["actors"]:
        aid=actor["id"]
        for sid,sp in pack["actor_scenario"][aid].items():
            paths={r["id"]:_path(actor,sid,r,main,pack,override,collect) for r in pack["routes"]}
            for p in paths.values(): ttp.extend(p["ttp_rows"])
            raw={r["id"]:pack["actor_route"][aid][r["id"]]*pack["scenario_route"][sid][r["id"]] if paths[r["id"]]["valid"] else 0 for r in pack["routes"]}; denom=sum(raw.values())
            if denom<=0: continue
            for route in pack["routes"]:
                rid=route["id"]; path=paths[rid]; attempt=base*shares[aid]*sp*(raw[rid]/denom)*path["opportunity"]; factors={"Duration":path["duration_factor"]*path["impact_factor"],"Exposure":path["exposure_factor"]*path["impact_factor"],"Direct":path["consequence_factor"]*path["impact_factor"]}; drivers=[]; blocks_by_channel=defaultdict(lambda:{"p50":0.0,"p99":0.0,"mean":0.0,"mu_parts":[],"sigma":None})
                for did,dist in severity[(aid,sid)].items():
                    if dist["mean"]<=0 or dist["mu"] is None: continue
                    block=dist["block"]; f=factors[block]
                    drivers.append({"driver_id":did,"block":block,"p50":dist["p50"]*f,"p99":dist["p99"]*f,"mu":dist["mu"]+math.log(max(f,1e-12)),"sigma":dist["sigma"],"mean":dist["mean"]*f})
                    bb=blocks_by_channel[block]; bb["p50"]+=dist["p50"]*f; bb["p99"]+=dist["p99"]*f; bb["mean"]+=dist["mean"]*f
                # Diagnostic block rows for sensitivity scaling only (not Balbix categories).
                blocks=[]
                for block,bb in blocks_by_channel.items():
                    if bb["mean"]<=0: continue
                    # Refit channel aggregate for CRN-compatible block_scale shocks
                    mu,sigma,mean=_lognormal(bb["p50"],bb["p99"]) if bb["p50"]>0 or bb["p99"]>0 else (None,0,0)
                    if mean<=0 or mu is None: continue
                    blocks.append({"block":block,"p50":bb["p50"],"p99":bb["p99"],"mu":mu,"sigma":sigma,"mean":mean})
                mean=sum(d["mean"] for d in drivers); cells.append({"actor":actor["name"],"actor_id":aid,"scenario":pack["scenario_names"][sid],"scenario_id":sid,"route":route["name"],"route_id":rid,"attempt":attempt,"success":path["success"],"event_rate":attempt*path["success"],"p50":sum(d["p50"] for d in drivers),"p99":sum(d["p99"] for d in drivers),"mean":mean,"aal":attempt*path["success"]*mean,"drivers":drivers,"blocks":blocks,"path":path})
    return cells,ttp,{"reference_lambda":ref,"weighted_activity":weighted,"lambda_base":base}


def _simulate(cells,n,seed,factor,actors,scenarios,rho,block_scale=None,actor_scale=None):
    """Monte Carlo annual aggregates with channel-comonotonic driver dependency.

    Dependency treatment (documented working approach):
    - Within a successful event, one latent standard normal is drawn per control
      channel (Duration, Exposure, Direct).
    - Channels share an equicorrelated Gaussian factor with correlation ``rho``
      (= sector-pack ``LOSS_BLOCK_RHO``), matching the pre-unified three-block model.
    - Every impact driver in a channel uses that channel's latent (comonotonic in
      Gaussian space): loss_d = exp(mu_d + sigma_d * Z_channel). There is no
      within-channel idiosyncratic shock, so related drivers do not diversify away
      merely because they are reported separately.
    - This is not a calibrated pairwise correlation matrix or copula; Balbix does
      not define one. It is the structural default that keeps within-event economic
      co-movement while preserving per-driver identity for reporting.
    """
    rng=np.random.default_rng(seed); annual=np.zeros(n); oep=np.zeros(n); actor={a:np.zeros(n) for a in actors}; scenario={s:np.zeros(n) for s in scenarios}
    blocks_annual={b:np.zeros(n) for b in ("Duration","Exposure","Direct")}; drivers_annual={}; block_scale=block_scale or {}; actor_scale=actor_scale or {}
    channels=("Duration","Exposure","Direct")
    for cell in cells:
        rate=cell["event_rate"]*factor*float(actor_scale.get(cell["actor"],1.0))
        counts=rng.poisson(rate,n); total=int(counts.sum())
        if not total: continue
        years=np.repeat(np.arange(n,dtype=np.int32),counts); common=rng.standard_normal(total); losses=np.zeros(total)
        # Prefer per-driver draws; fall back to legacy block draws for unit tests.
        draw_specs=cell.get("drivers") or [{"driver_id":b.get("block") or "Direct","block":b.get("block") or "Direct","mu":b["mu"],"sigma":b["sigma"]} for b in cell.get("blocks") or []]
        # Shared channel latents (comonotonic drivers within channel).
        channel_z={}
        for ch in channels:
            channel_z[ch]=math.sqrt(rho)*common+math.sqrt(max(1.0-rho,0.0))*rng.standard_normal(total)
        for spec in draw_specs:
            block_name=spec.get("block") or "Direct"
            scale=float(block_scale.get(block_name,1.0))
            z=channel_z.get(block_name)
            if z is None:
                z=math.sqrt(rho)*common+math.sqrt(max(1.0-rho,0.0))*rng.standard_normal(total)
            draw=np.exp(spec["mu"]+spec["sigma"]*z)*scale
            losses+=draw
            did=spec.get("driver_id") or block_name
            if did not in drivers_annual:
                drivers_annual[did]=np.zeros(n)
            drivers_annual[did]+=np.bincount(years,weights=draw,minlength=n)
            if block_name in blocks_annual:
                blocks_annual[block_name]+=np.bincount(years,weights=draw,minlength=n)
        contribution=np.bincount(years,weights=losses,minlength=n); annual+=contribution; actor[cell["actor"]]+=contribution; scenario[cell["scenario"]]+=contribution; np.maximum.at(oep,years,losses)
    return {"annual":annual,"oep":oep,"actor":actor,"scenario":scenario,"blocks":blocks_annual,"drivers":drivers_annual,"metrics":_metrics(annual),"dependency":{"model":"channel_comonotonic_gaussian","rho":float(rho),"within_channel":"comonotonic","across_channel":"equicorrelated_gaussian"}}


def _clear(ws,area):
    for row in ws[area]:
        for cell in row: cell.value=None


def _write(ws,start,rows):
    for r,row in enumerate(rows,start):
        for c,v in enumerate(row,1): ws.cell(r,c,v)


def _head(ws,row,cols):
    apply_column_headers(ws, row, cols, ROLE_CALC)
    apply_freeze_at_headers(ws, row)


def _outputs(wb,main,pack,audit,cells,ttp,freq,views,whatifs,tests):
    primary=_reporting_view(main["tail"].get("REPORTING_VIEW")); best,prudent,pview=views["Best Estimate"],views["Prudent"],views[primary]; pf=_num(main["freq"].get("PRUDENCE_FREQUENCY_FACTOR"),"prudence",0); actors=[a["name"] for a in pack["actors"]]; scenarios=[pack["scenario_names"][s] for s in pack["scenarios"]]
    ws=wb["03 Organisation Inputs"]; ws["C34"]=pack["meta"]["PACK_ID"]; ws["C35"]=pack["meta"]["PACK_VERSION"]; ws["C36"]=pack["meta"]["STATUS"]
    ws=wb["04 Exposure Adjustments"]; use=main["use_exposure"]=="Yes"
    for r in range(16,116):
        rid=str(ws.cell(r,1).value or "").strip()
        if not rid: continue
        x=main["exposure"][rid]; applicable=str(x["org_applicable"] or "").lower()!="no"; ws.cell(r,16,(x["opportunity"] if use else 1) if applicable else 0)
        for i,s in enumerate(STAGES,17): ws.cell(r,i,x["stage"][s] if use else 1)
        ws.cell(r,22,x["impact"] if use else 1)
    ws=wb["06 Impact and BIA"]
    for r,key in enumerate(["ANNUAL_REVENUE_AT_RISK","ANNUAL_PAYMENT_VALUE","BI_LOSS_FACTOR","CUSTOMERS","SENSITIVE_RECORDS","CRITICAL_ENDPOINTS","CRITICAL_SERVERS"],7): ws.cell(r,2,main["org"].get(key))
    for (typ,sid,pid),r in main["assumption_rows"].items():
        if typ=="RATE": x=main["rates"][pid]; vals=(x["p50"],x["p99"],x["used_p50"],x["used_p99"])
        else: vals=(pack["scenario_params"][sid][pid+"_P50"],pack["scenario_params"][sid][pid+"_P99"],main["params"][sid][pid+"_P50"],main["params"][sid][pid+"_P99"])
        for c,v in zip((7,8,11,12),vals): ws.cell(r,c,v)
    ws=wb["07 Frequency Assumptions"]
    for r,v in zip(range(26,31),[freq["reference_lambda"],freq["weighted_activity"],freq["lambda_base"],freq["lambda_base"]*pf,sum(c["attempt"] for c in cells)/max(freq["lambda_base"],1e-12)]): ws.cell(r,2,v)
    ws=wb["10 TTP Relevance Calc"]; _clear(ws,"A15:P800"); h=["Threat actor","Scenario","Route","TTP ID","Technique","Stage","Actor applicable","Scenario applicable","Relevant","Weight","Actual barrier","Reference barrier","Platforms","MITRE URL"]; _write(ws,15,[h]+ttp); _head(ws,15,len(h))
    ws=wb["11 Attack Path Calc"]; _clear(ws,"A15:Z250"); h=["Threat actor","Scenario","Route","Path valid","Attempts / yr","Opportunity","P(success)","Duration factor","Exposure factor","Direct factor","Impact factor"]
    for s in STAGES: h += [f"{s} required",f"{s} reference",f"{s} actual barrier",f"{s} reference barrier",f"{s} exposure",f"{s} through"]
    rows=[]
    for c in cells:
        p=c["path"]; row=[c["actor"],c["scenario"],c["route"],p["valid"],c["attempt"],p["opportunity"],c["success"],p["duration_factor"],p["exposure_factor"],p["consequence_factor"],p["impact_factor"]]
        for s in STAGES:
            v=p["stages"][s]; row += [v["required"],v["reference"],v["actual_barrier"],v["reference_barrier"],v["exposure"],v["through"]]
        rows.append(row)
    _write(ws,15,[h]+rows); _head(ws,15,len(h))
    ws=wb["12 Frequency and Success"]; _clear(ws,"A15:P250"); h=["Threat actor","Scenario","Route","Attempts / yr","P(success)","Events / yr","P50 loss severity per successful event","P99 loss severity per successful event","Expected event loss","AAL","Exposure model"]; _write(ws,15,[h]+[[c["actor"],c["scenario"],c["route"],c["attempt"],c["success"],c["event_rate"],c["p50"],c["p99"],c["mean"],c["aal"],main["use_exposure"]] for c in cells]); _head(ws,15,len(h))
    ws=wb["13 Consequence Calc"]; _clear(ws,"A15:P500"); h=["Threat actor","Scenario ID","Scenario","Loss block","Driver","Quantile","Amount","Applicable","Evidence grade","Source ID"]; _write(ws,15,[h]+audit); _head(ws,15,len(h))
    factor=pf if primary=="Prudent" else 1; matrix=np.zeros((len(actors),len(scenarios)))
    for c in cells: matrix[actors.index(c["actor"]),scenarios.index(c["scenario"])]+=c["aal"]*factor
    ws=wb["14 Actor Scenario Matrix"]; _clear(ws,"A15:F40"); _write(ws,15,[[f"AAL — {primary}",*scenarios]]+[[actors[i],*matrix[i].tolist()] for i in range(len(actors))]); _head(ws,15,5)
    ws=wb["15 Tail Risk Metrics"]; _clear(ws,"A15:J60"); h=["Metric","Best Estimate","Prudent","Definition"]; defs=[(LABEL_AAL,"AAL","Mean annual aggregate loss including zero-loss years"),(LABEL_VAR_95,"VaR95",f"95th percentile of annual aggregate loss (quantile method={VAR_QUANTILE_METHOD}; ≈ 1-in-20)"),(LABEL_TVAR_95,"TVaR95","Mean of the worst 5% of annual aggregate trials"),(LABEL_VAR_99,"VaR99",f"99th percentile of annual aggregate loss (quantile method={VAR_QUANTILE_METHOD}; ≈ 1-in-100)"),(LABEL_TVAR_99,"TVaR99","Mean of the worst 1% of annual aggregate trials"),("P(any event)","PAny","Positive-loss year share")]; _write(ws,15,[h]+[[a,best["metrics"][k],prudent["metrics"][k],d] for a,k,d in defs]); _head(ws,15,4)
    ws=wb["16 Control What-If"]; _clear(ws,"A15:Y80"); h=["Control ID","Control family","Current maturity","What-if maturity","Baseline AAL","What-If AAL","AAL Reduction","AAL Reduction %","Baseline AAL","What-If AAL","AAL Reduction","AAL Reduction %","Method","View"]; _write(ws,15,[h]+[[w["cid"],w["name"],w["current"],w["next"],w["be"]["baseline"],w["be"]["aal"],w["be"]["reduction"],w["be"]["pct"],w["prudent"]["baseline"],w["prudent"]["aal"],w["prudent"]["reduction"],w["prudent"]["pct"],"One-level analytic uplift",primary] for w in whatifs]); _head(ws,15,14)
    lec=[]
    for rp in RETURN_PERIODS:
        q=1-1/rp; lec.append([rp,1/rp,float(np.quantile(best["annual"],q,method="higher")),float(np.quantile(best["oep"],q,method="higher")),float(np.quantile(prudent["annual"],q,method="higher")),float(np.quantile(prudent["oep"],q,method="higher"))])
    ws=wb["17 Aggregate LECs"]; _clear(ws,"A15:I40"); h=["Return period","Exceedance probability","Best AEP","Best OEP","Prudent AEP","Prudent OEP"]; _write(ws,15,[h]+lec); _head(ws,15,6)
    ws=wb["18 Scenario LECs"]; _clear(ws,"A15:K40"); _write(ws,15,[["Return period","Exceedance probability",*scenarios]]+[[rp,1/rp,*[float(np.quantile(pview["scenario"][s],1-1/rp,method="higher")) for s in scenarios]] for rp in RETURN_PERIODS]); _head(ws,15,6)
    ws=wb["19 Actor LECs"]; _clear(ws,"A15:G40"); _write(ws,15,[["Return period","Exceedance probability",*actors]]+[[rp,1/rp,*[float(np.quantile(pview["actor"][a],1-1/rp,method="higher")) for a in actors]] for rp in RETURN_PERIODS]); _head(ws,15,6)
    ws=wb["02 Executive Summary"]; selected="TVaR99" if normalize_tail_basis(main["tail"].get("TVAR_SELECTION"))==LABEL_TVAR_99 else "TVaR95"; retention=_num(main["tail"].get("INSURANCE_RETENTION"),"retention",0)
    for r,k in enumerate(["AAL","VaR95","TVaR95","VaR99","TVaR99"],7): ws.cell(r,2,best["metrics"][k]); ws.cell(r,3,prudent["metrics"][k])
    ws.cell(12,2,best["metrics"][selected]); ws.cell(12,3,prudent["metrics"][selected]); ws.cell(13,2,retention); ws.cell(13,3,retention); ws.cell(14,2,max(best["metrics"][selected]-retention,0)); ws.cell(14,3,max(prudent["metrics"][selected]-retention,0))
    attempts=sum(c["attempt"] for c in cells); events=sum(c["event_rate"] for c in cells); now=datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"); ctx=[(best["metrics"]["PAny"],prudent["metrics"]["PAny"]),(events,events*pf),(attempts,attempts*pf),(main["use_exposure"],main["use_exposure"]),(pack["meta"]["SECTOR"],pack["meta"]["SECTOR"]),(pack["meta"]["PACK_ID"],pack["meta"]["PACK_ID"]),(pack["meta"]["STATUS"],pack["meta"]["STATUS"]),(len(best["annual"]),len(prudent["annual"])),(now,now)]
    for r,(a,b) in enumerate(ctx,18): ws.cell(r,2,a); ws.cell(r,3,b)
    total=pview["metrics"]["AAL"]; vf=pf if primary=="Prudent" else 1; actor_rows=[]; scenario_rows=[]
    for name in actors:
        arr=pview["actor"][name]; aal=float(arr.mean()); att=sum(c["attempt"]*vf for c in cells if c["actor"]==name); ev=sum(c["event_rate"]*vf for c in cells if c["actor"]==name); actor_rows.append([name,aal,att,ev,ev/att if att else 0,aal/total if total else 0,float(np.quantile(arr,.95,method="higher")),float(np.quantile(arr,.99,method="higher"))])
    for name in scenarios:
        arr=pview["scenario"][name]; aal=float(arr.mean()); att=sum(c["attempt"]*vf for c in cells if c["scenario"]==name); ev=sum(c["event_rate"]*vf for c in cells if c["scenario"]==name); scenario_rows.append([name,aal,att,ev,ev/att if att else 0,aal/total if total else 0,float(np.quantile(arr,.95,method="higher")),float(np.quantile(arr,.99,method="higher"))])
    _write(ws,30,actor_rows); _write(ws,37,scenario_rows)
    for r,w in enumerate(whatifs[:5],44): _write(ws,r,[[w["name"],w["current"],w["next"],w["current_aal"],w["whatif_aal"],w["reduction"]]])
    ws=wb["20 Risk Charts"]; ws._charts=[]; _clear(ws,"A15:H45"); _write(ws,15,[["Threat actor","AAL"]]+[[r[0],r[1]] for r in actor_rows]); _write(ws,22,[["Scenario","AAL"]]+[[r[0],r[1]] for r in scenario_rows]); _write(ws,29,[["Return period","Best AEP","Prudent AEP"]]+[[r[0],r[2],r[4]] for r in lec])
    bar=BarChart(); bar.type="bar"; bar.title=f"AAL by threat actor — {primary}"; bar.add_data(Reference(ws,min_col=2,min_row=15,max_row=19),titles_from_data=True); bar.set_categories(Reference(ws,min_col=1,min_row=16,max_row=19)); ws.add_chart(bar,"D6")
    bar2=BarChart(); bar2.type="bar"; bar2.title=f"AAL by scenario — {primary}"; bar2.add_data(Reference(ws,min_col=2,min_row=22,max_row=26),titles_from_data=True); bar2.set_categories(Reference(ws,min_col=1,min_row=23,max_row=26)); ws.add_chart(bar2,"D22")
    line=LineChart(); line.title="Annual Aggregate Loss Exceedance Curve"; line.add_data(Reference(ws,min_col=2,max_col=3,min_row=29,max_row=39),titles_from_data=True); line.set_categories(Reference(ws,min_col=1,min_row=30,max_row=39)); ws.add_chart(line,"L6")
    ws=wb["22 Validation Tests"]; _clear(ws,"A16:E100"); _write(ws,16,tests); _head(ws,15,5)


def _optional_cost(value):
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def refresh(input_path,output_path=None,sector_pack_dir=None,appetite_inputs=None,insurance_programme=None,run_packages=True,run_sensitivity=False,run_whatifs=True):
    input_path=Path(input_path).resolve(); output_path=Path(output_path or input_path).resolve(); wb=openpyxl.load_workbook(input_path,data_only=False); pack=_load_pack(wb,input_path,sector_pack_dir); _prepare(wb,pack); main=_load_main(wb,pack); severity,audit=_severity(main,pack); cells,ttp,freq=_cells(main,pack,severity,collect=True)
    n=int(_num(main["org"].get("SIMULATION_YEARS"),"simulation years",10000,2000000)); seed=int(_num(main["org"].get("RANDOM_SEED"),"seed",0)); pf=_num(main["freq"].get("PRUDENCE_FREQUENCY_FACTOR"),"prudence",0); actors=[a["name"] for a in pack["actors"]]; scenarios=[pack["scenario_names"][s] for s in pack["scenarios"]]; rho=pack["settings"]["LOSS_BLOCK_RHO"]
    views={"Best Estimate":_simulate(cells,n,seed,1,actors,scenarios,rho),"Prudent":_simulate(cells,n,seed,pf,actors,scenarios,rho)}
    best, prudent = views["Best Estimate"], views["Prudent"]
    primary=_reporting_view(main["tail"].get("REPORTING_VIEW")); aal_be=sum(c["aal"] for c in cells); aal_pr=aal_be*pf; freq_be=sum(c["event_rate"] for c in cells); freq_pr=freq_be*pf; whatifs=[]
    maturity_before={cid:control["maturity"] for cid,control in main["controls"].items()}

    def effective_level(c):
        m=c.get("maturity") or "Not Assessed"
        return "Developing" if m=="Not Assessed" else m

    def next_level(level):
        if level not in MATURITY_ORDER: return "Managed"
        return MATURITY_ORDER[min(MATURITY_ORDER.index(level)+1,len(MATURITY_ORDER)-1)]

    for cid,control in main["controls"].items():
        cur=effective_level(control); nxt=next_level(cur)
        mapped=any(t[0]==cid for maps in pack["control_ttp"].values() for t in maps)
        if (not run_whatifs) or (not mapped) or nxt==cur:
            be=metric_bundle_from_sim(best,best["metrics"],freq_be,freq_be)
            pr=metric_bundle_from_sim(prudent,prudent["metrics"],freq_pr,freq_pr)
            for d in (be,pr):
                d["aal"]=d["baseline"]; d["reduction"]=0; d["pct"]=0; d["event_freq"]=d["freq_base"]; d["freq_change"]=0
                for k in ("var95","var99","tvar95","tvar99"):
                    d[k]=d[f"{k}_base"]; d[f"{k}_change"]=0
        else:
            wi_cells=_cells(main,pack,severity,{cid:nxt})[0]
            wi_freq_be=sum(c["event_rate"] for c in wi_cells); wi_freq_pr=wi_freq_be*pf
            sim_be=_simulate(wi_cells,n,seed,1,actors,scenarios,rho); sim_pr=_simulate(wi_cells,n,seed,pf,actors,scenarios,rho)
            be=metric_bundle_from_sim(sim_be,best["metrics"],wi_freq_be,freq_be)
            pr=metric_bundle_from_sim(sim_pr,prudent["metrics"],wi_freq_pr,freq_pr)
        costs=treatment_cost_metrics(
            aal_reduction=pr["reduction"] if primary=="Prudent" else be["reduction"],
            one_off_cost=_optional_cost(control.get("one_off_cost") or control.get("implementation_cost")),
            annual_cost=_optional_cost(control.get("annual_cost") or control.get("recurring_cost")),
            evaluation_years=_optional_cost(control.get("evaluation_years") or main.get("tail",{}).get("EVALUATION_YEARS")),
        )
        whatifs.append({"cid":cid,"name":control["name"],"channel":control.get("family") or control["name"],"current":control["maturity"],"next":nxt,"whatif":nxt,"be":be,"prudent":pr,"reduction":pr["reduction"] if primary=="Prudent" else be["reduction"],"current_aal":prudent["metrics"]["AAL"] if primary=="Prudent" else best["metrics"]["AAL"],"whatif_aal":pr["aal"] if primary=="Prudent" else be["aal"],"reduction_pct":pr["pct"] if primary=="Prudent" else be["pct"],"mapped":mapped,"costs":costs})
    whatifs.sort(key=lambda x:x["reduction"],reverse=True)

    # Combined control packages — genuine multi-control re-simulation (not summed what-ifs)
    packages=[]
    if run_packages:
        for pname in ("Foundation","Priority","Target","User-defined"):
            ov=package_maturity_overrides(main["controls"],package=pname,levels=MATURITY_ORDER,effective_level=effective_level,next_level=next_level,whatifs=whatifs)
            if not ov: continue
            pkg_cells=_cells(main,pack,severity,ov)[0]
            pkg_freq=sum(c["event_rate"] for c in pkg_cells)*(pf if primary=="Prudent" else 1)
            pkg_sim=_simulate(pkg_cells,n,seed,(pf if primary=="Prudent" else 1),actors,scenarios,rho)
            base_m=views[primary]["metrics"]
            packages.append({
                "name":pname,"overrides":{k:ov[k] for k in list(ov)[:40]},"n_controls":len(ov),
                "event_freq":pkg_freq,"aal":pkg_sim["metrics"]["AAL"],"var95":pkg_sim["metrics"]["VaR95"],
                "tvar95":pkg_sim["metrics"]["TVaR95"],"var99":pkg_sim["metrics"]["VaR99"],"tvar99":pkg_sim["metrics"]["TVaR99"],
                "baseline_aal":base_m["AAL"],"baseline_var95":base_m["VaR95"],"baseline_tvar95":base_m["TVaR95"],
                "baseline_var99":base_m["VaR99"],"baseline_tvar99":base_m["TVaR99"],
                "baseline_event_freq":freq_pr if primary=="Prudent" else freq_be,
                "aal_reduction":base_m["AAL"]-pkg_sim["metrics"]["AAL"],
                "tvar99_reduction":base_m["TVaR99"]-pkg_sim["metrics"]["TVaR99"],
                "simulated":True,
                "costs":treatment_cost_metrics(aal_reduction=base_m["AAL"]-pkg_sim["metrics"]["AAL"],one_off_cost=None,annual_cost=None,evaluation_years=None),
            })

    tests=[]
    def add(name,ok,actual,expected,fix): tests.append([name,"PASS" if ok else "FAIL",actual,expected,fix])
    add("Sector-pack schema",pack["meta"]["SCHEMA_VERSION"]==PACK_SCHEMA_VERSION,pack["meta"]["SCHEMA_VERSION"],PACK_SCHEMA_VERSION,"Update pack or engine.")
    add("Evidence grades valid",all(x["grade"] in GRADES for x in pack["driver_rates"].values()) and all(x["grade"] in GRADES for x in pack["impact_matrix"].values()),"Valid","A-D","Correct pack grade.")
    neutral=main["use_exposure"]=="Yes" or all(abs(c["path"]["opportunity"]-1)<1e-12 and abs(c["path"]["impact_factor"]-1)<1e-12 and all(abs(v["exposure"]-1)<1e-12 for v in c["path"]["stages"].values()) for c in cells); add("Exposure switch neutralisation",neutral,"Neutral" if neutral else "Non-neutral","Neutral when No","Review exposure interface.")
    attempts=sum(c["attempt"] for c in cells); add("Baseline attempt frequency preserved",main["use_exposure"]=="Yes" or abs(attempts-freq["lambda_base"])<1e-9,attempts,freq["lambda_base"],"Review route renormalisation.")
    pairs=[(a["id"],s) for a in pack["actors"] for s,p in pack["actor_scenario"][a["id"]].items() if p>0]; covered=all(any(c["actor_id"]==a and c["scenario_id"]==s and c["path"]["valid"] for c in cells) for a,s in pairs); add("Positive actor-scenario cells have valid paths",covered,"Covered" if covered else "Gap","Covered","Review stage and ATT&CK applicability.")
    dead=[]
    for rid,maps in pack["route_ttp"].items():
        for tid,_ in maps:
            if not any(bool(int(float(pack["ttp"][tid].get(a["name"]) or 0))) and bool(int(float(pack["ttp"][tid].get(s) or 0))) for a in pack["actors"] for s in pack["scenarios"]): dead.append(f"{rid}/{tid}")
    add("No dead route-technique mappings",not dead,", ".join(dead) or "None","None","Correct mapping.")
    active_zero=[]
    for (mask_actor,sid,rid),required in pack["stage_requirements"].items():
        if any(required.values()):
            continue
        applicable_actors=[a["id"] for a in pack["actors"]] if mask_actor=="ALL" else [mask_actor]
        if any(pack["actor_route"].get(a,{}).get(rid,0)*pack["scenario_route"].get(sid,{}).get(rid,0)>0 for a in applicable_actors):
            active_zero.append(f"{mask_actor}/{sid}/{rid}")
    add("All-zero stage masks inactive",not active_zero,", ".join(active_zero) or "None","None with non-zero pooled route propensity","Set at least one required stage or zero the route propensity.")
    ordered=all(all(v["p99"]>=v["p50"] for v in drivers.values()) for drivers in severity.values()); add("Impact-driver P99 >= P50",ordered,"Ordered" if ordered else "Not ordered","Ordered","Review impact assumptions.")
    bounded=all(pack["settings"]["DURATION_FACTOR_FLOOR"]-1e-12<=c["path"]["duration_factor"]<=pack["settings"]["CONSEQUENCE_FACTOR_CEILING"]+1e-12 and pack["settings"]["EXPOSURE_FACTOR_FLOOR"]-1e-12<=c["path"]["exposure_factor"]<=pack["settings"]["CONSEQUENCE_FACTOR_CEILING"]+1e-12 for c in cells); add("Control ratios bounded",bounded,"Bounded" if bounded else "Out of bounds","Within pack bounds","Review aggregation.")
    crn=np.array_equal(_simulate(cells,10000,seed,1,actors,scenarios,rho)["annual"],_simulate(cells,10000,seed,1,actors,scenarios,rho)["annual"]); add("Common random numbers reproducible",crn,"Exact match" if crn else "Mismatch","Exact match","Use same seed/order.")
    analytic=sum(c["aal"] for c in cells); sim=views["Best Estimate"]["metrics"]["AAL"]; se=float(np.std(views["Best Estimate"]["annual"], ddof=1)/np.sqrt(n)); add("Simulation AAL reconciles to analytic AAL",abs(sim-analytic)<=max(.06*analytic,2*se,1),sim,analytic,"Increase years or review means (tolerance: 6% of analytic or 2×AAL SE).")
    tailok=all(v["metrics"]["TVaR95"]>=v["metrics"]["VaR95"] and v["metrics"]["TVaR99"]>=v["metrics"]["VaR99"] for v in views.values()); add("TVaR not below VaR",tailok,"Valid" if tailok else "Invalid","TVaR >= VaR","Review tail metrics.")
    sparse=np.array([0.0]*99+[100.0]); sparse_metrics=_metrics(sparse); sparse_ok=abs(sparse_metrics["TVaR95"]-20.0)<1e-12 and sparse_metrics["VaR95"]==0 and sparse_metrics["TVaR95"]>sparse_metrics["AAL"]
    add("Sparse-loss exact TVaR",sparse_ok,sparse_metrics["TVaR95"],20.0,"Restore fractional upper-tail expected shortfall.")
    agg_actual=_agg([.5]*6,.25,.95); agg_expected=.5805092839436838
    add("Diminishing-return control aggregation",abs(agg_actual-agg_expected)<1e-12,agg_actual,agg_expected,"Use remaining-gap OT recursion and barrier cap.")
    add("Reporting-view resolution",_reporting_view("Both")=="Prudent",_reporting_view("Both"),"Prudent","Use one reporting-view resolver for outputs and what-ifs.")
    curveok=True
    for v in views.values():
        a=[np.quantile(v["annual"],1-1/r,method="higher") for r in RETURN_PERIODS]; o=[np.quantile(v["oep"],1-1/r,method="higher") for r in RETURN_PERIODS]; curveok &= all(x<=y for x,y in zip(a,a[1:])) and all(x<=y for x,y in zip(o,o[1:])) and all(x<=y+1e-9 for x,y in zip(o,a))
    add("AEP/OEP valid",curveok,"Valid" if curveok else "Invalid","Monotonic; OEP <= AEP","Review curves.")
    add("Control what-ifs non-increasing",min(w["reduction"] for w in whatifs)>=-1e-6,min(w["reduction"] for w in whatifs),">= 0","Review control channels.")
    actor_sum=sum(float(v.mean()) for v in views["Best Estimate"]["actor"].values()); scenario_sum=sum(float(v.mean()) for v in views["Best Estimate"]["scenario"].values()); add("Actor/scenario AAL reconcile",abs(actor_sum-sim)<1e-6 and abs(scenario_sum-sim)<1e-6,max(abs(actor_sum-sim),abs(scenario_sum-sim)),0,"Review decomposition.")
    finite=all(np.isfinite(v["annual"]).all() for v in views.values()); add("Simulation finite",finite,"Finite" if finite else "Non-finite","Finite","Review inputs.")
    validation="PASS" if all(r[1]=="PASS" for r in tests) else "FAIL"; wb["03 Organisation Inputs"]["C37"]=validation; _outputs(wb,main,pack,audit,cells,ttp,freq,views,whatifs,tests)
    if getattr(wb,"calculation",None) is None: wb.calculation=CalcProperties()
    wb.calculation.calcMode="auto"; wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; _save_safe(wb,output_path)
    best, prudent = views["Best Estimate"], views["Prudent"]
    pview = views[primary]
    actor_aal = {a: float(pview["actor"][a].mean()) for a in actors}
    scenario_aal = {s: float(pview["scenario"][s].mean()) for s in scenarios}
    retention = _num(main["tail"].get("INSURANCE_RETENTION"), "retention", 0)
    selected = "TVaR99" if normalize_tail_basis(main["tail"].get("TVAR_SELECTION")) == LABEL_TVAR_99 else "TVaR95"
    scen_t95 = tvar_tail_contributions(pview["annual"], {s: pview["scenario"][s] for s in scenarios}, 0.95)
    scen_t99 = tvar_tail_contributions(pview["annual"], {s: pview["scenario"][s] for s in scenarios}, 0.99)
    actor_t95 = tvar_tail_contributions(pview["annual"], {a: pview["actor"][a] for a in actors}, 0.95)
    actor_t99 = tvar_tail_contributions(pview["annual"], {a: pview["actor"][a] for a in actors}, 0.99)
    top_aal_name = max(scenario_aal.items(), key=lambda kv: kv[1])[0] if scenario_aal else None
    top_tvar99_name = max(scen_t99.items(), key=lambda kv: kv[1])[0] if scen_t99 else None
    p_any = pview["metrics"]["PAny"]
    appetite = parse_appetite_inputs(appetite_inputs if appetite_inputs is not None else (main.get("appetite") or {}))
    appet = evaluate_appetite(
        appetite=appetite,
        aal=pview["metrics"]["AAL"],
        p_any=p_any,
        tvar95=pview["metrics"]["TVaR95"],
        tvar99=pview["metrics"]["TVaR99"],
        annual_losses=pview["annual"],
    )
    narrative = executive_risk_narrative(
        p_any_1y=p_any,
        p_any_5y=multi_year_event_probability(p_any, 5),
        aal=pview["metrics"]["AAL"],
        var95=pview["metrics"]["VaR95"],
        tvar95=pview["metrics"]["TVaR95"],
        var99=pview["metrics"]["VaR99"],
        tvar99=pview["metrics"]["TVaR99"],
        top_aal_name=top_aal_name,
        top_tvar99_name=top_tvar99_name,
        within_tolerance=(
            True if appet["appetite_status"] == "Within tolerance"
            else False if appet["appetite_status"] == "Above tolerance"
            else None
        ),
        tolerance=appetite.get("ANNUAL_LOSS_TOLERANCE"),
    )

    # Trial-level insurance (shared module with OT)
    programme_raw = insurance_programme if insurance_programme is not None else {"INSURANCE_RETENTION": retention}
    programme = parse_programme(programme_raw)
    insurance_analysis = None
    if programme.active():
        insurance_analysis = apply_programme(pview["annual"], programme)
        insurance_analysis = {
            k: v for k, v in insurance_analysis.items()
            if k not in {"ground_up", "retained", "insured", "residual", "uninsured_above_programme", "layer_recoveries", "programme"}
        }
        insurance_analysis["layers"] = [
            {"name": ly.name, "attachment": ly.attachment, "limit": ly.limit, "coinsurance": ly.coinsurance}
            for ly in programme.layers
        ]
        insurance_analysis["retention"] = programme.retention

    # Loss components: Balbix major categories from per-driver annual trials
    driver_annuals = pview.get("drivers") or {}
    loss_drivers = driver_contribution_rows(pview["annual"], driver_annuals)
    loss_components = category_contribution_rows(pview["annual"], driver_annuals)
    top_comp = leading_tvar99_component(loss_components)
    top_driver = leading_tvar99_component(loss_drivers)
    impact_reconcile = reconcile_impact_hierarchy(pview["annual"], driver_annuals)

    # Actor / scenario analysis with TVaR contributions
    vf = pf if primary == "Prudent" else 1
    scenario_analysis = []
    for s in scenarios:
        arr = pview["scenario"][s]
        att = sum(c["attempt"] * vf for c in cells if c["scenario"] == s)
        ev = sum(c["event_rate"] * vf for c in cells if c["scenario"] == s)
        m = annual_aggregate_metrics(arr)
        scenario_analysis.append({
            "name": s,
            "campaign_frequency": att,
            "successful_event_frequency": ev,
            "annual_event_probability": float(np.mean(arr > 0)),
            "aal": float(arr.mean()),
            "var95": m["VaR95"], "tvar95": m["TVaR95"], "var99": m["VaR99"], "tvar99": m["TVaR99"],
            "pct_total_aal": float(arr.mean()) / pview["metrics"]["AAL"] if pview["metrics"]["AAL"] else 0,
            "contrib_tvar95": scen_t95.get(s, 0.0),
            "contrib_tvar99": scen_t99.get(s, 0.0),
            "primary_operational_driver": "Records / endpoints / service disruption",
            "primary_financial_driver": (
                top_comp["name"] if top_comp else "Balbix impact categories"
            ),
        })
    actor_analysis = []
    for a in actors:
        arr = pview["actor"][a]
        att = sum(c["attempt"] * vf for c in cells if c["actor"] == a)
        ev = sum(c["event_rate"] * vf for c in cells if c["actor"] == a)
        m = annual_aggregate_metrics(arr)
        actor_analysis.append({
            "name": a,
            "campaign_frequency": att,
            "successful_event_frequency": ev,
            "annual_event_probability": float(np.mean(arr > 0)),
            "aal": float(arr.mean()),
            "var95": m["VaR95"], "tvar95": m["TVaR95"], "var99": m["VaR99"], "tvar99": m["TVaR99"],
            "pct_total_aal": float(arr.mean()) / pview["metrics"]["AAL"] if pview["metrics"]["AAL"] else 0,
            "contrib_tvar95": actor_t95.get(a, 0.0),
            "contrib_tvar99": actor_t99.get(a, 0.0),
        })
    actor_scenario_rows = []
    for c in cells:
        actor_scenario_rows.append({
            "actor": c["actor"], "scenario": c["scenario"], "route": c["route"],
            "campaign_frequency": c["attempt"] * vf,
            "applicable_frequency": c["attempt"] * vf,
            "success_probability": c["success"],
            "successful_event_frequency": c["event_rate"] * vf,
            "annual_event_probability": 1.0 - math.exp(-c["event_rate"] * vf) if c["event_rate"] * vf >= 0 else 0.0,
            "aal": c["aal"] * vf,
        })

    # Sensitivity reruns
    sensitivities = []
    if run_sensitivity:
        base_aal = pview["metrics"]["AAL"]
        base_tvar99 = pview["metrics"]["TVaR99"]
        base_factor = pf if primary == "Prudent" else 1.0

        def _sens_row(label, kind, factor, sim):
            m = sim["metrics"]
            return {
                "parameter": label, "kind": kind, "factor": factor, "simulation_years": n,
                "aal": m["AAL"], "var95": m["VaR95"], "tvar95": m["TVaR95"], "var99": m["VaR99"], "tvar99": m["TVaR99"],
                "delta_aal": m["AAL"] - base_aal, "delta_tvar99": m["TVaR99"] - base_tvar99,
                "pct_aal": (m["AAL"] - base_aal) / base_aal if base_aal else None,
                "pct_tvar99": (m["TVaR99"] - base_tvar99) / base_tvar99 if base_tvar99 else None,
                "limitation": None,
            }

        for label, factor in (("Campaign frequency +20%", 1.2), ("Campaign frequency -20%", 0.8)):
            sensitivities.append(_sens_row(label, "frequency", factor, _simulate(cells, n, seed, base_factor * factor, actors, scenarios, rho)))
        for label, factor in (("Business interruption (Duration) +20%", 1.2), ("Business interruption (Duration) -20%", 0.8)):
            sensitivities.append(_sens_row(label, "downtime", factor, _simulate(cells, n, seed, base_factor, actors, scenarios, rho, block_scale={"Duration": factor})))
        for label, factor in (("Incident / restoration (Direct) +20%", 1.2), ("Incident / restoration (Direct) -20%", 0.8)):
            sensitivities.append(_sens_row(label, "restoration", factor, _simulate(cells, n, seed, base_factor, actors, scenarios, rho, block_scale={"Direct": factor})))
        for label, factor in (("Data/privacy exposure +20%", 1.2), ("Data/privacy exposure -20%", 0.8)):
            sensitivities.append(_sens_row(label, "exposure", factor, _simulate(cells, n, seed, base_factor, actors, scenarios, rho, block_scale={"Exposure": factor})))
        # Actor-mix shock: boost first actor by +20% relative weight, renormalize to 100%
        if len(actors) >= 2:
            weights = np.array([sum(c["attempt"] for c in cells if c["actor"] == a) for a in actors], dtype=float)
            if weights.sum() > 0:
                weights = weights / weights.sum()
                boosted = weights.copy()
                boosted[0] *= 1.2
                boosted = boosted / boosted.sum()
                scales = {actors[i]: (boosted[i] / weights[i]) if weights[i] > 0 else 1.0 for i in range(len(actors))}
                sensitivities.append(_sens_row("Actor mix: +20% relative weight on lead actor (renormalised)", "actor_mix", 1.2, _simulate(cells, n, seed, base_factor, actors, scenarios, rho, actor_scale=scales)))
        sensitivities.sort(key=lambda r: abs(r.get("delta_aal") or 0), reverse=True)

    # IT operational diagnostics — applicability gated (omit non-applicable as None)
    revenue = _num(main["org"].get("ANNUAL_REVENUE_AT_RISK"), "revenue", 0) if main["org"].get("ANNUAL_REVENUE_AT_RISK") not in (None, "") else None
    records = _num(main["org"].get("SENSITIVE_RECORDS"), "records", 0) if main["org"].get("SENSITIVE_RECORDS") not in (None, "") else None
    endpoints = _num(main["org"].get("CRITICAL_ENDPOINTS"), "endpoints", 0) if main["org"].get("CRITICAL_ENDPOINTS") not in (None, "") else None
    servers = _num(main["org"].get("CRITICAL_SERVERS"), "servers", 0) if main["org"].get("CRITICAL_SERVERS") not in (None, "") else None
    operational_diagnostics = {
        "downtime_p50": None, "downtime_p95": None, "downtime_p99": None,
        "recovery_p50": None, "recovery_p95": None, "recovery_p99": None,
        "capacity_p50": None, "capacity_p95": None, "capacity_p99": None,
        "records_p50": records, "records_p95": records, "records_p99": records,
        "endpoints_p50": endpoints, "endpoints_p95": endpoints, "endpoints_p99": endpoints,
        "services_p50": servers, "services_p95": servers, "services_p99": servers,
        "revenue_per_day": (revenue / 365.0) if revenue else None,
        "p_down_7": None, "p_down_15": None, "p_down_30": None, "p_down_60": None,
        "p_cap_25": None, "p_cap_50": None, "p_cap_75": None,
    }

    return {
        "output": str(output_path),
        "pack_id": pack["meta"]["PACK_ID"],
        "pack_version": str(pack["meta"].get("PACK_VERSION") or ""),
        "pack_status": pack["meta"]["STATUS"],
        "use_exposure_model": main["use_exposure"],
        "engine_version": ENGINE_VERSION,
        "simulation_years": n,
        "random_seed": seed,
        "reporting_view": primary,
        "detailed_view": primary,
        "AAL": pview["metrics"]["AAL"],
        "VaR95": pview["metrics"]["VaR95"],
        "TVaR95": pview["metrics"]["TVaR95"],
        "VaR99": pview["metrics"]["VaR99"],
        "TVaR99": pview["metrics"]["TVaR99"],
        "P_any": p_any,
        "best_aal": best["metrics"]["AAL"],
        "prudent_aal": prudent["metrics"]["AAL"],
        "best_aal_se": float(np.std(best["annual"], ddof=1) / np.sqrt(n)),
        "prudent_aal_se": float(np.std(prudent["annual"], ddof=1) / np.sqrt(n)),
        "best_var95": best["metrics"]["VaR95"],
        "prudent_var95": prudent["metrics"]["VaR95"],
        "best_tvar95": best["metrics"]["TVaR95"],
        "prudent_tvar95": prudent["metrics"]["TVaR95"],
        "best_var99": best["metrics"]["VaR99"],
        "prudent_var99": prudent["metrics"]["VaR99"],
        "best_tvar99": best["metrics"]["TVaR99"],
        "prudent_tvar99": prudent["metrics"]["TVaR99"],
        "best_pany": best["metrics"]["PAny"],
        "prudent_pany": prudent["metrics"]["PAny"],
        "best_pany_3y": multi_year_event_probability(best["metrics"]["PAny"], 3),
        "prudent_pany_3y": multi_year_event_probability(prudent["metrics"]["PAny"], 3),
        "best_pany_5y": multi_year_event_probability(best["metrics"]["PAny"], 5),
        "prudent_pany_5y": multi_year_event_probability(prudent["metrics"]["PAny"], 5),
        "best_p_exceed_tolerance": evaluate_appetite(appetite=appetite, annual_losses=best["annual"], aal=best["metrics"]["AAL"])["p_exceed_tolerance"],
        "prudent_p_exceed_tolerance": appet["p_exceed_tolerance"],
        "risk_tolerance": appetite.get("ANNUAL_LOSS_TOLERANCE"),
        "appetite_inputs": appetite,
        "appetite_status": appet["appetite_status"],
        "appetite_breaches": appet["appetite_breaches"],
        "top_aal_contributor": top_aal_name,
        "top_tvar99_contributor": top_tvar99_name,
        "top_tvar99_component": None if top_comp is None else top_comp["label"],
        "top_tvar99_component_name": None if top_comp is None else top_comp["name"],
        "top_tvar99_component_contrib": None if top_comp is None else top_comp["contrib_tvar99"],
        "top_tvar99_component_pct": None if top_comp is None else top_comp["pct_tvar99"],
        "top_tvar99_driver": None if top_driver is None else top_driver["label"],
        "top_tvar99_driver_name": None if top_driver is None else top_driver["name"],
        "top_tvar99_driver_contrib": None if top_driver is None else top_driver["contrib_tvar99"],
        "top_tvar99_driver_pct": None if top_driver is None else top_driver["pct_tvar99"],
        "impact_reconcile": impact_reconcile,
        "impact_dependency": {
            "model": "channel_comonotonic_gaussian",
            "rho": float(rho),
            "rho_source": "LOSS_BLOCK_RHO (sector pack)",
            "within_channel": "comonotonic (shared latent Z per Duration/Exposure/Direct)",
            "across_channel": "equicorrelated Gaussian with correlation rho",
            "not_used": "pairwise correlation matrix; calibrated copula; independent driver sampling",
        },        "executive_narrative": narrative,
        "scenario_tvar95_contribution": scen_t95,
        "scenario_tvar99_contribution": scen_t99,
        "actor_tvar95_contribution": actor_t95,
        "actor_tvar99_contribution": actor_t99,
        "event_frequency": float(sum(c["event_rate"] for c in cells)) * vf,
        "best_event_frequency": float(sum(c["event_rate"] for c in cells)),
        "prudent_event_frequency": float(sum(c["event_rate"] for c in cells)) * pf,
        "attempt_frequency": float(sum(c["attempt"] for c in cells)) * vf,
        "best_attempt_frequency": float(sum(c["attempt"] for c in cells)),
        "prudent_attempt_frequency": float(sum(c["attempt"] for c in cells)) * pf,
        "actor_aal": actor_aal,
        "scenario_aal": scenario_aal,
        "actor_aal_best": {a: float(best["actor"][a].mean()) for a in actors},
        "actor_aal_prudent": {a: float(prudent["actor"][a].mean()) for a in actors},
        "scenario_aal_best": {s: float(best["scenario"][s].mean()) for s in scenarios},
        "scenario_aal_prudent": {s: float(prudent["scenario"][s].mean()) for s in scenarios},
        "whatifs": whatifs,
        "whatifs_ran": bool(run_whatifs),
        "control_maturities": maturity_before,
        "validation": validation,
        "tests_pass": sum(r[1]=="PASS" for r in tests),
        "tests_total": len(tests),
        "sector_pack_id": pack["meta"]["PACK_ID"],
        "sector_pack_status": pack["meta"]["STATUS"],
        "control_packages": packages,
        "insurance_analysis": insurance_analysis,
        "sensitivities": sensitivities,
        "formation": {
            "campaigns_per_year": float(sum(c["attempt"] for c in cells)) * vf,
            "applicable_campaigns_per_year": float(sum(c["attempt"] for c in cells)) * vf,
            "conditional_success_probability": (
                float(sum(c["event_rate"] for c in cells)) / float(sum(c["attempt"] for c in cells))
                if sum(c["attempt"] for c in cells) else 0.0
            ),
            "successful_events_per_year": float(sum(c["event_rate"] for c in cells)) * vf,
            "p_any_successful_event": p_any,
            "aal": pview["metrics"]["AAL"],
        },
        "actor_scenario_rows": actor_scenario_rows,
        "scenario_analysis": scenario_analysis,
        "actor_analysis": actor_analysis,
        "loss_components": loss_components,
        "loss_drivers": loss_drivers,
        "operational_diagnostics": operational_diagnostics,
        "annual_revenue_at_risk": revenue,
        "trials_available": True,
        "primary_annual_aal": pview["metrics"]["AAL"],
        "primary_annual_var95": pview["metrics"]["VaR95"],
        "primary_annual_tvar95": pview["metrics"]["TVaR95"],
        "primary_annual_var99": pview["metrics"]["VaR99"],
        "primary_annual_tvar99": pview["metrics"]["TVaR99"],
        "selected_tail_metric": selected,
    }
